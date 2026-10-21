#!/usr/bin/env python3
"""
parametro_excel_layout_final.py

Script para aplicar o layout Excel da Parametro Global a ficheiros .xlsx.

Objetivo:
- Manter o conteudo original do Excel.
- Nao adicionar linhas, colunas, textos ou formulas quando usado com --apply.
- Alterar apenas aspeto visual: cores, fontes, alinhamento, grelha, larguras e alturas.
- Usar a paleta vista nos ficheiros Parametro: barra superior azul-escuro,
  cabecalhos azul, linhas alternadas azul-claro/branco.

Uso:
  python parametro_excel_layout_final.py --apply input.xlsx output.xlsx
  python parametro_excel_layout_final.py --new Template_Parametro_Excel_Vazio.xlsx

Dependencia:
  pip install openpyxl
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# -----------------------------------------------------------------------------
# PALETA PARAMETRO GLOBAL
# -----------------------------------------------------------------------------
NAVY = "111827"        # Barra superior / titulo principal
BLUE = "00A3D7"        # Cabecalhos de tabela
LIGHT_BLUE = "EAF6FB"  # Linhas alternadas
LIGHT_BLUE_2 = "DDF1FA"  # Destaque secundario
WHITE = "FFFFFF"
GRID = "C9D1D9"
TEXT = "111827"
MUTED = "64748B"

FONT_NAME = "Aptos"  # Excel moderno; se nao existir, o Excel substitui automaticamente

THIN_GRID = Side(style="thin", color=GRID)
BORDER_ALL = Border(left=THIN_GRID, right=THIN_GRID, top=THIN_GRID, bottom=THIN_GRID)


# -----------------------------------------------------------------------------
# HELPERS
# -----------------------------------------------------------------------------
def solid(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_color)


def non_empty(value) -> bool:
    return value is not None and str(value).strip() != ""


def get_used_bounds(ws: Worksheet) -> tuple[int, int]:
    """Return actual used bounds, ignoring blank trailing cells as far as possible."""
    max_row = 1
    max_col = 1

    for row in range(1, ws.max_row + 1):
        row_has_content = False
        for col in range(1, ws.max_column + 1):
            if non_empty(ws.cell(row=row, column=col).value):
                row_has_content = True
                max_col = max(max_col, col)
        if row_has_content:
            max_row = row

    return max_row, max_col


def row_values(ws: Worksheet, row: int, max_col: int) -> list[str]:
    values: list[str] = []
    for col in range(1, max_col + 1):
        value = ws.cell(row=row, column=col).value
        if non_empty(value):
            values.append(str(value).strip())
    return values


def is_title_row(ws: Worksheet, row: int, max_col: int) -> bool:
    """Detect rows like 'REVISAO DO BACKLOG' or existing merged title rows."""
    values = row_values(ws, row, max_col)
    if len(values) != 1:
        return False

    text = values[0]
    has_merged_cell = any(rng.min_row <= row <= rng.max_row for rng in ws.merged_cells.ranges)
    looks_like_title = len(text) >= 8 and text.upper() == text
    return has_merged_cell or looks_like_title


def is_header_row(ws: Worksheet, row: int, max_col: int) -> bool:
    """Detect table header rows without modifying data."""
    values = row_values(ws, row, max_col)
    if len(values) < 2:
        return False

    header_keywords = (
        "#", "pedido", "app", "solicitado", "responsavel", "responsável",
        "estado", "decisao", "decisão", "prioridade", "data", "notas",
        "impacto", "esforco", "esforço", "categoria", "budget", "cliente",
        "fase", "dono", "owner", "status"
    )

    lower_values = [v.lower() for v in values]
    hits = 0
    for value in lower_values:
        if value == "#":
            hits += 1
        elif any(keyword in value for keyword in header_keywords if keyword != "#"):
            hits += 1

    return hits >= 2


def is_section_row(ws: Worksheet, row: int, max_col: int) -> bool:
    """Detect section labels such as 'Must-Have -- Fase 1'."""
    values = row_values(ws, row, max_col)
    if len(values) != 1:
        return False
    text = values[0]
    return 3 <= len(text) <= 80 and ("--" in text or "—" in text or text.endswith(":"))


def row_contains_multiline(ws: Worksheet, row: int, max_col: int) -> bool:
    for col in range(1, max_col + 1):
        value = ws.cell(row=row, column=col).value
        if isinstance(value, str) and "\n" in value:
            return True
    return False


# -----------------------------------------------------------------------------
# STYLES
# -----------------------------------------------------------------------------
def apply_title_style(ws: Worksheet, row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = solid(NAVY)
        cell.font = Font(name=FONT_NAME, bold=True, color=WHITE, size=14)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[row].height = 28


def apply_header_style(ws: Worksheet, row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = solid(BLUE)
        cell.font = Font(name=FONT_NAME, bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[row].height = 24


def apply_section_style(ws: Worksheet, row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = solid(LIGHT_BLUE)
        cell.font = Font(name=FONT_NAME, bold=True, color=BLUE, size=10)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[row].height = 22


def apply_body_style(ws: Worksheet, row: int, max_col: int) -> None:
    fill_color = LIGHT_BLUE if row % 2 == 0 else WHITE
    height = 34 if row_contains_multiline(ws, row, max_col) else 22

    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = solid(fill_color)
        cell.font = Font(name=FONT_NAME, color=TEXT, size=10)
        horizontal = "center" if col == 1 else "left"
        cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)
        cell.border = BORDER_ALL
    ws.row_dimensions[row].height = height


def apply_column_widths(ws: Worksheet, max_row: int, max_col: int) -> None:
    """Adjust visible widths without inserting/removing anything."""
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        max_len = 8

        for row in range(1, min(max_row, 120) + 1):
            value = ws.cell(row=row, column=col).value
            if value is None:
                continue
            # For multiline cells, measure the longest line.
            longest = max(len(part) for part in str(value).splitlines())
            max_len = max(max_len, longest)

        # Narrow ID columns; keep text columns readable but bounded.
        if col == 1:
            width = min(max(max_len + 2, 5), 8)
        else:
            width = min(max(max_len + 3, 12), 38)

        ws.column_dimensions[letter].width = width


# -----------------------------------------------------------------------------
# PUBLIC FUNCTIONS
# -----------------------------------------------------------------------------
def apply_parametro_layout(input_xlsx: str | Path, output_xlsx: str | Path) -> None:
    """
    Apply Parametro colors/layout to an existing workbook.

    Important: this function does NOT add, remove or change cell values/formulas.
    It only changes formatting and sheet display settings.
    """
    wb = load_workbook(input_xlsx)

    for ws in wb.worksheets:
        max_row, max_col = get_used_bounds(ws)

        ws.sheet_view.showGridLines = True
        ws.freeze_panes = "A4" if max_row >= 4 else None

        for row in range(1, max_row + 1):
            if is_title_row(ws, row, max_col):
                apply_title_style(ws, row, max_col)
            elif is_header_row(ws, row, max_col):
                apply_header_style(ws, row, max_col)
            elif is_section_row(ws, row, max_col):
                apply_section_style(ws, row, max_col)
            else:
                apply_body_style(ws, row, max_col)

        apply_column_widths(ws, max_row, max_col)

    wb.save(output_xlsx)


def create_empty_template(output_xlsx: str | Path) -> None:
    """Create a blank Parametro-style workbook template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Modelo"
    ws.sheet_view.showGridLines = True

    # Minimal placeholders are included so the formatting is visible.
    # Delete/replace these normally in Excel if needed.
    ws.merge_cells("A1:H1")
    ws["A1"] = "TITULO DO RELATORIO"
    headers = ["#", "Descricao", "Categoria", "Responsavel", "Estado", "Prioridade", "Data", "Notas"]
    for idx, header in enumerate(headers, start=1):
        ws.cell(row=3, column=idx).value = header

    for row in range(4, 14):
        ws.cell(row=row, column=1).value = row - 3

    apply_title_style(ws, 1, 8)
    apply_header_style(ws, 3, 8)
    for row in range(4, 14):
        apply_body_style(ws, row, 8)

    widths = [6, 32, 18, 20, 18, 18, 14, 36]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A4"

    cfg = wb.create_sheet("Config_Cores")
    cfg.append(["Elemento", "Hex", "Uso"])
    rows = [
        ["Azul escuro", NAVY, "Barra superior / titulo"],
        ["Azul Parametro", BLUE, "Cabecalho de tabelas"],
        ["Azul claro", LIGHT_BLUE, "Linhas alternadas"],
        ["Azul claro secundario", LIGHT_BLUE_2, "Destaques secundarios"],
        ["Branco", WHITE, "Fundo"],
        ["Grelha", GRID, "Bordas"],
        ["Texto", TEXT, "Texto principal"],
    ]
    for item in rows:
        cfg.append(item)

    apply_header_style(cfg, 1, 3)
    for row in range(2, 2 + len(rows)):
        apply_body_style(cfg, row, 3)
        cfg.cell(row=row, column=2).fill = solid(str(cfg.cell(row=row, column=2).value))
    cfg.column_dimensions["A"].width = 26
    cfg.column_dimensions["B"].width = 16
    cfg.column_dimensions["C"].width = 38

    wb.save(output_xlsx)


# -----------------------------------------------------------------------------
# CLI
# -----------------------------------------------------------------------------
def main() -> None:
    parser = argparse.ArgumentParser(description="Aplica o layout Excel da Parametro Global.")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--apply", nargs=2, metavar=("INPUT_XLSX", "OUTPUT_XLSX"), help="aplicar layout a um Excel existente")
    group.add_argument("--new", metavar="OUTPUT_XLSX", help="criar template vazio com layout Parametro")
    args = parser.parse_args()

    if args.apply:
        input_xlsx, output_xlsx = args.apply
        apply_parametro_layout(input_xlsx, output_xlsx)
    elif args.new:
        create_empty_template(args.new)


if __name__ == "__main__":
    main()
