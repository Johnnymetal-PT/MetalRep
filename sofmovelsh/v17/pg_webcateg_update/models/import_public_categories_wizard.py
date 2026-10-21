"""Wizard de importacao de categorias eCommerce/POS a partir de XLSX.

Pertence ao modulo `pg_webcateg_update`. Recebe um XLSX com `default_code` + as
colunas de categorias de website e/ou POS, faz match pelo codigo do produto e
cria/atribui as categorias correspondentes (com suporte a hierarquia "A / B / C").

Pre-requisitos:
- `openpyxl` no Python do servidor.
- Permissoes para criar `product.public.category` e/ou `pos.category` quando a
  opcao "Create missing categories" estiver ativa.
"""
import base64
import io
import logging
import re

from odoo import fields, models, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except Exception:  # pragma: no cover - so e None se openpyxl nao estiver instalado
    openpyxl = None


class ImportPublicCategoriesWizard(models.TransientModel):
    """Wizard que le um XLSX e regrava categorias por produto via default_code."""

    _name = "import.public.categories.wizard"
    _description = "Import Website/POS Categories from XLSX (match by default_code)"

    file_data = fields.Binary(string="XLSX File", required=True)
    file_name = fields.Char(string="Filename")

    # What to import
    do_website_categories = fields.Boolean(string="Update Website Categories", default=True)
    do_pos_categories = fields.Boolean(string="Update POS Categories", default=False)

    website_id = fields.Many2one(
        "website",
        string="Website (optional)",
        help="If set and the category model supports website_id, categories may be searched/created for this website.",
    )

    only_active_products = fields.Boolean(
        string="Only update active products",
        default=True,
        help="If enabled, products with active=False will be skipped.",
    )

    create_missing_categories = fields.Boolean(
        string="Create missing categories",
        default=True,
        help="If disabled, only existing categories will be linked.",
    )

    replace_categories = fields.Boolean(
        string="Replace categories",
        default=True,
        help="If enabled, categories are replaced by XLSX values. If disabled, categories are added.",
    )

    # Column names expected in XLSX header
    col_default_code = fields.Char(string="Column for default_code", default="default_code")

    # Separate columns for each category type
    col_website_categories = fields.Char(
        string="Website categories column",
        default="website_categories",
        help="Column containing website category paths (comma-separated). "
             "If your sheet uses 'public_categ_ids', set this to public_categ_ids.",
    )
    col_pos_categories = fields.Char(
        string="POS categories column",
        default="pos_categ_ids",
        help="Column containing POS categories (comma-separated). If you exported leaf-only, this is perfect.",
    )

    # -------------------------
    # Helpers
    # -------------------------
    def _norm(self, s):
        """Normalize XLSX cell values to a clean string (handles int/float cells from Excel)."""
        if s is None:
            return ""
        # Excel may store default_code as a number
        if isinstance(s, (int, float)):
            # Avoid "123.0"
            if isinstance(s, float) and s.is_integer():
                s = int(s)
            return str(s).strip()
        return str(s).strip()

    def _collapse_spaces(self, s: str) -> str:
        s = (s or "").strip()
        s = re.sub(r"\s+", " ", s)
        return s

    def _normalize_path(self, path: str) -> str:
        """
        Normalize a single category path:
        - trim spaces
        - remove trailing/leading slashes
        - collapse multiple spaces
        - normalize separators around '/'
        """
        path = self._collapse_spaces(path)
        path = path.strip().strip("/")
        if not path:
            return ""
        # normalize " / " spacing
        parts = [self._collapse_spaces(p) for p in path.split("/") if self._collapse_spaces(p)]
        return " / ".join(parts)

    def _split_category_paths(self, s):
        """
        Input example:
          "Produtos Kave,Todas / Móveis / Móvel TV"
        Output:
          ["Produtos Kave", "Todas / Móveis / Móvel TV"]
        """
        s = self._norm(s)
        if not s:
            return []
        raw_parts = [p.strip() for p in str(s).split(",")]
        cleaned = []
        for p in raw_parts:
            p = self._normalize_path(p)
            if p:
                cleaned.append(p)
        return cleaned

    def _get_sheet_and_header_map(self, wb):
        """
        Use the first sheet. Detect header row (within first 50 rows).
        Requires:
          - default_code always
          - website column only if do_website_categories
          - pos column only if do_pos_categories
        """
        ws = wb.worksheets[0]

        target_dc = (self.col_default_code or "default_code").strip().lower()
        target_web = (self.col_website_categories or "website_categories").strip().lower()
        target_pos = (self.col_pos_categories or "pos_categ_ids").strip().lower()

        needed = {target_dc}
        if self.do_website_categories:
            needed.add(target_web)
        if self.do_pos_categories:
            needed.add(target_pos)

        for r in range(1, min(ws.max_row, 50) + 1):
            row_vals = []
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                key = str(v).strip().lower() if v is not None else ""
                row_vals.append((key, c))

            keys = {k for k, _c in row_vals if k}
            if needed.issubset(keys):
                colmap = {"default_code": None, "website": None, "pos": None}
                for k, c in row_vals:
                    if k == target_dc:
                        colmap["default_code"] = c
                    elif self.do_website_categories and k == target_web:
                        colmap["website"] = c
                    elif self.do_pos_categories and k == target_pos:
                        colmap["pos"] = c
                return ws, r, colmap

        raise UserError(
            _("Could not find header row containing required columns: %s") % ", ".join(sorted(needed))
        )

    def _category_search_domain_for_leaf(self, Model, name, parent=None, website=False):
        domain = [
            ("name", "=", name),
            ("parent_id", "=", parent.id if parent else False),
        ]
        if website and self.website_id and "website_id" in Model._fields:
            domain.append(("website_id", "in", [False, self.website_id.id]))
        return domain

    def _get_or_create_category_by_path(self, model_name, path, website=False):
        """
        Find or create hierarchical categories for:
          - website: product.public.category
          - pos: pos.category
        Strategy:
          1) Try exact match on complete_name if exists
          2) Else walk/create by segments (name + parent_id)
        """
        Model = self.env[model_name].sudo()

        path = self._normalize_path(path)
        if not path:
            return False

        # 1) exact match on complete_name if available
        if "complete_name" in Model._fields:
            cat = Model.search([("complete_name", "=", path)], limit=1)
            if cat:
                return cat

        # 2) walk segments
        segments = [seg.strip() for seg in path.split(" / ") if seg.strip()]
        parent = False
        current = False

        for name in segments:
            domain = self._category_search_domain_for_leaf(Model, name, parent=parent, website=website)
            current = Model.search(domain, limit=1)
            if not current:
                if not self.create_missing_categories:
                    return False
                vals = {
                    "name": name,
                    "parent_id": parent.id if parent else False,
                }
                if website and self.website_id and "website_id" in Model._fields:
                    vals["website_id"] = self.website_id.id
                current = Model.create(vals)
            parent = current

        return current

    # -------------------------
    # Main action
    # -------------------------
    def action_import(self):
        if not (self.do_website_categories or self.do_pos_categories):
            raise UserError(_("Select at least one option: Website categories and/or POS categories."))

        if openpyxl is None:
            raise UserError(_("openpyxl is not installed on the server."))

        if not self.file_data:
            raise UserError(_("Please upload an XLSX file."))

        data = base64.b64decode(self.file_data)
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)

        ws, header_row, colmap = self._get_sheet_and_header_map(wb)

        Product = self.env["product.template"].sudo()

        updated_web = 0
        updated_pos = 0
        no_product = 0
        no_web = 0
        no_pos = 0

        created_before_public = set(self.env["product.public.category"].sudo().search([]).ids) if self.do_website_categories else set()
        created_before_pos = set(self.env["pos.category"].sudo().search([]).ids) if self.do_pos_categories else set()

        for r in range(header_row + 1, ws.max_row + 1):
            dc_val = ws.cell(r, colmap["default_code"]).value
            default_code = self._norm(dc_val)
            if not default_code:
                continue

            domain = [("default_code", "=", default_code)]
            if self.only_active_products and "active" in Product._fields:
                domain.append(("active", "=", True))

            product = Product.search(domain, limit=1)
            if not product:
                no_product += 1
                continue

            # --- Website categories ---
            if self.do_website_categories:
                web_val = ws.cell(r, colmap["website"]).value if colmap["website"] else None
                paths = self._split_category_paths(web_val)
                if not paths:
                    no_web += 1
                else:
                    cat_ids = []
                    for p in paths:
                        cat = self._get_or_create_category_by_path("product.public.category", p, website=True)
                        if cat:
                            cat_ids.append(cat.id)
                    cat_ids = sorted(set(cat_ids))
                    if cat_ids:
                        if self.replace_categories:
                            product.write({"public_categ_ids": [(6, 0, cat_ids)]})
                        else:
                            product.write({"public_categ_ids": [(4, cid) for cid in cat_ids]})
                        updated_web += 1
                    else:
                        no_web += 1

            # --- POS categories ---
            if self.do_pos_categories:
                pos_val = ws.cell(r, colmap["pos"]).value if colmap["pos"] else None
                paths = self._split_category_paths(pos_val)
                if not paths:
                    no_pos += 1
                else:
                    cat_ids = []
                    for p in paths:
                        cat = self._get_or_create_category_by_path("pos.category", p, website=False)
                        if cat:
                            cat_ids.append(cat.id)
                    cat_ids = sorted(set(cat_ids))
                    if cat_ids:
                        if self.replace_categories:
                            product.write({"pos_categ_ids": [(6, 0, cat_ids)]})
                        else:
                            product.write({"pos_categ_ids": [(4, cid) for cid in cat_ids]})
                        updated_pos += 1
                    else:
                        no_pos += 1

        created_public = 0
        created_pos = 0

        if self.do_website_categories:
            created_after_public = set(self.env["product.public.category"].sudo().search([]).ids)
            created_public = len(created_after_public - created_before_public)

        if self.do_pos_categories:
            created_after_pos = set(self.env["pos.category"].sudo().search([]).ids)
            created_pos = len(created_after_pos - created_before_pos)

        msg_lines = [
            _("Import finished."),
            _("- Products not found (default_code): %d") % no_product,
        ]

        if self.do_website_categories:
            msg_lines += [
                _("- Website categories updated: %d") % updated_web,
                _("- Rows with no website categories applied: %d") % no_web,
                _("- Website categories created: %d") % created_public,
            ]

        if self.do_pos_categories:
            msg_lines += [
                _("- POS categories updated: %d") % updated_pos,
                _("- Rows with no POS categories applied: %d") % no_pos,
                _("- POS categories created: %d") % created_pos,
            ]

        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("Category import (XLSX)"),
                "message": "\n".join(msg_lines),
                "type": "success",
                "sticky": True,
            },
        }

