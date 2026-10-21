import base64
import io
import logging
from datetime import datetime
import xlrd
import openpyxl
import csv
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class BIMStatementParser:
    """Parser for BIM (Millennium BIM) bank statements"""
    
    def __init__(self):
        self.header_lines_to_skip = 11  # Skip first 11 lines (header info)
        
    def parse_file(self, file_content, file_name):
        """
        Parse BIM bank statement file and return formatted data for Odoo import
        
        :param file_content: Base64 encoded file content
        :param file_name: Name of the file
        :return: Formatted data ready for Odoo import
        """
        try:
            _logger.info(f"BIM Parser: Starting to parse file {file_name}")
            
            # Decode file content
            file_data = base64.b64decode(file_content)
            _logger.info(f"BIM Parser: File decoded, size: {len(file_data)} bytes")
            
            # Determine file type and parse accordingly
            if file_name.lower().endswith('.xlsx'):
                _logger.info("BIM Parser: Processing as XLSX file")
                return self._parse_xlsx(file_data)
            elif file_name.lower().endswith('.xls'):
                _logger.info("BIM Parser: Processing as XLS file")
                return self._parse_xls(file_data)
            elif file_name.lower().endswith('.csv'):
                _logger.info("BIM Parser: Processing as CSV file")
                return self._parse_csv(file_data)
            else:
                raise UserError("Formato de arquivo não suportado. Use XLSX, XLS ou CSV.")
                
        except Exception as e:
            _logger.error(f"Erro ao processar arquivo BIM: {str(e)}")
            raise UserError(f"Erro ao processar arquivo: {str(e)}")
    
    def _parse_xlsx(self, file_data):
        """Parse XLSX file format"""
        try:
            # Create a BytesIO object from file data
            file_stream = io.BytesIO(file_data)
            
            # Load workbook
            wb = openpyxl.load_workbook(file_stream, data_only=True)
            sheet = wb.active
            
            _logger.info(f"Processing BIM file: Sheet '{sheet.title}' with {sheet.max_row} rows")
            
            transactions = []
            
            # Process transactions starting from row 12
            row_count = 0
            for row_idx, row in enumerate(sheet.iter_rows(min_row=12, values_only=True), start=12):
                row_count += 1
                # Skip completely empty rows
                if not row or all(v is None for v in row):
                    continue
                    
                transaction = self._process_row(row, row_idx)
                if transaction:
                    transactions.append(transaction)
                    _logger.debug(f"Transaction added from row {row_idx}")
            
            _logger.info(f"BIM Parser: Processed {row_count} rows, found {len(transactions)} transactions")
            return self._format_for_odoo_import(transactions)
            
        except Exception as e:
            _logger.error(f"Erro ao processar XLSX: {str(e)}")
            raise UserError(f"Erro ao processar arquivo XLSX: {str(e)}")
    
    def _parse_xls(self, file_data):
        """Parse XLS file format"""
        try:
            # Create a BytesIO object from file data
            file_stream = io.BytesIO(file_data)
            
            # Open workbook
            wb = xlrd.open_workbook(file_contents=file_stream.read())
            sheet = wb.sheet_by_index(0)
            
            transactions = []
            
            # Start from row 12 (index 11 in xlrd)
            for row_idx in range(11, sheet.nrows):
                row = sheet.row_values(row_idx)
                
                # Skip empty rows
                if not row or not row[0]:  # Check if first column (date) is empty
                    continue
                    
                transaction = self._process_row(row, row_idx + 1)
                if transaction:
                    transactions.append(transaction)
            
            return self._format_for_odoo_import(transactions)
            
        except Exception as e:
            _logger.error(f"Erro ao processar XLS: {str(e)}")
            raise UserError(f"Erro ao processar arquivo XLS: {str(e)}")
    
    def _parse_csv(self, file_data):
        """Parse CSV file format"""
        try:
            # Decode bytes to string
            text_data = file_data.decode('utf-8-sig')  # Handle BOM if present
            
            # Create StringIO object
            file_stream = io.StringIO(text_data)
            
            # Create CSV reader
            csv_reader = csv.reader(file_stream, delimiter=',')
            
            transactions = []
            
            # Skip header lines
            for _ in range(self.header_lines_to_skip):
                next(csv_reader, None)
            
            # Process transactions
            for row_idx, row in enumerate(csv_reader, start=12):
                # Skip empty rows
                if not row or not row[0]:  # Check if first column (date) is empty
                    continue
                    
                transaction = self._process_row(row, row_idx)
                if transaction:
                    transactions.append(transaction)
            
            return self._format_for_odoo_import(transactions)
            
        except Exception as e:
            _logger.error(f"Erro ao processar CSV: {str(e)}")
            raise UserError(f"Erro ao processar arquivo CSV: {str(e)}")
    
    def _process_row(self, row, row_number):
        """
        Process a single row from BIM statement
        
        Two possible formats:
        Format 1 (Test file):
        0: Data da Transação (Transaction Date)
        1: Data Valor (Value Date) 
        2: Descrição (Description)
        3: Débito (Debit - positive values)
        4: Crédito (Credit - positive values)
        5: Saldo (Balance)
        
        Format 2 (Real file - data starts at column C/index 2):
        2: Data da Transação (Transaction Date)
        3: Data Valor (Value Date)
        4: Descrição (Description)
        5: None
        6: Débito (Debit - already negative values)
        7: None
        8: Crédito (Credit - positive values or 0)
        9: Saldo (Balance)
        """
        try:
            # Auto-detect format based on where data starts
            # Check if first columns are None/empty (real file format)
            if len(row) > 2 and row[0] is None and row[1] is None and row[2] is not None:
                # Format 2: Real file (data starts at index 2/column C)
                transaction_date = row[2] if len(row) > 2 else None
                value_date = row[3] if len(row) > 3 else None
                description = row[4] if len(row) > 4 else ""
                # Skip column F (index 5) - it's None
                debit = row[6] if len(row) > 6 else 0
                # Skip column H (index 7) - it's None
                credit = row[8] if len(row) > 8 else 0
                balance = row[9] if len(row) > 9 else 0
            else:
                # Format 1: Test file (data starts at index 0/column A)
                transaction_date = row[0] if len(row) > 0 else None
                value_date = row[1] if len(row) > 1 else None
                description = row[2] if len(row) > 2 else ""
                debit = row[3] if len(row) > 3 else 0
                credit = row[4] if len(row) > 4 else 0
                balance = row[5] if len(row) > 5 else 0
            
            # Skip if no date
            if not transaction_date:
                return None
            
            # Skip header row if it contains text like "Data" 
            if isinstance(transaction_date, str) and 'data' in transaction_date.lower():
                return None
            
            # Parse date (handle different formats)
            date_obj = self._parse_date(transaction_date)
            if not date_obj:
                _logger.debug(f"Could not parse date in row {row_number}: {transaction_date}")
                return None
            
            # Calculate amount
            # Note: In real files, debits come as negative values already
            # In test files, debits come as positive values
            amount = 0.0
            
            # Process debit
            if debit and debit != 0:
                debit_parsed = self._parse_amount(debit)
                # If debit is already negative (real file), use as is
                # If debit is positive (test file), make it negative
                if debit_parsed < 0:
                    amount = debit_parsed  # Already negative
                else:
                    amount = -debit_parsed  # Make it negative
            
            # Process credit (always positive when present)
            if credit and credit != 0:
                credit_parsed = self._parse_amount(credit)
                if credit_parsed != 0:
                    amount = abs(credit_parsed)  # Ensure positive
            
            # Skip if amount is zero
            if amount == 0:
                return None
            
            transaction = {
                'date': date_obj,
                'payment_ref': str(description).strip() if description else "",
                'amount': amount,
                'balance': self._parse_amount(balance) if balance else 0,
                'row_number': row_number
            }
            
            return transaction
            
        except Exception as e:
            _logger.debug(f"Error processing row {row_number}: {str(e)}")
            return None
    
    def _parse_date(self, date_value):
        """Parse date from various formats"""
        if not date_value:
            return None
            
        # If it's already a datetime object
        if isinstance(date_value, datetime):
            return date_value.date()
        
        # If it's a date object from datetime
        from datetime import date
        if isinstance(date_value, date):
            return date_value
        
        # If it's a string, try different formats
        if isinstance(date_value, str):
            date_formats = [
                '%d/%m/%Y',
                '%d-%m-%Y', 
                '%Y-%m-%d',
                '%d/%m/%y',
                '%d-%m-%y',
                '%m/%d/%Y',  # US format
                '%m-%d-%Y',  # US format
                '%Y/%m/%d',  # ISO format variant
            ]
            
            for fmt in date_formats:
                try:
                    return datetime.strptime(date_value.strip(), fmt).date()
                except ValueError:
                    continue
        
        # If it's a number (Excel date)
        if isinstance(date_value, (int, float)):
            try:
                # Excel date starts from 1900-01-01
                from datetime import timedelta
                excel_base = datetime(1900, 1, 1)
                # Excel incorrectly considers 1900 as leap year, so subtract 2
                delta = timedelta(days=int(date_value) - 2)
                return (excel_base + delta).date()
            except:
                pass
        
        return None
    
    def _parse_amount(self, amount_value):
        """Parse amount from various formats"""
        if not amount_value:
            return 0.0
            
        # If it's already a number
        if isinstance(amount_value, (int, float)):
            return float(amount_value)
        
        # If it's a string, clean and parse
        if isinstance(amount_value, str):
            # Remove currency symbols, spaces
            cleaned = amount_value.strip()
            cleaned = cleaned.replace(' ', '')
            cleaned = cleaned.replace('MZN', '')
            cleaned = cleaned.replace('MT', '')
            cleaned = cleaned.replace('Mts', '')
            
            # Handle different decimal/thousand separators
            # Check if we have both . and ,
            if '.' in cleaned and ',' in cleaned:
                # Determine which is thousand separator
                dot_pos = cleaned.rfind('.')
                comma_pos = cleaned.rfind(',')
                
                if dot_pos > comma_pos:
                    # . is decimal separator, , is thousand separator
                    cleaned = cleaned.replace(',', '')
                else:
                    # , is decimal separator, . is thousand separator
                    cleaned = cleaned.replace('.', '')
                    cleaned = cleaned.replace(',', '.')
            elif ',' in cleaned:
                # Only comma present - could be decimal or thousand separator
                # If there are 2 digits after comma, likely decimal
                parts = cleaned.split(',')
                if len(parts) == 2 and len(parts[1]) == 2:
                    cleaned = cleaned.replace(',', '.')
                else:
                    # Likely thousand separator
                    cleaned = cleaned.replace(',', '')
            
            # Handle parentheses for negative numbers
            if cleaned.startswith('(') and cleaned.endswith(')'):
                cleaned = '-' + cleaned[1:-1]
            
            try:
                return float(cleaned)
            except ValueError:
                return 0.0
        
        return 0.0
    
    def _format_for_odoo_import(self, transactions):
        """
        Format transactions for Odoo import
        Returns data in format expected by base_import.import
        """
        if not transactions:
            raise UserError("Nenhuma transação válida encontrada no arquivo.")
        
        # Sort transactions by date
        transactions.sort(key=lambda x: x['date'])
        
        # Create CSV data for Odoo import
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Write header (field names for account.bank.statement.line)
        writer.writerow(['date', 'payment_ref', 'amount'])
        
        # Write transactions
        for trans in transactions:
            writer.writerow([
                trans['date'].strftime('%Y-%m-%d'),
                trans['payment_ref'],
                str(trans['amount'])
            ])
        
        # Get CSV content
        csv_content = output.getvalue()
        
        # Encode to base64
        csv_base64 = base64.b64encode(csv_content.encode('utf-8'))
        
        return {
            'file': csv_base64,
            'file_name': 'bim_formatted_statement.csv',
            'file_type': 'text/csv',
            'transaction_count': len(transactions),
            'total_debit': sum(t['amount'] for t in transactions if t['amount'] < 0),
            'total_credit': sum(t['amount'] for t in transactions if t['amount'] > 0),
        }