import base64
import io
import logging
from datetime import datetime

# Test if libraries are available
try:
    import xlrd
    import openpyxl
    HAS_EXCEL_LIBS = True
except ImportError as e:
    HAS_EXCEL_LIBS = False
    import_error = str(e)

import csv
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

# Log library availability at module load
if HAS_EXCEL_LIBS:
    _logger.info("BIM Parser: Excel libraries (openpyxl, xlrd) loaded successfully")
else:
    _logger.error(f"BIM Parser: Failed to import Excel libraries: {import_error}")


class BIMStatementParser:
    """Parser for BIM (Millennium BIM) bank statements"""
    
    def __init__(self):
        self.header_lines_to_skip = 11
        if not HAS_EXCEL_LIBS:
            raise UserError(f"Excel libraries not available: {import_error}")
        
    def parse_file(self, file_content, file_name):
        """
        Parse BIM bank statement file and return formatted data for Odoo import
        """
        try:
            _logger.info(f"BIM Parser: Starting to parse file {file_name}")
            
            # Decode file content
            file_data = base64.b64decode(file_content)
            _logger.info(f"BIM Parser: File decoded, size: {len(file_data)} bytes")
            
            # Determine file type and parse accordingly
            if file_name.lower().endswith('.xlsx'):
                _logger.info("BIM Parser: Processing as XLSX file")
                result = self._parse_xlsx(file_data)
                _logger.info(f"BIM Parser: XLSX parsing complete, got result: {result is not None}")
                return result
            elif file_name.lower().endswith('.xls'):
                _logger.info("BIM Parser: Processing as XLS file")
                return self._parse_xls(file_data)
            elif file_name.lower().endswith('.csv'):
                _logger.info("BIM Parser: Processing as CSV file")
                return self._parse_csv(file_data)
            else:
                raise UserError("Formato de arquivo não suportado. Use XLSX, XLS ou CSV.")
                
        except Exception as e:
            _logger.error(f"Erro ao processar arquivo BIM: {str(e)}")
            import traceback
            _logger.error(traceback.format_exc())
            raise UserError(f"Erro ao processar arquivo: {str(e)}")
    
    def _parse_xlsx(self, file_data):
        """Parse XLSX file format"""
        try:
            _logger.info("BIM Parser: _parse_xlsx called")
            # Create a BytesIO object from file data
            file_stream = io.BytesIO(file_data)
            
            # Load workbook
            wb = openpyxl.load_workbook(file_stream, data_only=True)
            sheet = wb.active
            
            _logger.info(f"BIM Parser: Sheet '{sheet.title}' with {sheet.max_row} rows")
            
            transactions = []
            
            # Process transactions starting from row 12
            row_count = 0
            for row_idx, row in enumerate(sheet.iter_rows(min_row=12, values_only=True), start=12):
                row_count += 1
                # Skip completely empty rows
                if not row or all(v is None for v in row):
                    continue
                    
                transaction = self._process_row(row, row_idx)
                if transaction:
                    transactions.append(transaction)
                    _logger.info(f"BIM Parser: Transaction added from row {row_idx}: {transaction}")
            
            _logger.info(f"BIM Parser: Processed {row_count} rows, found {len(transactions)} transactions")
            
            if not transactions:
                _logger.warning("BIM Parser: No transactions found!")
                # Log some sample rows for debugging
                for i, row in enumerate(sheet.iter_rows(min_row=12, max_row=15, values_only=True), start=12):
                    _logger.warning(f"  Sample row {i}: {row[:10] if row else 'Empty'}")
            
            result = self._format_for_odoo_import(transactions)
            _logger.info(f"BIM Parser: Formatted result with {result.get('transaction_count')} transactions")
            return result
            
        except Exception as e:
            _logger.error(f"Erro ao processar XLSX: {str(e)}")
            import traceback
            _logger.error(traceback.format_exc())
            raise UserError(f"Erro ao processar arquivo XLSX: {str(e)}")
    
    def _process_row(self, row, row_number):
        """Process a single row from BIM statement"""
        try:
            _logger.debug(f"Processing row {row_number}")
            
            # Auto-detect format based on where data starts
            # Check if first columns are None/empty (real file format)
            if len(row) > 2 and row[0] is None and row[1] is None and row[2] is not None:
                # Format 2: Real file (data starts at index 2/column C)
                transaction_date = row[2] if len(row) > 2 else None
                value_date = row[3] if len(row) > 3 else None
                description = row[4] if len(row) > 4 else ""
                # Skip column F (index 5) - it's None
                debit = row[6] if len(row) > 6 else 0
                # Skip column H (index 7) - it's None
                credit = row[8] if len(row) > 8 else 0
                balance = row[9] if len(row) > 9 else 0
                _logger.debug(f"Row {row_number}: Real format - date={transaction_date}, debit={debit}, credit={credit}")
            else:
                # Format 1: Test file (data starts at index 0/column A)
                transaction_date = row[0] if len(row) > 0 else None
                value_date = row[1] if len(row) > 1 else None
                description = row[2] if len(row) > 2 else ""
                debit = row[3] if len(row) > 3 else 0
                credit = row[4] if len(row) > 4 else 0
                balance = row[5] if len(row) > 5 else 0
                _logger.debug(f"Row {row_number}: Test format - date={transaction_date}, debit={debit}, credit={credit}")
            
            # Skip if no date
            if not transaction_date:
                return None
            
            # Skip header row if it contains text like "Data" 
            if isinstance(transaction_date, str) and 'data' in transaction_date.lower():
                return None
            
            # Parse date (handle different formats)
            date_obj = self._parse_date(transaction_date)
            if not date_obj:
                _logger.debug(f"Could not parse date in row {row_number}: {transaction_date}")
                return None
            
            # Calculate amount
            # Note: In real files, debits come as negative values already
            # In test files, debits come as positive values
            amount = 0.0
            
            # Process debit
            if debit and debit != 0:
                debit_parsed = self._parse_amount(debit)
                # If debit is already negative (real file), use as is
                # If debit is positive (test file), make it negative
                if debit_parsed < 0:
                    amount = debit_parsed  # Already negative
                else:
                    amount = -debit_parsed  # Make it negative
            
            # Process credit (always positive when present)
            if credit and credit != 0:
                credit_parsed = self._parse_amount(credit)
                if credit_parsed != 0:
                    amount = abs(credit_parsed)  # Ensure positive
            
            # Skip if amount is zero
            if amount == 0:
                return None
            
            transaction = {
                'date': date_obj,
                'payment_ref': str(description).strip() if description else "",
                'amount': amount,
                'balance': self._parse_amount(balance) if balance else 0,
                'row_number': row_number
            }
            
            return transaction
            
        except Exception as e:
            _logger.debug(f"Error processing row {row_number}: {str(e)}")
            return None
    
    def _parse_date(self, date_value):
        """Parse date from various formats"""
        if not date_value:
            return None
            
        # If it's already a datetime object
        if isinstance(date_value, datetime):
            return date_value.date()
        
        # If it's a date object from datetime
        from datetime import date
        if isinstance(date_value, date):
            return date_value
        
        # If it's a string, try different formats
        if isinstance(date_value, str):
            date_formats = [
                '%d/%m/%Y',
                '%d-%m-%Y', 
                '%Y-%m-%d',
                '%d/%m/%y',
                '%d-%m-%y',
                '%m/%d/%Y',  # US format
                '%m-%d-%Y',  # US format
                '%Y/%m/%d',  # ISO format variant
            ]
            
            for fmt in date_formats:
                try:
                    return datetime.strptime(date_value.strip(), fmt).date()
                except ValueError:
                    continue
        
        # If it's a number (Excel date)
        if isinstance(date_value, (int, float)):
            try:
                # Excel date starts from 1900-01-01
                from datetime import timedelta
                excel_base = datetime(1900, 1, 1)
                # Excel incorrectly considers 1900 as leap year, so subtract 2
                delta = timedelta(days=int(date_value) - 2)
                return (excel_base + delta).date()
            except:
                pass
        
        return None
    
    def _parse_amount(self, amount_value):
        """Parse amount from various formats"""
        if not amount_value:
            return 0.0
            
        # If it's already a number
        if isinstance(amount_value, (int, float)):
            return float(amount_value)
        
        # If it's a string, clean and parse
        if isinstance(amount_value, str):
            # Remove currency symbols, spaces
            cleaned = amount_value.strip()
            cleaned = cleaned.replace(' ', '')
            cleaned = cleaned.replace('MZN', '')
            cleaned = cleaned.replace('MT', '')
            cleaned = cleaned.replace('Mts', '')
            
            # Handle different decimal/thousand separators
            # Check if we have both . and ,
            if '.' in cleaned and ',' in cleaned:
                # Determine which is thousand separator
                dot_pos = cleaned.rfind('.')
                comma_pos = cleaned.rfind(',')
                
                if dot_pos > comma_pos:
                    # . is decimal separator, , is thousand separator
                    cleaned = cleaned.replace(',', '')
                else:
                    # , is decimal separator, . is thousand separator
                    cleaned = cleaned.replace('.', '')
                    cleaned = cleaned.replace(',', '.')
            elif ',' in cleaned:
                # Only comma present - could be decimal or thousand separator
                # If there are 2 digits after comma, likely decimal
                parts = cleaned.split(',')
                if len(parts) == 2 and len(parts[1]) == 2:
                    cleaned = cleaned.replace(',', '.')
                else:
                    # Likely thousand separator
                    cleaned = cleaned.replace(',', '')
            
            # Handle parentheses for negative numbers
            if cleaned.startswith('(') and cleaned.endswith(')'):
                cleaned = '-' + cleaned[1:-1]
            
            try:
                return float(cleaned)
            except ValueError:
                return 0.0
        
        return 0.0
    
    def _format_for_odoo_import(self, transactions):
        """
        Format transactions for Odoo import
        Returns data in format expected by base_import.import
        """
        if not transactions:
            raise UserError("Nenhuma transação válida encontrada no arquivo.")
        
        # Sort transactions by date
        transactions.sort(key=lambda x: x['date'])
        
        # Create CSV data for Odoo import
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Write header (field names for account.bank.statement.line)
        writer.writerow(['date', 'payment_ref', 'amount'])
        
        # Write transactions
        for trans in transactions:
            writer.writerow([
                trans['date'].strftime('%Y-%m-%d'),
                trans['payment_ref'],
                str(trans['amount'])
            ])
        
        # Get CSV content
        csv_content = output.getvalue()
        
        # Encode to base64
        csv_base64 = base64.b64encode(csv_content.encode('utf-8'))
        
        return {
            'file': csv_base64,
            'file_name': 'bim_formatted_statement.csv',
            'file_type': 'text/csv',
            'transaction_count': len(transactions),
            'total_debit': sum(t['amount'] for t in transactions if t['amount'] < 0),
            'total_credit': sum(t['amount'] for t in transactions if t['amount'] > 0),
        }
    
    # Stub methods for XLS and CSV (copy from original)
    def _parse_xls(self, file_data):
        """Parse XLS file format"""
        raise UserError("XLS format not yet implemented in debug version")
    
    def _parse_csv(self, file_data):
        """Parse CSV file format"""
        raise UserError("CSV format not yet implemented in debug version")