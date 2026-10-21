#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para extrair informações específicas do arquivo BIM
"""

import openpyxl
from datetime import datetime

def extract_bim_info(file_path):
    """Extrai informações importantes do arquivo BIM"""
    
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    
    print("="*80)
    print("EXTRAÇÃO DE INFORMAÇÕES DO EXTRATO BIM")
    print("="*80)
    
    # 1. Extrai período do extrato (linha 6)
    periodo_cell = sheet.cell(row=6, column=3).value
    if periodo_cell and 'de' in periodo_cell and 'até' in periodo_cell:
        print(f"\n[PERIODO DO EXTRATO]")
        print(f"   {periodo_cell}")
        # Extrai as datas
        import re
        dates = re.findall(r'\d{2}-\d{2}-\d{4}', periodo_cell)
        if len(dates) == 2:
            print(f"   Data Inicial: {dates[0]}")
            print(f"   Data Final: {dates[1]}")
    
    # 2. Extrai saldo inicial (linha 7, coluna 9)
    print(f"\n[SALDO INICIAL]")
    saldo_inicial = None
    # Verifica várias colunas possíveis para o saldo
    for col in [9, 8, 7, 6, 5]:
        val = sheet.cell(row=7, column=col).value
        if val:
            print(f"   Coluna {col}: {val}")
            saldo_inicial = val
            break
    
    # 3. Extrai saldo final (linha 10, coluna 9)
    print(f"\n[SALDO FINAL]")
    saldo_final = None
    for col in [9, 8, 7, 6, 5]:
        val = sheet.cell(row=10, column=col).value
        if val:
            print(f"   Coluna {col}: {val}")
            saldo_final = val
            break
    
    # 4. Extrai primeira transação (linha 12)
    print(f"\n[PRIMEIRA TRANSACAO]")
    primeira_trans = []
    for col in range(1, 12):
        val = sheet.cell(row=12, column=col).value
        if val:
            primeira_trans.append((col, val))
    
    for col, val in primeira_trans:
        if isinstance(val, datetime):
            val = val.strftime('%d/%m/%Y')
        print(f"   Coluna {col}: {val}")
    
    # 5. Encontra última transação (última linha com dados)
    print(f"\n[ULTIMA TRANSACAO]")
    ultima_linha = None
    for row in range(sheet.max_row, 11, -1):  # Começa do fim e vai até linha 12
        date_cell = sheet.cell(row=row, column=3).value  # Coluna C tem a data
        if date_cell and (isinstance(date_cell, datetime) or (isinstance(date_cell, str) and date_cell.strip())):
            ultima_linha = row
            break
    
    if ultima_linha:
        print(f"   Encontrada na linha {ultima_linha}:")
        ultima_trans = []
        for col in range(1, 12):
            val = sheet.cell(row=ultima_linha, column=col).value
            if val:
                if isinstance(val, datetime):
                    val = val.strftime('%d/%m/%Y')
                ultima_trans.append((col, val))
        
        for col, val in ultima_trans:
            print(f"   Coluna {col}: {val}")
    
    # 6. Extrai totais (linhas 8 e 9)
    print(f"\n[TOTAIS]")
    creditos = None
    debitos = None
    for col in [9, 8, 7, 6, 5]:
        val_cred = sheet.cell(row=8, column=col).value
        val_deb = sheet.cell(row=9, column=col).value
        if val_cred:
            creditos = val_cred
            print(f"   Total Creditos (linha 8, coluna {col}): {val_cred}")
        if val_deb:
            debitos = val_deb
            print(f"   Total Debitos (linha 9, coluna {col}): {val_deb}")
        if creditos and debitos:
            break
    
    # 7. Conta total de transações
    print(f"\n[ESTATISTICAS]")
    total_trans = 0
    for row in range(12, sheet.max_row + 1):
        date_cell = sheet.cell(row=row, column=3).value
        if date_cell and (isinstance(date_cell, datetime) or (isinstance(date_cell, str) and date_cell.strip())):
            total_trans += 1
    print(f"   Total de transacoes: {total_trans}")
    print(f"   Primeira transacao na linha: 12")
    print(f"   Ultima transacao na linha: {ultima_linha}")
    
    # 8. Mostra estrutura das colunas de transação
    print(f"\n[ESTRUTURA DAS COLUNAS] (linha 11 - cabecalho):")
    for col in range(1, 12):
        header = sheet.cell(row=11, column=col).value
        if header:
            print(f"   Coluna {col}: {header}")
    
    wb.close()
    
    # Retorna dicionário com informações
    return {
        'saldo_inicial': saldo_inicial,
        'saldo_final': saldo_final,
        'total_creditos': creditos,
        'total_debitos': debitos,
        'total_transacoes': total_trans,
        'primeira_linha': 12,
        'ultima_linha': ultima_linha
    }

if __name__ == "__main__":
    file_path = r"C:\Users\custodio.massingue\Downloads\docs\co.xlsx"
    info = extract_bim_info(file_path)
    
    print("\n" + "="*80)
    print("RESUMO PARA IMPLEMENTAÇÃO")
    print("="*80)
    print(f"""
    Para o preview do BIM, precisamos mostrar:
    
    1. SALDO INICIAL: Linha 7, Coluna 9 (ou procurar nas colunas 5-9)
    2. SALDO FINAL: Linha 10, Coluna 9 (ou procurar nas colunas 5-9)
    3. PRIMEIRA TRANSAÇÃO: Linha 12 (Data na coluna 3, Descrição na coluna 5, Valor na coluna 7 ou 8)
    4. ÚLTIMA TRANSAÇÃO: Varredura reversa a partir do fim do arquivo
    5. PERÍODO: Extrair da linha 6, coluna 3
    6. TOTAIS: Créditos na linha 8, Débitos na linha 9
    
    Estrutura das transações (começam na linha 12):
    - Coluna 3: Data da Transação
    - Coluna 4: Data Valor
    - Coluna 5: Descrição
    - Coluna 7: Débito
    - Coluna 8: Crédito
    - Coluna 9: Saldo (quando presente)
    """)