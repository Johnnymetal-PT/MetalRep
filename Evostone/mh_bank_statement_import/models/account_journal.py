# -*- coding: utf-8 -*-
import io
import re
import logging
from datetime import datetime, date

from odoo import models, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    _logger.warning("openpyxl not installed. BIM import will not work.")

try:
    import xlrd
    XLRD_AVAILABLE = True
except ImportError:
    XLRD_AVAILABLE = False
    _logger.warning("xlrd not installed. BCI import will not work.")


class AccountJournal(models.Model):
    _inherit = 'account.journal'

    def _get_bank_statements_available_import_formats(self):
        """Adiciona formatos BIM, BCI, BCI Cartão Crédito, BIM Pré-Pago e Mozabanco à lista de formatos disponíveis"""
        rslt = super()._get_bank_statements_available_import_formats()
        rslt.extend(['BIM (xlsx)', 'BCI (xls)', 'BCI Cartão Crédito (xls)', 'BIM Pré-Pago (csv)', 'Mozabanco (csv)'])
        return rslt

    # ==========================================
    # DETECÇÃO DE FORMATO
    # ==========================================

    def _check_bim(self, attachment):
        """
        Verifica se o ficheiro é um extracto BIM.
        BIM usa formato xlsx e tem padrões específicos.
        """
        if not OPENPYXL_AVAILABLE:
            return False

        # Verificar extensão
        filename = attachment.name or ''
        if not filename.lower().endswith('.xlsx'):
            return False

        try:
            workbook = openpyxl.load_workbook(io.BytesIO(attachment.raw), read_only=True, data_only=True)
            sheet = workbook.active

            # BIM tem "Actividades da Conta" na linha 2
            # e "Lista de transacções da conta Nº" na linha 6
            for row_idx in range(1, 10):
                for col_idx in range(1, 5):
                    cell_value = sheet.cell(row=row_idx, column=col_idx).value
                    if cell_value:
                        cell_str = str(cell_value)
                        if 'Actividades da Conta' in cell_str or 'Lista de transac' in cell_str:
                            workbook.close()
                            _logger.info("Ficheiro detectado como BIM: %s", filename)
                            return True

            workbook.close()
            return False
        except Exception as e:
            _logger.debug("Erro ao verificar BIM: %s", e)
            return False

    def _check_bci(self, attachment):
        """
        Verifica se o ficheiro é um extracto BCI.
        BCI usa formato xls e tem padrões específicos.
        """
        if not XLRD_AVAILABLE:
            return False

        # Verificar extensão
        filename = attachment.name or ''
        if not filename.lower().endswith('.xls'):
            return False

        try:
            workbook = xlrd.open_workbook(file_contents=attachment.raw)
            sheet = workbook.sheet_by_index(0)

            # BCI tem "Conta" na primeira coluna e "Nib" logo abaixo
            for row_idx in range(min(5, sheet.nrows)):
                cell_value = sheet.cell_value(row_idx, 0)
                if cell_value:
                    cell_str = str(cell_value).strip()
                    if cell_str == 'Conta' or cell_str == 'Nib':
                        _logger.info("Ficheiro detectado como BCI: %s", filename)
                        return True

            return False
        except Exception as e:
            _logger.debug("Erro ao verificar BCI: %s", e)
            return False

    def _check_bci_credit_card(self, attachment):
        """
        Verifica se o ficheiro é um extracto de Cartão de Crédito BCI.
        Diferença do BCI normal: tem "Saldo Anterior" mas não tem "Nib".
        """
        if not XLRD_AVAILABLE:
            return False

        filename = attachment.name or ''
        if not filename.lower().endswith('.xls'):
            return False

        try:
            workbook = xlrd.open_workbook(file_contents=attachment.raw)
            sheet = workbook.sheet_by_index(0)

            has_conta = False
            has_saldo_anterior = False
            has_nib = False

            for row_idx in range(min(6, sheet.nrows)):
                cell_value = str(sheet.cell_value(row_idx, 0)).strip()
                if cell_value == 'Conta':
                    has_conta = True
                elif cell_value == 'Nib':
                    has_nib = True
                elif cell_value == 'Saldo Anterior':
                    has_saldo_anterior = True

            # Cartão de crédito: tem Conta e Saldo Anterior, mas NÃO tem Nib
            if has_conta and has_saldo_anterior and not has_nib:
                _logger.info("Ficheiro detectado como BCI Cartão Crédito: %s", filename)
                return True

            return False
        except Exception as e:
            _logger.debug("Erro ao verificar BCI Cartão Crédito: %s", e)
            return False

    def _check_bim_prepaid_card(self, attachment):
        """
        Verifica se o ficheiro é um extracto de Cartão Pré-Pago BIM.
        Formato CSV com separador ; e encoding latin-1.
        """
        filename = attachment.name or ''
        if not filename.lower().endswith('.csv'):
            return False

        try:
            # Tentar ler como CSV com encoding latin-1
            content = attachment.raw.decode('latin-1')
            lines = content.strip().split('\n')

            if len(lines) < 2:
                return False

            # Verificar header do BIM Pré-Pago
            header = lines[0].lower()
            if 'lançamento' in header or 'lancamento' in header:
                if 'montante' in header or 'descri' in header:
                    _logger.info("Ficheiro detectado como BIM Cartão Pré-Pago: %s", filename)
                    return True

            return False
        except Exception as e:
            _logger.debug("Erro ao verificar BIM Cartão Pré-Pago: %s", e)
            return False

    def _check_mozabanco(self, attachment):
        """
        Verifica se o ficheiro é um extracto do Mozabanco.
        Formato CSV com separador ; e estrutura específica:
        - Linha 1: sep=;
        - Linha 2: Número da conta
        - Linha 3: Saldo de Abertura
        - Linha 4: Saldo de fecho
        - Linha 5: Moeda
        - Linha 6: Headers (Data;Data-Valor;Descrição;Nº de Ref;Débito;Crédito;Saldo)
        """
        filename = attachment.name or ''
        if not filename.lower().endswith('.csv'):
            return False

        try:
            # Tentar diferentes encodings
            content = None
            for encoding in ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252']:
                try:
                    content = attachment.raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue

            if not content:
                return False

            lines = content.strip().split('\n')

            if len(lines) < 6:
                return False

            # Verificar padrões do Mozabanco
            has_sep = 'sep=' in lines[0].lower()
            has_saldo_abertura = 'saldo de abertura' in lines[2].lower() if len(lines) > 2 else False
            has_saldo_fecho = 'saldo de fecho' in lines[3].lower() if len(lines) > 3 else False

            # Verificar header na linha 6 (índice 5)
            header_line = lines[5].lower() if len(lines) > 5 else ''
            has_correct_header = 'data' in header_line and 'débito' in header_line or 'debito' in header_line

            if (has_saldo_abertura or has_saldo_fecho) and has_correct_header:
                _logger.info("Ficheiro detectado como Mozabanco: %s", filename)
                return True

            return False
        except Exception as e:
            _logger.debug("Erro ao verificar Mozabanco: %s", e)
            return False

    # ==========================================
    # PARSER BIM
    # ==========================================

    def _parse_bim_amount(self, value):
        """
        Converte valor BIM para float.
        BIM usa formato europeu: 1.234,56 ou -1.234,56
        """
        if value is None or value == '' or value == 0:
            return 0.0

        if isinstance(value, (int, float)):
            return float(value)

        # String: remover espaços e converter
        value_str = str(value).strip()
        if not value_str:
            return 0.0

        # Remover pontos de milhar e trocar vírgula por ponto
        value_str = value_str.replace('.', '').replace(',', '.')
        try:
            return float(value_str)
        except ValueError:
            return 0.0

    def _parse_bim_date(self, value):
        """Converte data BIM para date object"""
        if value is None:
            return None

        if isinstance(value, datetime):
            return value.date()

        if isinstance(value, str):
            # Formato: 2026-01-12 00:00:00
            try:
                return datetime.strptime(value.split()[0], '%Y-%m-%d').date()
            except ValueError:
                pass

            # Formato: 12-01-2026
            try:
                return datetime.strptime(value, '%d-%m-%Y').date()
            except ValueError:
                pass

        return None

    def _extract_bim_info_from_description(self, description):
        """
        Extrai informações da descrição BIM.

        Exemplo: "ALLPORT STONE TRADINGRef. 260112104523186Código do Canal SCI IPNome do Canal..."

        Retorna: {
            'partner_name': 'ALLPORT STONE TRADING',
            'ref': '260112104523186',
            'ref_adicional': '00005 FE1',
            'canal': 'INTERNET',
            'note': 'descrição completa limpa'
        }
        """
        result = {
            'partner_name': None,
            'ref': None,
            'ref_adicional': None,
            'canal': None,
            'note': None,
        }

        if not description:
            return result

        desc = str(description).strip()

        # Extrair referência principal: "Ref. 260112104523186"
        ref_match = re.search(r'Ref\.\s*(\d+)', desc)
        if ref_match:
            result['ref'] = ref_match.group(1)

        # Extrair referência adicional: "Ref. Adicional 00005 FE1"
        ref_adic_match = re.search(r'Ref\.\s*Adicional\s*([A-Z0-9\s]+?)(?:C[oó]digo|$)', desc, re.IGNORECASE)
        if ref_adic_match:
            result['ref_adicional'] = ref_adic_match.group(1).strip()

        # Extrair nome do canal: "Nome do Canal SCI INTERNET"
        canal_match = re.search(r'Nome do Canal SCI\s*(\w+)', desc, re.IGNORECASE)
        if canal_match:
            result['canal'] = canal_match.group(1)

        # Extrair nome do parceiro (texto antes de "Ref.")
        if 'Ref.' in desc:
            partner_part = desc.split('Ref.')[0].strip()
            # Limpar prefixos comuns
            for prefix in ['COM.', 'COM ', 'TRF', 'LEV', 'POS', 'IS ', 'CCC']:
                if partner_part.upper().startswith(prefix):
                    partner_part = partner_part[len(prefix):].strip()

            # Se ainda tem texto significativo, usar como partner_name
            if len(partner_part) > 3 and not partner_part.startswith('0'):
                result['partner_name'] = partner_part[:100]

        # Se não encontrou parceiro antes de Ref, tentar extrair do início
        if not result['partner_name']:
            # Padrões conhecidos de nomes
            patterns = [
                r'^([A-Z][A-Z\s]+(?:TRADING|STONE|LDA|SA))(?:Ref|C[oó]digo)',
                r'^([A-Z][A-Z\s]{5,})(?:Ref|C[oó]digo)',
            ]
            for pattern in patterns:
                match = re.search(pattern, desc)
                if match:
                    result['partner_name'] = match.group(1).strip()[:100]
                    break

        # Criar nota limpa (remover códigos técnicos)
        note = desc
        # Remover partes técnicas
        note = re.sub(r'C[oó]digo do Canal SCI\s*\w+', '', note, flags=re.IGNORECASE)
        note = re.sub(r'Nome do Canal SCI\s*\w+', '', note, flags=re.IGNORECASE)
        note = re.sub(r'C[oó]digo Transac[çc][aã]o SCI\s*\w+', '', note, flags=re.IGNORECASE)
        note = re.sub(r'Ref\.\s*Adicional\s*[A-Z0-9\s]+', '', note, flags=re.IGNORECASE)
        note = re.sub(r'Ref\.\s*\d+', '', note)
        note = re.sub(r'\s+', ' ', note).strip()
        result['note'] = note if note else None

        return result

    def _parse_bank_statement_file_bim(self, attachment):
        """
        Parseia extracto BIM (xlsx).

        Estrutura BIM:
        - Linha 6: Info da conta e período
        - Linha 7: Saldo Inicial
        - Linha 8: Créditos
        - Linha 9: Débitos
        - Linha 10: Saldo Final
        - Linha 11: Headers
        - Linha 12+: Dados
        - Colunas: C(3)=Data Trans, D(4)=Data Valor, E(5)=Descrição, G(7)=Débito, I(9)=Crédito, J(10)=Saldo
        """
        _logger.info("Iniciando parse BIM: %s", attachment.name)

        workbook = openpyxl.load_workbook(io.BytesIO(attachment.raw), read_only=True, data_only=True)
        sheet = workbook.active

        # Extrair informações do cabeçalho
        account_number = None
        currency_code = 'MZN'
        balance_start = 0.0
        balance_end_header = 0.0
        period_start = None
        period_end = None
        statement_name = None

        # Procurar linha com info da conta (normalmente linha 6)
        for row_idx in range(1, 15):
            cell_value = sheet.cell(row=row_idx, column=3).value  # Coluna C
            if cell_value and 'Lista de transac' in str(cell_value):
                cell_str = str(cell_value)

                # Extrair número da conta: "conta Nº 638537033"
                match = re.search(r'conta\s*N[º°]\s*(\d+)', cell_str)
                if match:
                    account_number = match.group(1)

                # Extrair moeda
                if 'MZN' in cell_str:
                    currency_code = 'MZN'
                elif 'USD' in cell_str:
                    currency_code = 'USD'
                elif 'EUR' in cell_str:
                    currency_code = 'EUR'

                # Extrair período: "de 01-01-2026 até 12-01-2026"
                period_match = re.search(r'de\s*(\d{2}-\d{2}-\d{4})\s*at[eé]\s*(\d{2}-\d{2}-\d{4})', cell_str)
                if period_match:
                    try:
                        period_start = datetime.strptime(period_match.group(1), '%d-%m-%Y').date()
                        period_end = datetime.strptime(period_match.group(2), '%d-%m-%Y').date()
                    except ValueError:
                        pass
                break

        # Procurar Saldo Inicial e Saldo Final
        for row_idx in range(1, 15):
            cell_value = sheet.cell(row=row_idx, column=3).value
            if cell_value:
                cell_str = str(cell_value)
                if 'Saldo Inicial' in cell_str:
                    balance_val = sheet.cell(row=row_idx, column=9).value
                    _logger.info("BIM - Saldo Inicial raw (linha %s): %s (tipo: %s)",
                                 row_idx, balance_val, type(balance_val).__name__)
                    balance_start = self._parse_bim_amount(balance_val)
                elif 'Saldo Final' in cell_str:
                    balance_val = sheet.cell(row=row_idx, column=9).value
                    _logger.info("BIM - Saldo Final raw (linha %s): %s (tipo: %s)",
                                 row_idx, balance_val, type(balance_val).__name__)
                    balance_end_header = self._parse_bim_amount(balance_val)

        # Gerar nome do statement
        if period_end:
            statement_name = f"BIM Extracto {period_end.strftime('%Y-%m-%d')}"
        else:
            statement_name = f"BIM Extracto {date.today().strftime('%Y-%m-%d')}"

        _logger.info("BIM - Conta: %s, Moeda: %s, Saldo Inicial: %s, Saldo Final: %s",
                     account_number, currency_code, balance_start, balance_end_header)

        # Encontrar linha de headers
        header_row = None
        for row_idx in range(1, 20):
            cell_value = sheet.cell(row=row_idx, column=3).value
            _logger.debug("BIM - Linha %s, Coluna C: %s", row_idx, cell_value)
            if cell_value and 'Data' in str(cell_value) and 'Transa' in str(cell_value):
                header_row = row_idx
                _logger.info("BIM - Header encontrado na linha %s", header_row)
                break

        if not header_row:
            _logger.warning("BIM - Header não encontrado, tentando alternativa...")
            # Tentar encontrar só "Data Trans"
            for row_idx in range(1, 20):
                cell_value = sheet.cell(row=row_idx, column=3).value
                if cell_value and 'Data' in str(cell_value):
                    header_row = row_idx
                    _logger.info("BIM - Header alternativo encontrado na linha %s", header_row)
                    break

        if not header_row:
            workbook.close()
            raise UserError(_("Não foi possível encontrar a linha de cabeçalho no ficheiro BIM."))

        # Parse transações
        transactions = []
        data_start_row = header_row + 1
        first_transaction_balance = None  # Saldo da primeira transação (é o saldo final no BIM)
        last_transaction_saldo = None     # Saldo da última transação
        last_transaction_amount = None    # Amount da última transação

        _logger.info("BIM - Iniciando parse de transações a partir da linha %s", data_start_row)

        empty_rows_count = 0
        max_empty_rows = 20  # Parar após 20 linhas vazias consecutivas

        # Usar iter_rows para leitura eficiente (muito mais rápido que cell())
        # Colunas: C(3)=Data Trans, D(4)=Data Valor, E(5)=Descrição, G(7)=Débito, I(9)=Crédito, J(10)=Saldo
        for row_idx, row in enumerate(sheet.iter_rows(min_row=data_start_row, max_col=11), start=data_start_row):
            # Coluna C (3): Data Transação (índice 2 no array)
            date_trans = row[2].value if len(row) > 2 else None
            if not date_trans:
                empty_rows_count += 1
                if empty_rows_count >= max_empty_rows:
                    _logger.info("BIM - Parando parse após %s linhas vazias consecutivas na linha %s",
                                 max_empty_rows, row_idx)
                    break
                continue

            # Reset contador de linhas vazias
            empty_rows_count = 0

            date_parsed = self._parse_bim_date(date_trans)
            if not date_parsed:
                continue

            # Coluna D (4): Data Valor (índice 3)
            date_valor_raw = row[3].value if len(row) > 3 else None
            date_valor = self._parse_bim_date(date_valor_raw) if date_valor_raw else date_parsed

            # Coluna E (5): Descrição (índice 4)
            description = row[4].value if len(row) > 4 else ''

            # Coluna G (7): Débito (índice 6)
            debit = self._parse_bim_amount(row[6].value if len(row) > 6 else None)

            # Coluna I (9): Crédito (índice 8)
            credit = self._parse_bim_amount(row[8].value if len(row) > 8 else None)

            # Coluna J (10): Saldo (índice 9)
            saldo = self._parse_bim_amount(row[9].value if len(row) > 9 else None)

            # Guardar saldo da primeira transação (é o saldo final no BIM)
            if first_transaction_balance is None and saldo:
                first_transaction_balance = saldo
                _logger.info("BIM - Saldo da primeira transação (saldo final): %s", first_transaction_balance)

            # Calcular amount (débito é negativo no BIM)
            amount = credit + debit  # debit já vem negativo

            # Guardar dados da última transação (para calcular saldo inicial)
            if saldo:
                last_transaction_saldo = saldo
                last_transaction_amount = amount

            if amount == 0:
                continue

            # Extrair informações da descrição
            info = self._extract_bim_info_from_description(description)

            # Gerar unique_import_id usando referência se disponível
            if info['ref']:
                unique_id = f"BIM-{account_number}-{info['ref']}"
            else:
                unique_id = f"BIM-{account_number}-{date_parsed.strftime('%Y%m%d')}-{row_idx}-{abs(amount):.2f}"

            # Criar payment_ref limpo
            payment_ref = info['note'] or str(description).strip()
            if len(payment_ref) > 200:
                payment_ref = payment_ref[:197] + '...'

            transactions.append({
                'date': date_parsed,
                'date_valor': date_valor,
                'payment_ref': payment_ref or 'Sem descrição',
                'amount': amount,
                'debit': abs(debit) if debit < 0 else 0.0,
                'credit': credit if credit > 0 else 0.0,
                'balance': saldo,
                'unique_import_id': unique_id,
                'ref': info['ref'],
                'partner_name': info['partner_name'],
                'note': f"Canal: {info['canal'] or 'N/A'}",
                'sequence': len(transactions) + 1,
            })

        workbook.close()

        _logger.info("BIM - %d transações encontradas", len(transactions))

        if not transactions:
            raise UserError(_("Nenhuma transação encontrada no ficheiro BIM."))

        # SALDO FINAL: primeira transação (no BIM a primeira linha é a mais recente)
        final_balance = first_transaction_balance or 0.0
        _logger.info("BIM - Saldo final (primeira transação): %s", final_balance)

        # SALDO INICIAL: calculado a partir da última transação
        # Fórmula: saldo_inicial = saldo_ultima_linha - amount_ultima_linha
        if last_transaction_saldo is not None and last_transaction_amount is not None:
            balance_start = last_transaction_saldo - last_transaction_amount
            _logger.info("BIM - Saldo inicial calculado: %s (saldo=%s - amount=%s)",
                        balance_start, last_transaction_saldo, last_transaction_amount)
        else:
            balance_start = 0.0
            _logger.warning("BIM - Não foi possível calcular saldo inicial, usando 0.0")

        vals_bank_statement = [{
            'name': statement_name,
            'date': period_end or date.today(),
            'balance_start': balance_start,
            'balance_end_real': final_balance,
            'transactions': transactions,
        }]

        return currency_code, account_number, vals_bank_statement

    # ==========================================
    # PARSER BCI
    # ==========================================

    def _parse_bci_amount(self, value):
        """
        Converte valor BCI para float.
        BCI usa formato: +50.000,00 ou -2.548,91
        """
        if value is None or value == '':
            return 0.0

        if isinstance(value, (int, float)):
            return float(value)

        value_str = str(value).strip()
        if not value_str:
            return 0.0

        # Remover MZN, USD, etc.
        value_str = re.sub(r'[A-Za-z]', '', value_str).strip()

        # Remover espaços
        value_str = value_str.replace(' ', '')

        # Manter sinal
        is_negative = value_str.startswith('-')
        is_positive = value_str.startswith('+')

        # Remover sinal para processar
        value_str = value_str.lstrip('+-')

        # Converter formato europeu: 50.000,00 -> 50000.00
        value_str = value_str.replace('.', '').replace(',', '.')

        try:
            result = float(value_str)
            if is_negative:
                result = -abs(result)
            return result
        except ValueError:
            return 0.0

    def _parse_bci_date(self, value):
        """Converte data BCI para date object"""
        if value is None:
            return None

        if isinstance(value, datetime):
            return value.date()

        if isinstance(value, str):
            value = value.strip()

            # Formato: 01-01-2026
            try:
                return datetime.strptime(value, '%d-%m-%Y').date()
            except ValueError:
                pass

            # Formato: 01/01/2026
            try:
                return datetime.strptime(value, '%d/%m/%Y').date()
            except ValueError:
                pass

        return None

    def _extract_bci_info_from_description(self, description):
        """
        Extrai informações da descrição BCI.

        Exemplos de descrições BCI:
        - "TRF.NAC./MULTIPLO LOTES: 0002 REF:25118054033631"
        - "TRF.AL./MULTIPLO LOTES: 0"
        - "COMPRA POS-PAGAMENTO-MPESA T00"
        - "LEVI MH LDA"

        Retorna: {
            'partner_name': nome do parceiro se identificável,
            'ref': referência da transação,
            'note': descrição limpa
        }
        """
        result = {
            'partner_name': None,
            'ref': None,
            'note': None,
        }

        if not description:
            return result

        desc = str(description).strip()

        # Extrair referência: "REF:25118054033631" ou "REF: 123456"
        ref_match = re.search(r'REF[:\s]*(\d+)', desc, re.IGNORECASE)
        if ref_match:
            result['ref'] = ref_match.group(1)

        # Extrair nome do parceiro
        partner_name = None

        # Padrão 1: Nome de empresa conhecida (LDA, SA, etc)
        company_match = re.search(r'([A-Z][A-Z\s]+(?:LDA|SA|TRADING|STONE))', desc, re.IGNORECASE)
        if company_match:
            partner_name = company_match.group(1).strip()

        # Padrão 2: TRF para nome: "TRF.NAC./MULTIPLO LOTES" -> nome pode estar em outra parte
        if not partner_name and 'TRF' in desc.upper():
            # Procurar texto após "LOTES:" ou similar
            lotes_match = re.search(r'LOTES[:\s]*\d+\s*(.+?)(?:REF|$)', desc, re.IGNORECASE)
            if lotes_match and len(lotes_match.group(1).strip()) > 3:
                partner_name = lotes_match.group(1).strip()

        # Padrão 3: Texto simples que parece ser nome
        if not partner_name:
            # Se não tem palavras-chave técnicas, pode ser nome direto
            technical_keywords = ['TRF', 'POS', 'COMPRA', 'PAGAMENTO', 'MPESA', 'ATM', 'LEV', 'DEP']
            is_technical = any(kw in desc.upper() for kw in technical_keywords)
            if not is_technical and len(desc) > 3 and desc[0].isalpha():
                partner_name = desc[:100]

        if partner_name and len(partner_name) > 3:
            result['partner_name'] = partner_name[:100]

        # Criar nota limpa
        note = desc
        # Remover REF:xxxxx
        note = re.sub(r'REF[:\s]*\d+', '', note, flags=re.IGNORECASE).strip()
        result['note'] = note if note else None

        return result

    def _parse_bank_statement_file_bci(self, attachment):
        """
        Parseia extracto BCI (xls).

        Estrutura BCI:
        - Linha 2: Conta (coluna B)
        - Linha 3: NIB (coluna B)
        - Linha 4: Saldo Anterior (coluna B)
        - Linha 9: Headers
        - Linha 10+: Dados
        - Colunas: A=Data Mov, B=Num Doc, C=Num Oper, D=Data Valor, E=Descrição, F=Valor, G=Saldo, H=Moeda
        """
        _logger.info("Iniciando parse BCI: %s", attachment.name)

        workbook = xlrd.open_workbook(file_contents=attachment.raw)
        sheet = workbook.sheet_by_index(0)

        # Extrair informações do cabeçalho
        account_number = None
        nib = None
        currency_code = 'MZN'
        balance_start = 0.0
        statement_name = None
        first_date = None
        last_date = None

        for row_idx in range(min(10, sheet.nrows)):
            cell_label = str(sheet.cell_value(row_idx, 0)).strip()

            if cell_label == 'Conta':
                account_number = str(sheet.cell_value(row_idx, 1)).strip()
            elif cell_label == 'Nib':
                nib = str(sheet.cell_value(row_idx, 1)).strip()
            elif cell_label == 'Saldo Anterior':
                balance_val = sheet.cell_value(row_idx, 1)
                balance_start = self._parse_bci_amount(balance_val)

        _logger.info("BCI - Conta: %s, NIB: %s, Moeda: %s, Saldo Anterior: %s",
                     account_number, nib, currency_code, balance_start)

        # Encontrar linha de headers (procura "Data Mov.")
        header_row = None
        for row_idx in range(sheet.nrows):
            cell_value = str(sheet.cell_value(row_idx, 0)).strip()
            if 'Data Mov' in cell_value or 'Data mov' in cell_value.lower():
                header_row = row_idx
                break

        if header_row is None:
            raise UserError(_("Não foi possível encontrar a linha de cabeçalho no ficheiro BCI."))

        # Parse transações
        transactions = []
        data_start_row = header_row + 1
        balance_end = balance_start
        first_transaction_saldo = None    # Saldo da primeira transação
        first_transaction_amount = None   # Amount da primeira transação

        for row_idx in range(data_start_row, sheet.nrows):
            # Coluna A (0): Data Movimento
            date_mov = sheet.cell_value(row_idx, 0)
            if not date_mov:
                continue

            date_parsed = self._parse_bci_date(date_mov)
            if not date_parsed:
                continue

            # Registrar primeira e última data
            if first_date is None:
                first_date = date_parsed
            last_date = date_parsed

            # Coluna B (1): Número Documento
            num_doc = str(sheet.cell_value(row_idx, 1)).strip() if sheet.cell_value(row_idx, 1) else ''

            # Coluna C (2): Número Operação
            num_oper = str(sheet.cell_value(row_idx, 2)).strip() if sheet.cell_value(row_idx, 2) else ''

            # Coluna D (3): Data Valor
            date_valor_raw = sheet.cell_value(row_idx, 3)
            date_valor = self._parse_bci_date(date_valor_raw) if date_valor_raw else date_parsed

            # Coluna E (4): Descrição
            description = str(sheet.cell_value(row_idx, 4)).strip() if sheet.cell_value(row_idx, 4) else ''

            # Coluna F (5): Valor (já com sinal)
            amount = self._parse_bci_amount(sheet.cell_value(row_idx, 5))

            # Coluna G (6): Saldo
            saldo = self._parse_bci_amount(sheet.cell_value(row_idx, 6))

            # Guardar dados da primeira transação (para calcular saldo inicial)
            # BCI tem ordem cronológica: primeira linha = mais antiga
            if first_transaction_saldo is None and saldo is not None:
                first_transaction_saldo = saldo
                first_transaction_amount = amount
                _logger.info("BCI - Primeira transação: saldo=%s, amount=%s", saldo, amount)

            # Coluna H (7): Moeda
            moeda_cell = sheet.cell_value(row_idx, 7) if sheet.ncols > 7 else None
            if moeda_cell:
                moeda = str(moeda_cell).strip().upper()
                if moeda in ['MZN', 'USD', 'EUR']:
                    currency_code = moeda

            if amount == 0:
                continue

            # Extrair informações da descrição
            info = self._extract_bci_info_from_description(description)

            # Gerar unique_import_id usando referência se disponível
            if info['ref']:
                unique_id = f"BCI-{account_number}-{info['ref']}"
            elif num_oper:
                unique_id = f"BCI-{account_number}-{num_doc}-{num_oper}"
            else:
                unique_id = f"BCI-{account_number}-{date_parsed.strftime('%Y%m%d')}-{row_idx}-{abs(amount):.2f}"

            # Compor descrição completa para payment_ref
            payment_ref = info['note'] or description
            if num_doc:
                payment_ref = f"[{num_doc}] {payment_ref}"
            if len(payment_ref) > 200:
                payment_ref = payment_ref[:197] + '...'

            transactions.append({
                'date': date_parsed,
                'date_valor': date_valor,
                'payment_ref': description or 'Sem descrição',
                'amount': amount,
                'debit': abs(amount) if amount < 0 else 0.0,
                'credit': amount if amount > 0 else 0.0,
                'balance': saldo,
                'num_doc': num_doc,
                'num_oper': num_oper,
                'currency': currency_code,
                'unique_import_id': unique_id,
                'ref': info['ref'] or num_oper,
                'partner_name': info['partner_name'],
                'note': '',
                'sequence': len(transactions) + 1,
            })

            balance_end = saldo

        _logger.info("BCI - %d transações encontradas", len(transactions))

        if not transactions:
            raise UserError(_("Nenhuma transação encontrada no ficheiro BCI."))

        # SALDO INICIAL: calculado a partir da primeira transação
        # BCI ordem cronológica: primeira linha = mais antiga
        # Fórmula: saldo_inicial = saldo_primeira_linha - amount_primeira_linha
        if first_transaction_saldo is not None and first_transaction_amount is not None:
            balance_start = first_transaction_saldo - first_transaction_amount
            _logger.info("BCI - Saldo inicial calculado: %s (saldo=%s - amount=%s)",
                         balance_start, first_transaction_saldo, first_transaction_amount)
        else:
            _logger.warning("BCI - Não foi possível calcular saldo inicial, usando cabeçalho: %s", balance_start)

        # Gerar nome do statement
        if last_date:
            statement_name = f"BCI Extracto {last_date.strftime('%Y-%m-%d')}"
        else:
            statement_name = f"BCI Extracto {date.today().strftime('%Y-%m-%d')}"

        vals_bank_statement = [{
            'name': statement_name,
            'date': last_date or date.today(),
            'balance_start': balance_start,
            'balance_end_real': balance_end,
            'transactions': transactions,
        }]

        return currency_code, account_number, vals_bank_statement

    # ==========================================
    # PARSER BCI CARTÃO DE CRÉDITO
    # ==========================================

    def _parse_bank_statement_file_bci_credit_card(self, attachment):
        """
        Parseia extracto de Cartão de Crédito BCI (xls).

        Estrutura:
        - Linha 0: Período "DE 01/01/2026 ATÉ 01/12/2026"
        - Linha 1: Conta (número do cartão)
        - Linha 2: Saldo Anterior
        - Linha 5: Headers
        - Linha 6+: Dados
        - Colunas: A=Data Mov, B=Data Valor, C=Descrição, D=Valor, E=Moeda, F=Saldo, G=Moeda
        """
        _logger.info("Iniciando parse BCI Cartão Crédito: %s", attachment.name)

        workbook = xlrd.open_workbook(file_contents=attachment.raw)
        sheet = workbook.sheet_by_index(0)

        # Extrair informações do cabeçalho
        account_number = None
        currency_code = 'MZN'
        balance_start = 0.0
        period_start = None
        period_end = None
        statement_name = None

        for row_idx in range(min(6, sheet.nrows)):
            cell_label = str(sheet.cell_value(row_idx, 0)).strip()

            if row_idx == 0 and 'DE ' in cell_label and ' AT' in cell_label:
                # Extrair período: "DE 01/01/2026 ATÉ 01/12/2026"
                period_match = re.search(r'DE\s*(\d{2}/\d{2}/\d{4})\s*AT[EÉ]\s*(\d{2}/\d{2}/\d{4})', cell_label, re.IGNORECASE)
                if period_match:
                    try:
                        period_start = datetime.strptime(period_match.group(1), '%d/%m/%Y').date()
                        period_end = datetime.strptime(period_match.group(2), '%d/%m/%Y').date()
                    except ValueError:
                        pass
            elif cell_label == 'Conta':
                account_number = str(sheet.cell_value(row_idx, 1)).strip()
            elif cell_label == 'Saldo Anterior':
                balance_val = sheet.cell_value(row_idx, 1)
                balance_start = self._parse_bci_amount(balance_val)

        _logger.info("BCI Cartão - Conta: %s, Moeda: %s, Saldo Anterior: %s",
                     account_number, currency_code, balance_start)

        # Encontrar linha de headers (procura "Data Mov.")
        header_row = None
        for row_idx in range(sheet.nrows):
            cell_value = str(sheet.cell_value(row_idx, 0)).strip()
            if 'Data Mov' in cell_value:
                header_row = row_idx
                break

        if header_row is None:
            raise UserError(_("Não foi possível encontrar a linha de cabeçalho no ficheiro BCI Cartão."))

        # Parse transações
        transactions = []
        data_start_row = header_row + 1
        balance_end = balance_start
        last_date = None
        first_transaction_saldo = None    # Saldo da primeira transação
        first_transaction_amount = None   # Amount da primeira transação

        for row_idx in range(data_start_row, sheet.nrows):
            # Coluna A (0): Data Movimento
            date_mov = sheet.cell_value(row_idx, 0)
            if not date_mov:
                continue

            date_parsed = self._parse_bci_date(date_mov)
            if not date_parsed:
                continue

            last_date = date_parsed

            # Coluna B (1): Data Valor
            date_valor_raw = sheet.cell_value(row_idx, 1)
            date_valor = self._parse_bci_date(date_valor_raw) if date_valor_raw else date_parsed

            # Coluna C (2): Descrição
            description = str(sheet.cell_value(row_idx, 2)).strip() if sheet.cell_value(row_idx, 2) else ''

            # Coluna D (3): Valor (com sinal +/-)
            amount = self._parse_bci_amount(sheet.cell_value(row_idx, 3))

            # Coluna E (4): Moeda
            moeda_cell = sheet.cell_value(row_idx, 4)
            if moeda_cell:
                moeda = str(moeda_cell).strip().upper()
                if moeda in ['MZN', 'USD', 'EUR']:
                    currency_code = moeda

            # Coluna F (5): Saldo
            saldo = self._parse_bci_amount(sheet.cell_value(row_idx, 5))

            # Guardar dados da primeira transação (para calcular saldo inicial)
            # BCI Cartão tem ordem cronológica: primeira linha = mais antiga
            if first_transaction_saldo is None and saldo is not None:
                first_transaction_saldo = saldo
                first_transaction_amount = amount
                _logger.info("BCI Cartão - Primeira transação: saldo=%s, amount=%s", saldo, amount)

            if amount == 0:
                continue

            # Gerar unique_import_id
            unique_id = f"BCI-CC-{account_number}-{date_parsed.strftime('%Y%m%d')}-{row_idx}-{abs(amount):.2f}"

            transactions.append({
                'date': date_parsed,
                'date_valor': date_valor,
                'payment_ref': description or 'Sem descrição',
                'amount': amount,
                'debit': abs(amount) if amount < 0 else 0.0,
                'credit': amount if amount > 0 else 0.0,
                'balance': saldo,
                'num_doc': '',
                'num_oper': '',
                'currency': currency_code,
                'unique_import_id': unique_id,
                'ref': '',
                'partner_name': '',
                'note': '',
                'sequence': len(transactions) + 1,
            })

            balance_end = saldo

        _logger.info("BCI Cartão - %d transações encontradas", len(transactions))

        if not transactions:
            raise UserError(_("Nenhuma transação encontrada no ficheiro BCI Cartão."))

        # SALDO INICIAL: calculado a partir da primeira transação
        # BCI Cartão ordem cronológica: primeira linha = mais antiga
        # Fórmula: saldo_inicial = saldo_primeira_linha - amount_primeira_linha
        if first_transaction_saldo is not None and first_transaction_amount is not None:
            balance_start = first_transaction_saldo - first_transaction_amount
            _logger.info("BCI Cartão - Saldo inicial calculado: %s (saldo=%s - amount=%s)",
                         balance_start, first_transaction_saldo, first_transaction_amount)
        else:
            _logger.warning("BCI Cartão - Não foi possível calcular saldo inicial, usando cabeçalho: %s", balance_start)

        # Gerar nome do statement
        if last_date:
            statement_name = f"BCI Cartão Crédito {account_number} {last_date.strftime('%Y-%m-%d')}"
        else:
            statement_name = f"BCI Cartão Crédito {account_number} {date.today().strftime('%Y-%m-%d')}"

        vals_bank_statement = [{
            'name': statement_name,
            'date': last_date or date.today(),
            'balance_start': balance_start,
            'balance_end_real': balance_end,
            'transactions': transactions,
        }]

        return currency_code, account_number, vals_bank_statement

    # ==========================================
    # PARSER BIM CARTÃO PRÉ-PAGO (CSV)
    # ==========================================

    def _parse_bim_prepaid_amount(self, value):
        """
        Converte valor do BIM Pré-Pago para float.
        Formato: +1.277,52 MZN ou -1.277,52 MZN
        """
        if not value:
            return 0.0

        value_str = str(value).strip()

        # Remover MZN, USD, EUR e espaços
        value_str = re.sub(r'[A-Za-z]', '', value_str).strip()

        # Remover caracteres especiais (possíveis problemas de encoding)
        value_str = re.sub(r'[^\d.,+-]', '', value_str)

        if not value_str:
            return 0.0

        # Manter sinal
        is_negative = '-' in value_str
        value_str = value_str.replace('+', '').replace('-', '')

        # Converter formato europeu: 1.277,52 -> 1277.52
        value_str = value_str.replace('.', '').replace(',', '.')

        try:
            result = float(value_str)
            return -result if is_negative else result
        except ValueError:
            return 0.0

    def _parse_bank_statement_file_bim_prepaid_card(self, attachment):
        """
        Parseia extracto de Cartão Pré-Pago BIM (CSV).

        Estrutura CSV:
        - Separador: ;
        - Encoding: latin-1
        - Colunas: Lançamento de dados;Valor dos dados;Descrição;Montante;Monte Original

        SALDOS:
        - Saldo Inicial: busca do último statement do diário
        - Saldo Final: Saldo Inicial + Créditos - Débitos
        """
        _logger.info("Iniciando parse BIM Cartão Pré-Pago: %s", attachment.name)

        # Ler conteúdo CSV
        content = attachment.raw.decode('latin-1')
        lines = content.strip().split('\n')

        if len(lines) < 2:
            raise UserError(_("Ficheiro CSV vazio ou inválido."))

        # Parse header
        header = lines[0].split(';')
        _logger.info("BIM Pré-Pago - Header: %s", header)

        # Informações do extracto
        currency_code = 'MZN'
        account_number = 'PREPAID'
        first_date = None
        total_credit = 0.0
        total_debit = 0.0

        # Parse transações
        transactions = []

        for row_idx, line in enumerate(lines[1:], start=2):
            if not line.strip():
                continue

            cols = line.split(';')
            if len(cols) < 4:
                continue

            # Coluna 0: Lançamento de dados (data de lançamento)
            date_lanc_str = cols[0].strip()
            try:
                date_lanc = datetime.strptime(date_lanc_str, '%d/%m/%Y').date()
            except ValueError:
                continue

            # Coluna 1: Valor dos dados (data valor)
            date_valor_str = cols[1].strip() if len(cols) > 1 else ''
            try:
                date_valor = datetime.strptime(date_valor_str, '%d/%m/%Y').date()
            except ValueError:
                date_valor = date_lanc

            # Registrar primeira data (será usada como data do extracto)
            if first_date is None:
                first_date = date_lanc

            # Coluna 2: Descrição
            description = cols[2].strip() if len(cols) > 2 else ''

            # Coluna 3: Montante
            amount_raw = self._parse_bim_prepaid_amount(cols[3]) if len(cols) > 3 else 0.0

            if amount_raw == 0:
                continue

            # INVERTER SINAIS: + passa para - e - passa para +
            # No extracto do cartão pré-pago, a lógica é inversa
            amount = -amount_raw

            # Acumular créditos e débitos
            if amount > 0:
                total_credit += amount
            else:
                total_debit += abs(amount)

            # Gerar unique_import_id
            unique_id = f"BIM-PP-{date_lanc.strftime('%Y%m%d')}-{row_idx}-{abs(amount):.2f}"

            transactions.append({
                'date': date_lanc,
                'date_valor': date_valor,
                'payment_ref': description or 'Sem descrição',
                'amount': amount,
                'debit': abs(amount) if amount < 0 else 0.0,
                'credit': amount if amount > 0 else 0.0,
                'balance': 0.0,
                'num_doc': '',
                'num_oper': '',
                'currency': currency_code,
                'unique_import_id': unique_id,
                'ref': '',
                'partner_name': '',
                'note': '',
                'sequence': len(transactions) + 1,
            })

        _logger.info("BIM Pré-Pago - %d transações encontradas", len(transactions))

        if not transactions:
            raise UserError(_("Nenhuma transação encontrada no ficheiro BIM Pré-Pago."))

        # Gerar nome do statement usando a data da primeira transação
        statement_date = first_date or date.today()
        statement_name = f"BIM Pré-Pago {statement_date.strftime('%Y-%m-%d')}"

        # SALDO INICIAL: buscar do último statement do diário
        last_statement = self.env['account.bank.statement'].search([
            ('journal_id', '=', self.id)
        ], order='date desc, id desc', limit=1)

        if last_statement:
            balance_start = last_statement.balance_end_real
            _logger.info("BIM Pré-Pago - Saldo inicial do último statement: %s", balance_start)
        else:
            balance_start = 0.0
            _logger.info("BIM Pré-Pago - Nenhum statement anterior, saldo inicial = 0.0")

        # SALDO FINAL: Saldo Inicial + Créditos - Débitos
        balance_end = balance_start + total_credit - total_debit
        _logger.info("BIM Pré-Pago - Saldo final: %s (inicial=%s + creditos=%s - debitos=%s)",
                    balance_end, balance_start, total_credit, total_debit)

        vals_bank_statement = [{
            'name': statement_name,
            'date': statement_date,
            'balance_start': balance_start,
            'balance_end_real': balance_end,
            'transactions': transactions,
        }]

        return currency_code, account_number, vals_bank_statement

    # ==========================================
    # PARSER MOZABANCO (CSV)
    # ==========================================

    def _parse_mozabanco_amount(self, value):
        """
        Converte valor do Mozabanco para float.
        Formato: -2 ou 50 ou 35431,06
        Débitos já vêm negativos, créditos positivos.
        """
        if not value:
            return 0.0

        if isinstance(value, (int, float)):
            return float(value)

        value_str = str(value).strip()
        if not value_str:
            return 0.0

        # Remover caracteres não numéricos excepto - , .
        value_str = re.sub(r'[^\d.,-]', '', value_str)

        if not value_str:
            return 0.0

        # Manter sinal
        is_negative = value_str.startswith('-')
        value_str = value_str.lstrip('-')

        # Converter formato europeu: 35.431,06 -> 35431.06
        # Se tem vírgula, assume formato europeu
        if ',' in value_str:
            value_str = value_str.replace('.', '').replace(',', '.')

        try:
            result = float(value_str)
            return -result if is_negative else result
        except ValueError:
            return 0.0

    def _parse_mozabanco_date(self, value):
        """Converte data Mozabanco para date object. Formato: DD/MM/YYYY"""
        if not value:
            return None

        if isinstance(value, datetime):
            return value.date()

        if isinstance(value, date):
            return value

        value_str = str(value).strip()
        if not value_str:
            return None

        # Formato: 27/01/2026
        try:
            return datetime.strptime(value_str, '%d/%m/%Y').date()
        except ValueError:
            pass

        # Formato alternativo: 27-01-2026
        try:
            return datetime.strptime(value_str, '%d-%m-%Y').date()
        except ValueError:
            pass

        return None

    def _parse_bank_statement_file_mozabanco(self, attachment):
        """
        Parseia extracto Mozabanco (CSV).

        Estrutura CSV:
        - Linha 1: sep=;
        - Linha 2: Número da conta (ex: 04381052010001)
        - Linha 3: Saldo de Abertura: ;4653,72
        - Linha 4: Saldo de fecho: ;2894,78
        - Linha 5: Moeda;MZN
        - Linha 6: Headers (Data;Data-Valor;Descrição;Nº de Ref;Débito;Crédito;Saldo)
        - Linha 7+: Transações

        Colunas:
        0: Data (DD/MM/YYYY)
        1: Data-Valor (DD/MM/YYYY)
        2: Descrição
        3: Nº de Ref
        4: Débito (valor negativo)
        5: Crédito (valor positivo)
        6: Saldo
        """
        _logger.info("Iniciando parse Mozabanco: %s", attachment.name)

        # Tentar diferentes encodings
        content = None
        for encoding in ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252']:
            try:
                content = attachment.raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if not content:
            raise UserError(_("Não foi possível ler o ficheiro. Encoding não suportado."))

        lines = content.strip().split('\n')

        if len(lines) < 7:
            raise UserError(_("Ficheiro CSV inválido. Número insuficiente de linhas."))

        # Extrair informações do cabeçalho
        account_number = None
        currency_code = 'MZN'
        balance_start = 0.0
        balance_end = 0.0

        # Linha 2: Número da conta (pode ter caracteres invisíveis)
        account_line = lines[1].replace(';', '').strip()
        # Remover caracteres não numéricos
        account_number = re.sub(r'[^\d]', '', account_line)
        if account_number:
            _logger.info("Mozabanco - Conta: %s", account_number)

        # Linha 3: Saldo de Abertura
        if len(lines) > 2:
            saldo_abertura_line = lines[2]
            parts = saldo_abertura_line.split(';')
            if len(parts) > 1:
                balance_start = self._parse_mozabanco_amount(parts[1])
                _logger.info("Mozabanco - Saldo Abertura: %s", balance_start)

        # Linha 4: Saldo de fecho
        if len(lines) > 3:
            saldo_fecho_line = lines[3]
            parts = saldo_fecho_line.split(';')
            if len(parts) > 1:
                balance_end = self._parse_mozabanco_amount(parts[1])
                _logger.info("Mozabanco - Saldo Fecho: %s", balance_end)

        # Linha 5: Moeda
        if len(lines) > 4:
            moeda_line = lines[4]
            parts = moeda_line.split(';')
            if len(parts) > 1:
                moeda = parts[1].strip().upper()
                if moeda in ['MZN', 'USD', 'EUR']:
                    currency_code = moeda
                _logger.info("Mozabanco - Moeda: %s", currency_code)

        # Linha 6: Headers (índice 5)
        # Linha 7+: Dados (índice 6+)
        data_start_row = 6

        # Parse transações
        transactions = []
        first_date = None
        last_date = None
        last_transaction_saldo = None     # Saldo da última transação
        last_transaction_amount = None    # Amount da última transação

        for row_idx, line in enumerate(lines[data_start_row:], start=data_start_row + 1):
            if not line.strip():
                continue

            cols = line.split(';')
            if len(cols) < 6:
                continue

            # Coluna 0: Data
            date_str = cols[0].strip()
            date_parsed = self._parse_mozabanco_date(date_str)
            if not date_parsed:
                continue

            # Registrar primeira e última data
            if first_date is None:
                first_date = date_parsed
            last_date = date_parsed

            # Coluna 1: Data-Valor
            date_valor_str = cols[1].strip() if len(cols) > 1 else ''
            date_valor = self._parse_mozabanco_date(date_valor_str) or date_parsed

            # Coluna 2: Descrição
            description = cols[2].strip() if len(cols) > 2 else ''

            # Coluna 3: Nº de Ref
            ref = cols[3].strip() if len(cols) > 3 else ''

            # Coluna 4: Débito (valor negativo)
            debit = self._parse_mozabanco_amount(cols[4]) if len(cols) > 4 else 0.0

            # Coluna 5: Crédito (valor positivo)
            credit = self._parse_mozabanco_amount(cols[5]) if len(cols) > 5 else 0.0

            # Coluna 6: Saldo
            saldo = self._parse_mozabanco_amount(cols[6]) if len(cols) > 6 else 0.0

            # Calcular amount (débito já é negativo)
            amount = debit + credit

            # Guardar dados da última transação (para calcular saldo inicial)
            if saldo:
                last_transaction_saldo = saldo
                last_transaction_amount = amount

            if amount == 0:
                continue

            # Gerar unique_import_id
            if ref:
                unique_id = f"MOZABANCO-{account_number}-{ref}-{abs(amount):.2f}"
            else:
                unique_id = f"MOZABANCO-{account_number}-{date_parsed.strftime('%Y%m%d')}-{row_idx}-{abs(amount):.2f}"

            transactions.append({
                'date': date_parsed,
                'date_valor': date_valor,
                'payment_ref': description or 'Sem descrição',
                'amount': amount,
                'debit': abs(debit) if debit < 0 else 0.0,
                'credit': credit if credit > 0 else 0.0,
                'balance': saldo,
                'num_doc': ref,
                'num_oper': '',
                'currency': currency_code,
                'unique_import_id': unique_id,
                'ref': ref,
                'partner_name': '',
                'note': '',
                'sequence': len(transactions) + 1,
            })

        _logger.info("Mozabanco - %d transações encontradas", len(transactions))

        if not transactions:
            raise UserError(_("Nenhuma transação encontrada no ficheiro Mozabanco."))

        # SALDO INICIAL: ordenar transações por (data, num_doc, saldo).
        # Mozabanco lista por Data-Valor desc, com Data podendo ser diferente
        # dentro do mesmo grupo. Para encontrar a op mais antiga:
        # 1) data ascendente
        # 2) num_doc ascendente (sequencial dentro do dia)
        # 3) saldo descendente para débito / ascendente para crédito
        #    (saldo só desce com débito, sobe com crédito)
        def _moza_sort_key(t):
            is_debit = t['amount'] < 0
            return (
                t['date'],
                t.get('num_doc') or '',
                -t['balance'] if is_debit else t['balance'],
            )

        transactions_sorted = sorted(transactions, key=_moza_sort_key)
        first_t = transactions_sorted[0]
        balance_start = round(first_t['balance'] - first_t['amount'], 2)
        _logger.info("Mozabanco - Saldo inicial calculado: %s (op mais antiga: saldo=%s, amount=%s)",
                     balance_start, first_t['balance'], first_t['amount'])

        # Gerar nome do statement
        if first_date:
            statement_name = f"Mozabanco Extracto {first_date.strftime('%Y-%m-%d')}"
        else:
            statement_name = f"Mozabanco Extracto {date.today().strftime('%Y-%m-%d')}"

        vals_bank_statement = [{
            'name': statement_name,
            'date': first_date or date.today(),
            'balance_start': balance_start,
            'balance_end_real': balance_end,
            'transactions': transactions,
        }]

        return currency_code, account_number, vals_bank_statement

    # ==========================================
    # MÉTODO PRINCIPAL DE PARSE
    # ==========================================

    def _parse_bank_statement_file(self, attachment):
        """
        Override do método principal para detectar e parsear
        extractos BIM, BCI, BCI Cartão de Crédito, BIM Pré-Pago e Mozabanco.
        """
        # Tentar BIM primeiro (xlsx)
        if self._check_bim(attachment):
            return self._parse_bank_statement_file_bim(attachment)

        # Tentar BCI Cartão de Crédito (antes do BCI normal)
        if self._check_bci_credit_card(attachment):
            return self._parse_bank_statement_file_bci_credit_card(attachment)

        # Tentar BCI
        if self._check_bci(attachment):
            return self._parse_bank_statement_file_bci(attachment)

        # Tentar Mozabanco (csv)
        if self._check_mozabanco(attachment):
            return self._parse_bank_statement_file_mozabanco(attachment)

        # Tentar BIM Pré-Pago (csv) - deve vir depois do Mozabanco
        if self._check_bim_prepaid_card(attachment):
            return self._parse_bank_statement_file_bim_prepaid_card(attachment)

        # Se não for nenhum formato conhecido, passar para o próximo parser
        return super()._parse_bank_statement_file(attachment)

    # ==========================================
    # PREVIEW DE IMPORTAÇÃO
    # ==========================================

    def _create_bank_statement_preview(self, attachment, currency_code, account_number, statements, bank_name):
        """
        Cria um wizard de preview com os dados parseados.

        Args:
            attachment: ir.attachment do ficheiro
            currency_code: código da moeda (MZN, USD, EUR)
            account_number: número da conta bancária
            statements: lista de statements parseados
            bank_name: nome do banco (BIM, BCI, etc)

        Returns:
            Action para abrir o wizard de preview
        """
        import json

        if not statements:
            raise UserError(_("Nenhuma transação encontrada no ficheiro."))

        stmt = statements[0]  # Primeiro statement
        transactions = stmt.get('transactions', [])

        # Calcular totais
        total_credit = sum(t['amount'] for t in transactions if t['amount'] > 0)
        total_debit = sum(t['amount'] for t in transactions if t['amount'] < 0)

        # Preparar dados para JSON (converter dates para strings)
        import_data = {
            'currency_code': currency_code,
            'account_number': account_number,
            'statements': []
        }

        for s in statements:
            stmt_copy = dict(s)
            if stmt_copy.get('date'):
                stmt_copy['date'] = stmt_copy['date'].strftime('%Y-%m-%d') if hasattr(stmt_copy['date'], 'strftime') else str(stmt_copy['date'])

            trans_copy = []
            for t in stmt_copy.get('transactions', []):
                t_copy = dict(t)
                if t_copy.get('date'):
                    t_copy['date'] = t_copy['date'].strftime('%Y-%m-%d') if hasattr(t_copy['date'], 'strftime') else str(t_copy['date'])
                if t_copy.get('date_valor'):
                    t_copy['date_valor'] = t_copy['date_valor'].strftime('%Y-%m-%d') if hasattr(t_copy['date_valor'], 'strftime') else str(t_copy['date_valor'])
                trans_copy.append(t_copy)
            stmt_copy['transactions'] = trans_copy
            import_data['statements'].append(stmt_copy)

        # Criar preview wizard
        preview = self.env['mh.bank.statement.import.preview'].create({
            'journal_id': self.id,
            'attachment_id': attachment.id,
            'filename': attachment.name,
            'bank_name': bank_name,
            'account_number': account_number or '',
            'currency_code': currency_code,
            'statement_name': stmt.get('name', ''),
            'statement_date': stmt.get('date'),
            'balance_start': stmt.get('balance_start', 0.0),
            'balance_end': stmt.get('balance_end_real', 0.0),
            'total_transactions': len(transactions),
            'total_credit': total_credit,
            'total_debit': abs(total_debit),
            'import_data': json.dumps(import_data),
        })

        # Criar linhas de preview
        for trans in transactions:
            self.env['mh.bank.statement.import.preview.line'].create({
                'preview_id': preview.id,
                'sequence': trans.get('sequence', 0),
                'date': trans.get('date'),
                'date_valor': trans.get('date_valor'),
                'payment_ref': trans.get('payment_ref', ''),
                'debit': trans.get('debit', 0.0),
                'credit': trans.get('credit', 0.0),
                'balance': trans.get('balance', 0.0),
                'num_doc': trans.get('num_doc', ''),
                'num_oper': trans.get('num_oper', ''),
                'currency': trans.get('currency', ''),
                'amount': trans.get('amount', 0.0),
                'partner_name': trans.get('partner_name', ''),
                'ref': trans.get('ref', ''),
                'note': trans.get('note', ''),
            })

        # Retornar action para abrir o wizard
        return {
            'name': _('Preview de Importação - %s') % bank_name,
            'type': 'ir.actions.act_window',
            'res_model': 'mh.bank.statement.import.preview',
            'view_mode': 'form',
            'views': [(False, 'form')],
            'res_id': preview.id,
            'target': 'new',
            'context': dict(self.env.context),
        }

    def create_document_from_attachment(self, attachment_ids=None):
        """
        Override para mostrar preview antes de importar ficheiros BIM/BCI/BCI Cartão.
        """
        journal = self or self.browse(self.env.context.get('default_journal_id'))

        if journal.type in ('bank', 'credit', 'cash'):
            attachments = self.env['ir.attachment'].browse(attachment_ids)
            if not attachments:
                raise UserError(_("No attachment was provided"))

            # Verificar se é BIM, BCI Cartão ou BCI para mostrar preview
            for attachment in attachments:
                if journal._check_bim(attachment):
                    currency_code, account_number, statements = journal._parse_bank_statement_file_bim(attachment)
                    return journal._create_bank_statement_preview(
                        attachment, currency_code, account_number, statements, 'BIM'
                    )

                # BCI Cartão de Crédito (verificar antes do BCI normal)
                if journal._check_bci_credit_card(attachment):
                    currency_code, account_number, statements = journal._parse_bank_statement_file_bci_credit_card(attachment)
                    return journal._create_bank_statement_preview(
                        attachment, currency_code, account_number, statements, 'BCI Cartão Crédito'
                    )

                if journal._check_bci(attachment):
                    currency_code, account_number, statements = journal._parse_bank_statement_file_bci(attachment)
                    return journal._create_bank_statement_preview(
                        attachment, currency_code, account_number, statements, 'BCI'
                    )

                # Mozabanco (csv)
                if journal._check_mozabanco(attachment):
                    currency_code, account_number, statements = journal._parse_bank_statement_file_mozabanco(attachment)
                    return journal._create_bank_statement_preview(
                        attachment, currency_code, account_number, statements, 'Mozabanco'
                    )

                # BIM Pré-Pago (csv) - deve vir depois do Mozabanco
                if journal._check_bim_prepaid_card(attachment):
                    currency_code, account_number, statements = journal._parse_bank_statement_file_bim_prepaid_card(attachment)
                    return journal._create_bank_statement_preview(
                        attachment, currency_code, account_number, statements, 'BIM Pré-Pago'
                    )

        # Se não for BIM, BCI Cartão, BCI, Mozabanco nem BIM Pré-Pago, usar importação padrão
        return super().create_document_from_attachment(attachment_ids)
