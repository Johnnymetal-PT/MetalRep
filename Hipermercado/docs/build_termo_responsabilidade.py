#!/usr/bin/env python3
"""Gera o Termo de Responsabilidade (acesso Odoo.sh) sobre o template Word da marca.

O template `doc_templates/assets/Template_Parametro_Word.docx` traz capa, cabeçalho,
rodapé e estilos da Parâmetro Global. Este script preenche os placeholders da capa,
descarta as páginas de exemplo e reconstrói o corpo com o conteúdo do Termo,
reutilizando os parágrafos-protótipo do template para herdar a formatação.
"""

from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = Path(r"c:\Users\Utilizador\Desktop\Repos\Hipermercado")
TEMPLATE = REPO / "doc_templates" / "assets" / "Template_Parametro_Word.docx"
OUT = REPO / "TERMO_RESPONSABILIDADE_ODOO_SH_v1.0.docx"

DARK = RGBColor(0x12, 0x1A, 0x26)
MUTED = RGBColor(0x5A, 0x69, 0x7D)

CLIENTE = "Prudentealgoritmo, Lda."
DATA = "13 de agosto de 2026"
VERSAO = "1.0"

# Índices no corpo do template (ver inspeção): 0-22 é a capa (22 tem o section break),
# 23.. são páginas de exemplo que servem de protótipo e depois são descartadas.
IDX_H1, IDX_DESC, IDX_EMPTY, IDX_BODY, IDX_PAGEBREAK = 23, 24, 25, 26, 28
FIRST_SAMPLE, LAST_SAMPLE = 23, 65


# --------------------------------------------------------------------------- utils
def substituir_em_runs(paragrafo, antigo: str, novo: str) -> bool:
    """Substitui texto preservando a formatação do primeiro run."""
    if antigo not in paragrafo.text:
        return False
    runs = paragrafo.runs
    if not runs:
        return False
    runs[0].text = paragrafo.text.replace(antigo, novo)
    for r in runs[1:]:
        r.text = ""
    return True


def substituir_no_documento(container, pares: dict[str, str]) -> None:
    for p in container.paragraphs:
        for antigo, novo in pares.items():
            substituir_em_runs(p, antigo, novo)
    for t in getattr(container, "tables", []):
        for row in t.rows:
            for cell in row.cells:
                substituir_no_documento(cell, pares)


def bordo_inferior(cell, cor="5A697D", tamanho="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), tamanho)
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), cor)
    borders.append(bottom)


# ------------------------------------------------------------------- construtores
class Corpo:
    """Acumula elementos e insere-os antes do sectPr final do documento."""

    def __init__(self, doc: Document, protos: dict):
        self._body = doc.element.body
        self._sect = self._body.find(qn("w:sectPr"))
        self._protos = protos
        self._doc = doc

    def _clonar(self, nome: str):
        el = copy.deepcopy(self._protos[nome])
        self._sect.addprevious(el)
        return el

    def _paragrafo(self, nome: str):
        from docx.text.paragraph import Paragraph

        el = self._clonar(nome)
        return Paragraph(el, self._doc._body)

    def _limpar_runs(self, p):
        proto_run = None
        for r in p.runs:
            if proto_run is None:
                proto_run = copy.deepcopy(r._r)
            r._r.getparent().remove(r._r)
        return proto_run

    def titulo_pagina(self, texto: str, descricao: str = "") -> None:
        p = self._paragrafo("h1")
        self._escrever(p, [(texto, False)])
        if descricao:
            d = self._paragrafo("desc")
            self._escrever(d, [(descricao, False)])

    def _escrever(self, p, partes: list[tuple[str, bool]]) -> None:
        """partes: lista de (texto, negrito)."""
        from docx.text.run import Run

        proto_run = self._limpar_runs(p)
        for texto, negrito in partes:
            r_el = copy.deepcopy(proto_run)
            for t in r_el.findall(qn("w:t")):
                r_el.remove(t)
            p._p.append(r_el)
            run = Run(r_el, p)
            run.text = texto
            if negrito:
                run.font.bold = True
        return p

    def clausula(self, numero: str, titulo: str) -> None:
        p = self._paragrafo("body")
        self._escrever(p, [(f"{numero}. {titulo}", True)])
        for r in p.runs:
            r.font.name = "Aptos Display"
            r.font.size = Pt(13)
            r.font.color.rgb = DARK
        pf = p.paragraph_format
        pf.space_before = Pt(14)
        pf.space_after = Pt(4)
        pf.keep_with_next = True

    def texto(self, partes, espaco_antes: int = 0, indent_cm: float = 0.0) -> None:
        if isinstance(partes, str):
            partes = [(partes, False)]
        p = self._paragrafo("body")
        self._escrever(p, partes)
        pf = p.paragraph_format
        pf.space_before = Pt(espaco_antes)
        pf.space_after = Pt(6)
        if indent_cm:
            pf.left_indent = Cm(indent_cm)
            pf.first_line_indent = Cm(-0.7)

    def alinea(self, marcador: str, texto: str) -> None:
        self.texto([(f"{marcador}\t{texto}", False)], indent_cm=0.8)

    def espaco(self) -> None:
        self._clonar("empty")

    def salto_pagina(self) -> None:
        self._clonar("pagebreak")

    def tabela_assinatura(self, entidade: str) -> None:
        self.texto([(entidade, True)], espaco_antes=10)
        tabela = self._doc.add_table(rows=4, cols=2)
        tabela.autofit = False
        for i, campo in enumerate(("Nome", "Cargo", "Data", "Assinatura")):
            etiqueta, valor = tabela.rows[i].cells
            etiqueta.width = Cm(2.6)
            valor.width = Cm(11.0)
            par = etiqueta.paragraphs[0]
            run = par.add_run(f"{campo}:")
            run.font.name = "Aptos"
            run.font.size = Pt(10.5)
            run.font.color.rgb = MUTED
            valor.paragraphs[0].add_run(" ")
            bordo_inferior(valor)
            for cell in (etiqueta, valor):
                cell.paragraphs[0].paragraph_format.space_after = Pt(10)
        self._sect.addprevious(tabela._tbl)
        self.espaco()


def main() -> int:
    if not TEMPLATE.exists():
        print(f"ERRO: template não encontrado em {TEMPLATE}", file=sys.stderr)
        return 1

    doc = Document(str(TEMPLATE))
    body = doc.element.body
    elementos = list(body)

    protos = {
        "h1": copy.deepcopy(elementos[IDX_H1]),
        "desc": copy.deepcopy(elementos[IDX_DESC]),
        "empty": copy.deepcopy(elementos[IDX_EMPTY]),
        "body": copy.deepcopy(elementos[IDX_BODY]),
        "pagebreak": copy.deepcopy(elementos[IDX_PAGEBREAK]),
    }

    # Capa: preencher placeholders antes de mexer no resto do corpo.
    substituir_no_documento(
        doc,
        {
            "[Título do Documento]": "Termo de Responsabilidade",
            "[Subtítulo opcional]": "Acesso à Plataforma Odoo.sh",
            "[Descrição curta do documento.]": (
                "Regula a concessão de acesso à plataforma Odoo.sh associada ao projeto "
                "do Cliente, as responsabilidades de utilização do ambiente e a "
                "propriedade intelectual dos desenvolvimentos realizados pela Parâmetro."
            ),
            "[Nome]": f"{CLIENTE} | Projeto Odoo",
            "[DD de Mês de AAAA]": DATA,
            "[X.X]": VERSAO,
        },
    )

    # Descartar as páginas de exemplo do template (mantendo a capa e o sectPr final).
    for el in elementos[FIRST_SAMPLE : LAST_SAMPLE + 1]:
        body.remove(el)

    # Rodapés das duas secções.
    for sec in doc.sections:
        substituir_no_documento(
            sec.footer, {"[Documento] | [Ano]": "Termo de Responsabilidade | Odoo.sh | 2026"}
        )

    c = Corpo(doc, protos)

    c.titulo_pagina("Termo de Responsabilidade", "Acesso à Plataforma Odoo.sh")

    c.texto("Entre:", espaco_antes=2)
    c.texto([("Parâmetro Global, Lda.", True), (", adiante designada por Parâmetro,", False)])
    c.texto("e")
    c.texto([(CLIENTE, True), (", adiante designada por Cliente.", False)])

    c.clausula("1", "Objeto")
    c.texto(
        "O presente Termo de Responsabilidade regula a concessão de acesso à plataforma "
        "Odoo.sh, associada ao projeto do Cliente, permitindo a gestão técnica, "
        "manutenção e desenvolvimento do respetivo ambiente."
    )

    c.clausula("2", "Propriedade da Base de Dados")
    c.texto(
        "A base de dados do sistema é propriedade exclusiva do Cliente, podendo este "
        "aceder, gerir, alterar ou utilizar a mesma conforme entender, no âmbito da sua "
        "atividade e de acordo com as condições da plataforma Odoo.sh."
    )

    c.clausula("3", "Responsabilidade de Utilização")
    c.texto(
        "O Cliente compromete-se a utilizar o acesso concedido de forma responsável, "
        "assegurando que:"
    )
    c.alinea("a)", "as credenciais de acesso são utilizadas apenas por pessoas autorizadas;")
    c.alinea(
        "b)",
        "não são realizadas ações que comprometam a segurança, integridade ou "
        "funcionamento do sistema;",
    )
    c.alinea(
        "c)",
        "qualquer alteração técnica efetuada no ambiente Odoo.sh é realizada sob sua "
        "responsabilidade.",
    )
    c.texto(
        "A Parâmetro não se responsabiliza por eventuais impactos decorrentes de "
        "alterações realizadas diretamente pelo Cliente ou por terceiros autorizados "
        "por este.",
        espaco_antes=4,
    )

    c.clausula("4", "Propriedade Intelectual")
    c.texto(
        "Todo o código, módulos, desenvolvimentos, personalizações e respetiva "
        "documentação técnica desenvolvidos pela Parâmetro no âmbito do projeto "
        "constituem propriedade intelectual da Parâmetro Global, Lda."
    )

    c.clausula("5", "Direito de Utilização")
    c.texto(
        "Ao Cliente é concedido um direito de utilização do código desenvolvido pela "
        "Parâmetro, exclusivamente para efeitos de funcionamento e utilização do seu "
        "sistema."
    )
    c.texto(
        "Este direito não inclui o direito de exploração comercial, revenda, "
        "distribuição, sublicenciamento ou reutilização do referido código em projetos "
        "de terceiros sem autorização prévia e por escrito da Parâmetro."
    )

    c.clausula("6", "Limitação de Responsabilidade")
    c.texto(
        "A Parâmetro não se responsabiliza por quaisquer danos, perdas de dados, falhas "
        "operacionais ou impactos técnicos decorrentes de alterações realizadas "
        "diretamente pelo Cliente ou por terceiros no ambiente Odoo.sh."
    )

    c.clausula("7", "Aceitação")
    c.texto(
        "Com a assinatura do presente documento, o Cliente declara ter tomado "
        "conhecimento e aceitar as condições estabelecidas no presente Termo de "
        "Responsabilidade."
    )

    c.espaco()
    c.tabela_assinatura("Pela Parâmetro Global, Lda.")
    c.tabela_assinatura(f"Pela {CLIENTE}")

    doc.save(str(OUT))
    print(f"OK: {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
