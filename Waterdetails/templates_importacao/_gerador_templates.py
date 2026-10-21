# -*- coding: utf-8 -*-
"""Gera os templates Excel de importacao para o cliente (Odoo 19.0 Enterprise).
Modelos alvo: res.partner (Contactos), product.template (Produtos),
product.pricelist / product.pricelist.item (Listas de Precos por marca).
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = os.path.dirname(os.path.abspath(__file__))

# ---- estilos ----
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
REQ_FILL = PatternFill("solid", fgColor="C00000")  # obrigatorio
OPT_FILL = PatternFill("solid", fgColor="1F4E78")
EX_FONT = Font(italic=True, color="808080")
EX_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
SUB_FONT = Font(bold=True, size=11, color="1F4E78")
NOTE_FONT = Font(size=10, color="404040")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_headers(ws, headers, row=1):
    for c, (label, required, width, comment) in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = HDR_FONT
        cell.fill = REQ_FILL if required else OPT_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = width
        if comment:
            from openpyxl.comments import Comment
            cell.comment = Comment(comment, "Parametro Global")
    ws.row_dimensions[row].height = 34
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def write_examples(ws, rows, start_row=2):
    for r, data in enumerate(rows, start=start_row):
        for c, val in enumerate(data, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = EX_FONT
            cell.fill = EX_FILL
            cell.border = BORDER
            cell.alignment = WRAP


def add_dropdown(ws, col_letter, options, first=2, last=2000):
    formula = '"%s"' % ",".join(options)
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    dv.error = "Valor invalido. Escolha uma das opcoes da lista."
    dv.errorTitle = "Opcao invalida"
    ws.add_data_validation(dv)
    dv.add("%s%d:%s%d" % (col_letter, first, col_letter, last))


def instr_sheet(wb, title, lines):
    ws = wb.create_sheet("Instrucoes", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    ws.cell(row=2, column=2, value=title).font = TITLE_FONT
    r = 4
    for kind, text in lines:
        cell = ws.cell(row=r, column=2, value=text)
        if kind == "h":
            cell.font = SUB_FONT
        else:
            cell.font = NOTE_FONT
        cell.alignment = WRAP
        r += 1
    return ws


COMMON_INTRO = [
    ("h", "Como preencher este ficheiro"),
    ("n", "1. Preencha uma linha por registo, a partir da linha abaixo dos titulos."),
    ("n", "2. As linhas a cinzento sao EXEMPLOS — apague-as antes de enviar."),
    ("n", "3. Colunas com cabecalho a VERMELHO sao de preenchimento obrigatorio."),
    ("n", "4. Nao altere, renomeie nem reordene as colunas de titulo (linha 1)."),
    ("n", "5. Passe o rato sobre cada titulo para ver a explicacao (comentario)."),
    ("n", "6. Datas no formato AAAA-MM-DD. Numeros sem simbolo de moeda (use ponto decimal)."),
    ("n", "7. Em caso de duvida, deixe a celula vazia em vez de inventar um valor."),
    ("n", ""),
]


# =====================================================================
# 1. CONTACTOS
# =====================================================================
def build_contactos():
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    headers = [
        ("Tipo de Contacto *", True, 20, "Cliente | Fornecedor | Cliente e Fornecedor"),
        ("Tipo de Entidade *", True, 18, "Empresa (pessoa coletiva) ou Individuo (pessoa singular)"),
        ("Nome *", True, 34, "Nome completo da empresa ou da pessoa. Campo obrigatorio."),
        ("Empresa-Mae", False, 28, "Preencher apenas se este contacto for uma pessoa dentro de uma empresa ja listada (ex.: o comprador da empresa X). Indicar o nome exato da empresa."),
        ("NIF / Nº Contribuinte", False, 18, "Numero de contribuinte. Para PT incluir prefixo do pais nao e necessario; o NIF e validado contra a localizacao."),
        ("Morada (Rua)", False, 30, "Rua e numero."),
        ("Morada (Linha 2)", False, 24, "Andar, lote, complemento (opcional)."),
        ("Codigo Postal", False, 14, "Formato 0000-000 para PT."),
        ("Localidade", False, 20, "Cidade / localidade."),
        ("Distrito", False, 18, "Distrito / regiao (ex.: Porto, Lisboa)."),
        ("Pais", False, 16, "Por defeito Portugal. Indicar o pais para contactos estrangeiros."),
        ("Telefone", False, 18, "Telefone fixo."),
        ("Telemovel", False, 18, "Telemovel."),
        ("Email", False, 28, "Email principal. Usado para envio de documentos."),
        ("Website", False, 22, "Sitio web (opcional)."),
        ("Funcao / Cargo", False, 20, "Cargo da pessoa (apenas se Tipo de Entidade = Individuo)."),
        ("Idioma", False, 18, "Idioma de comunicacao. Ex.: Portugues (PT)."),
        ("IBAN", False, 28, "Conta bancaria (opcional, util para pagamentos/recebimentos)."),
        ("Etiquetas", False, 22, "Classificacoes livres separadas por virgula (ex.: Revendedor, VIP)."),
        ("Notas", False, 30, "Observacoes internas (opcional)."),
    ]
    style_headers(ws, headers)
    examples = [
        ["Cliente", "Empresa", "Exemplo Cliente, Lda", "", "501234567", "Rua das Flores 12",
         "", "4000-001", "Porto", "Porto", "Portugal", "220000000", "", "geral@exemplocliente.pt",
         "www.exemplocliente.pt", "", "Portugues (PT)", "PT50000000000000000000000", "Revendedor", ""],
        ["Fornecedor", "Empresa", "Fornecedor Demo, S.A.", "", "509876543", "Av. Central 100",
         "", "1000-100", "Lisboa", "Lisboa", "Portugal", "210000000", "", "compras@fornecedordemo.pt",
         "", "", "Portugues (PT)", "", "", ""],
        ["Cliente e Fornecedor", "Individuo", "Maria Silva", "Exemplo Cliente, Lda", "", "",
         "", "", "Porto", "", "Portugal", "", "930000000", "maria.silva@exemplocliente.pt",
         "", "Diretora de Compras", "Portugues (PT)", "", "", "Contacto principal"],
    ]
    write_examples(ws, examples)
    add_dropdown(ws, "A", ["Cliente", "Fornecedor", "Cliente e Fornecedor"])
    add_dropdown(ws, "B", ["Empresa", "Individuo"])

    instr_sheet(wb, "Template de CONTACTOS (Clientes e Fornecedores) — Odoo 19", COMMON_INTRO + [
        ("h", "Notas especificas dos Contactos"),
        ("n", "- 'Tipo de Contacto': escolha se a entidade e Cliente, Fornecedor, ou ambos."),
        ("n", "- 'Tipo de Entidade': Empresa para pessoas coletivas; Individuo para pessoas singulares."),
        ("n", "- 'Empresa-Mae': use para ligar pessoas (compradores, contactos) a uma empresa ja listada. O nome tem de coincidir exatamente com a linha da empresa."),
        ("n", "- O NIF e validado pela localizacao portuguesa; confirme que esta correto."),
        ("n", "- Pode juntar clientes e fornecedores no mesmo ficheiro; a coluna 'Tipo de Contacto' distingue-os."),
    ])
    path = os.path.join(OUT, "01_Contactos_Template.xlsx")
    wb.save(path)
    return path


# =====================================================================
# 2. PRODUTOS
# =====================================================================
def build_produtos():
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    headers = [
        ("Referencia Interna / SKU *", True, 22, "Codigo unico do produto (default_code). Nao repetir entre produtos."),
        ("Nome do Produto *", True, 34, "Designacao comercial do produto."),
        ("Tipo *", True, 14, "Bem (produto fisico) | Servico | Combo. A maioria dos artigos de catalogo e 'Bem'."),
        ("Controla Inventario?", False, 16, "Sim/Nao. So aplicavel a 'Bem'. 'Sim' = gere stock/quantidades."),
        ("Marca", False, 18, "Marca do produto. Sera mapeada para Categoria de Produto ou Etiqueta no Odoo (ver Instrucoes)."),
        ("Categoria de Produto", False, 22, "Categoria interna (ex.: Lavagem > Ceras). Pode usar / para hierarquia."),
        ("Codigo de Barras", False, 18, "EAN/Barcode (opcional)."),
        ("Unidade de Medida", False, 16, "Ex.: Unidades, kg, L. Por defeito Unidades."),
        ("Preco de Venda *", True, 14, "Preco de venda base sem IVA (list_price)."),
        ("Custo", False, 12, "Custo do produto (standard_price)."),
        ("IVA de Venda", False, 16, "Imposto de venda. Ex.: IVA 23%. Tem de coincidir com um imposto configurado no Odoo."),
        ("IVA de Compra", False, 16, "Imposto de compra. Ex.: IVA 23%."),
        ("Pode Vender?", False, 12, "Sim/Nao. Disponivel para vendas."),
        ("Pode Comprar?", False, 13, "Sim/Nao. Disponivel para compras."),
        ("Fornecedor Principal", False, 24, "Nome do fornecedor (deve existir tambem no ficheiro de Contactos)."),
        ("Referencia do Fornecedor", False, 20, "Codigo do produto no catalogo do fornecedor (opcional)."),
        ("Preco de Compra", False, 14, "Preco de compra ao fornecedor principal (opcional)."),
        ("Descricao para Venda", False, 30, "Texto que aparece nos documentos de venda (opcional)."),
        ("Peso (kg)", False, 12, "Peso unitario (opcional)."),
        ("Notas Internas", False, 28, "Observacoes (opcional)."),
    ]
    style_headers(ws, headers)
    examples = [
        ["CER-001", "Cera Liquida Premium 500ml", "Bem", "Sim", "MarcaX", "Lavagem / Ceras",
         "5601234567890", "Unidades", "24.90", "12.50", "IVA 23%", "IVA 23%", "Sim", "Sim",
         "Fornecedor Demo, S.A.", "FX-500", "12.50", "Cera de alta durabilidade 500ml", "0.6", ""],
        ["SHA-010", "Shampoo Neutro 1L", "Bem", "Sim", "MarcaY", "Lavagem / Shampoos",
         "5609876543210", "Unidades", "9.90", "4.20", "IVA 23%", "IVA 23%", "Sim", "Sim",
         "Fornecedor Demo, S.A.", "SY-1L", "4.20", "", "1.0", ""],
        ["SRV-001", "Servico de Detalhe Completo", "Servico", "Nao", "", "Servicos",
         "", "Unidades", "120.00", "", "IVA 23%", "", "Sim", "Nao", "", "", "", "Detalhe interior e exterior", "", ""],
    ]
    write_examples(ws, examples)
    add_dropdown(ws, "C", ["Bem", "Servico", "Combo"])
    add_dropdown(ws, "D", ["Sim", "Nao"])
    add_dropdown(ws, "M", ["Sim", "Nao"])
    add_dropdown(ws, "N", ["Sim", "Nao"])

    # folha de referencia de impostos PT
    ref = wb.create_sheet("Valores Permitidos")
    ref.sheet_view.showGridLines = False
    ref.column_dimensions["A"].width = 28
    ref.column_dimensions["B"].width = 60
    ref.cell(row=1, column=1, value="IVA (exemplos PT)").font = SUB_FONT
    ref.cell(row=1, column=2, value="Notas").font = SUB_FONT
    taxes = [
        ("IVA 23%", "Taxa normal (Continente)"),
        ("IVA 13%", "Taxa intermedia"),
        ("IVA 6%", "Taxa reduzida"),
        ("IVA Isento", "Operacoes isentas (indicar artigo aplicavel na configuracao)"),
    ]
    for i, (t, n) in enumerate(taxes, start=2):
        ref.cell(row=i, column=1, value=t).font = NOTE_FONT
        ref.cell(row=i, column=2, value=n).font = NOTE_FONT
    ref.cell(row=7, column=1, value="ATENCAO: os nomes dos impostos devem coincidir EXATAMENTE com os impostos configurados na base Odoo do cliente.").font = Font(italic=True, color="C00000", size=10)

    instr_sheet(wb, "Template de PRODUTOS — Odoo 19", COMMON_INTRO + [
        ("h", "Notas especificas dos Produtos"),
        ("n", "- 'Referencia Interna / SKU' deve ser unica. E o identificador usado para ligar produtos as Listas de Precos."),
        ("n", "- 'Tipo': use 'Bem' para artigos fisicos, 'Servico' para servicos (sem stock)."),
        ("n", "- 'Controla Inventario?' = Sim apenas se quiser gerir quantidades em armazem desse produto."),
        ("n", "- IMPORTANTE sobre 'Marca': o Odoo 19 standard NAO tem um campo 'Marca' dedicado."),
        ("n", "   -> Recomendacao: a Marca sera mapeada para 'Categoria de Produto' (permite Listas de Precos por marca com uma so regra por marca)."),
        ("n", "   -> Em alternativa pode ser mapeada para 'Etiquetas de Produto'. A decisao final e da equipa Parametro na configuracao."),
        ("n", "   -> Preencha sempre a coluna 'Marca'; nos tratamos do mapeamento tecnico."),
        ("n", "- 'IVA de Venda'/'IVA de Compra': use os nomes exatos dos impostos do Odoo (ver folha 'Valores Permitidos')."),
        ("n", "- 'Fornecedor Principal' deve existir tambem no ficheiro de Contactos (como Fornecedor)."),
    ])
    path = os.path.join(OUT, "02_Produtos_Template.xlsx")
    wb.save(path)
    return path


# =====================================================================
# 3. LISTAS DE PRECOS POR MARCA
# =====================================================================
def build_precos():
    wb = Workbook()
    # Folha A: listas
    wsa = wb.active
    wsa.title = "Listas de Precos"
    headers_a = [
        ("Nome da Lista de Precos *", True, 30, "Nome unico da lista (ex.: 'Precos MarcaX', 'Revenda MarcaY')."),
        ("Marca associada", False, 20, "Marca a que esta lista diz respeito (opcional, organizativo)."),
        ("Moeda *", True, 12, "Por defeito EUR."),
        ("Aplicavel a", False, 26, "Todos os Clientes | Cliente especifico | Grupo de clientes."),
        ("Cliente especifico", False, 26, "Se 'Aplicavel a' = Cliente especifico, indicar o nome do cliente (do ficheiro de Contactos)."),
        ("Notas", False, 28, "Observacoes (opcional)."),
    ]
    style_headers(wsa, headers_a)
    ex_a = [
        ["Precos MarcaX", "MarcaX", "EUR", "Todos os Clientes", "", "Tabela base da MarcaX"],
        ["Revenda MarcaY", "MarcaY", "EUR", "Grupo de clientes", "", "Para revendedores"],
    ]
    write_examples(wsa, ex_a)
    add_dropdown(wsa, "C", ["EUR"])
    add_dropdown(wsa, "D", ["Todos os Clientes", "Cliente especifico", "Grupo de clientes"])

    # Folha B: regras de preco
    wsb = wb.create_sheet("Regras de Preco")
    headers_b = [
        ("Lista de Precos *", True, 28, "Nome da lista — tem de coincidir com a folha 'Listas de Precos'."),
        ("Aplica-se a *", True, 22, "Produto | Categoria de Produto | Todos os Produtos."),
        ("Referencia do Produto (SKU)", False, 24, "Preencher se 'Aplica-se a' = Produto. SKU do ficheiro de Produtos."),
        ("Categoria de Produto", False, 22, "Preencher se 'Aplica-se a' = Categoria de Produto (ex.: a marca)."),
        ("Quantidade Minima", False, 16, "Quantidade a partir da qual a regra aplica (por defeito 1)."),
        ("Tipo de Calculo *", True, 18, "Preco Fixo | Desconto (%) | Formula."),
        ("Preco Fixo", False, 14, "Preencher se 'Tipo de Calculo' = Preco Fixo. Valor sem IVA."),
        ("Desconto (%)", False, 14, "Preencher se 'Tipo de Calculo' = Desconto (%). Ex.: 10 = 10% de desconto."),
        ("Data Inicio", False, 14, "Inicio de validade (AAAA-MM-DD), opcional."),
        ("Data Fim", False, 14, "Fim de validade (AAAA-MM-DD), opcional."),
    ]
    style_headers(wsb, headers_b)
    ex_b = [
        ["Precos MarcaX", "Categoria de Produto", "", "MarcaX", "1", "Desconto (%)", "", "10", "", ""],
        ["Precos MarcaX", "Produto", "CER-001", "", "1", "Preco Fixo", "21.90", "", "", ""],
        ["Revenda MarcaY", "Produto", "SHA-010", "6", "Desconto (%)", "Desconto (%)", "", "15", "2026-01-01", "2026-12-31"],
    ]
    # corrigir 3a linha (10 colunas)
    ex_b[2] = ["Revenda MarcaY", "Produto", "SHA-010", "", "6", "Desconto (%)", "", "15", "2026-01-01", "2026-12-31"]
    write_examples(wsb, ex_b)
    add_dropdown(wsb, "B", ["Produto", "Categoria de Produto", "Todos os Produtos"])
    add_dropdown(wsb, "F", ["Preco Fixo", "Desconto (%)", "Formula"])

    instr_sheet(wb, "Template de LISTAS DE PRECOS (por marca) — Odoo 19", COMMON_INTRO + [
        ("h", "Como funciona"),
        ("n", "- Folha 'Listas de Precos': defina aqui CADA lista (tipicamente uma por marca)."),
        ("n", "- Folha 'Regras de Preco': defina as linhas de preco/desconto de cada lista."),
        ("n", "- A coluna 'Lista de Precos' (folha 2) tem de coincidir com o nome da folha 1."),
        ("n", ""),
        ("h", "Listas de Precos por marca — recomendacao"),
        ("n", "- Se a Marca estiver mapeada para Categoria de Produto, pode criar UMA regra por marca:"),
        ("n", "   'Aplica-se a' = Categoria de Produto, indicando a marca, com Preco Fixo ou Desconto (%)."),
        ("n", "- Para excecoes a produtos especificos, acrescente regras 'Aplica-se a' = Produto com o SKU."),
        ("n", "- 'Tipo de Calculo': Preco Fixo (preco final) ou Desconto (%) sobre o preco de venda base do produto."),
    ])
    path = os.path.join(OUT, "03_Listas_Precos_Template.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    for fn in (build_contactos, build_produtos, build_precos):
        print("Gerado:", fn())
