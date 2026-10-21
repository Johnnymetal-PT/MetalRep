"""
Actualiza ORCAMENTO_WATERDETAILS_V1.xlsx para a nova estrutura:
  Fase 1 — Faturação:     80h / 10 dias / 5 400 €
  Fase 2 — Operação:      64h /  8 dias / 4 360 € (indicativo)
  Fase 3 — Expansão:      72h /  9 dias / 4 680 € [POR VALIDAR]
  Total 1.º ano s/IVA:    7 752 €
  Recorrência anual:      5 340 €
"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

XLSX = "c:/Projetos/Waterdetails/docs/ORCAMENTO_WATERDETAILS_V1.xlsx"
wb = load_workbook(XLSX)

def w(ws, row, col, val):
    ws.cell(row=row, column=col).value = val

def clr(ws, row, max_col=11):
    for c in range(1, max_col + 1):
        ws.cell(row=row, column=c).value = None


# ===========================================================================
# RESUMO
# ===========================================================================
ws = wb["Resumo"]

# Desunir A25:C25 (nota) para poder usar as células B25/C25
ws.unmerge_cells("A25:C25")

# R2 — data (merged A2:C2 — escrever só em A2)
w(ws, 2, 1, "Parâmetro Global · Consultoria Informática · 25/05/2026")

# R4 — título secção 1 (merged A4:C4 — só A4)
w(ws, 4, 1, "1. IMPLEMENTAÇÃO — Custo Único (Fase 1 — Faturação)")

# R6 — serviços Fase 1
w(ws, 6, 1, "Serviços — Fase 1 (Faturação)")
w(ws, 6, 2, 5400)
w(ws, 6, 3, "Detalhe em sheet Serviços — 80h / 10 dias / ~6 semanas")

# R7 — licenciamento (valor mantém-se 1824, actualizar nota)
w(ws, 7, 3, "⚠ Placeholder — confirmar c/ Stefano Biglieri (Odoo Enterprise 5 users, Odoo SH)")

# R8 — plugin AT (valor mantém-se 528, actualizar nota)
w(ws, 8, 3, "⚠ Placeholder — confirmar c/ Stefano Biglieri (plugin AT certificação faturas)")

# R9 — total 1.º ano s/IVA: 5400+1824+528 = 7752
w(ws, 9, 2, 7752)

# R10 — IVA 23%
w(ws, 10, 2, round(7752 * 0.23, 2))

# R11 — total c/IVA
w(ws, 11, 2, round(7752 * 1.23, 2))

# R13 — título secção 2 (merged A13:C13 — só A13)
# Recorrência: manter valores correctos (2280, 660, 2400, 5340)

# R20 — título secção 3 (merged A20:C20 — só A20)
w(ws, 20, 1, "3. FASES SEGUINTES (Indicativo — a confirmar após go-live Fase 1)")

# R22 — Fase 2 serviços
w(ws, 22, 1, "Serviços — Fase 2: Operação Comercial (indicativo)")
w(ws, 22, 2, 4360)
w(ws, 22, 3, "CRM, Pipedrive [POR VALIDAR], Compras, Inventário — ver sheet Serviços")

# R23 — subtotal Fase 2
w(ws, 23, 1, "SUBTOTAL FASE 2 (sem IVA)")
w(ws, 23, 2, 4360)
w(ws, 23, 3, None)

# R24 — Fase 3 serviços (era linha vazia)
w(ws, 24, 1, "Serviços — Fase 3: Expansão (indicativo) [POR VALIDAR]")
w(ws, 24, 2, 4680)
w(ws, 24, 3, "BoM, Multiempresa (Starplast SRLE), Analítica, Módulos — ver sheet Serviços")

# R25 — subtotal Fase 3 (era nota — unmerged acima)
w(ws, 25, 1, "SUBTOTAL FASE 3 (sem IVA)")
w(ws, 25, 2, 4680)
w(ws, 25, 3, None)

# R26 — total fases seguintes
w(ws, 26, 1, "TOTAL FASES SEGUINTES (Fase 2 + Fase 3, indicativo)")
w(ws, 26, 2, 9040)
w(ws, 26, 3, None)

# R27 — nota
w(ws, 27, 1,
  "Nota: Licenciamento Odoo e plugin AT — valores placeholder, confirmar c/ Stefano Biglieri. "
  "Fase 2 e Fase 3 sujeitas a levantamento de requisitos adicional. "
  "Cada fase tem orçamento e proposta própria após go-live da anterior. "
  "Proposta válida até 20/06/2026.")

# Limpar R28 (por precaução)
clr(ws, 28, max_col=3)


# ===========================================================================
# SERVIÇOS
# ===========================================================================
ws = wb["Serviços"]

# R2 — título (merged A2:E2 — só A2)
w(ws, 2, 1, "FASE 1 — Faturação (Virgin Details, Lda)")

# Novos dados Fase 1: R4–R10 (7 tarefas, substituindo as antigas R4–R15)
fase1 = [
    ("01.Gestão de Projeto",          "Kick-off, acompanhamento quinzenal, reporting, encerramento de fase",           8,   1,    720),
    ("02.Kick-off + Configuração",    "Setup empresa, 5 utilizadores, grupos de acesso, e-mail e templates",          12,   1.5,  780),
    ("03.Contactos e Produtos",       "Clientes, fornecedores, produtos e serviços; importação de PHC GO e Excel",     8,   1,    520),
    ("04.Vendas e Faturação",         "Orçamentos, encomendas, faturas, notas de crédito, faturação certificada AT",  16,   2,   1040),
    ("05.Financeiro l10n_pt",         "Plano contas PT, IVA, SAF-T, plugin AT, condições pagamento, banca",           16,   2,   1040),
    ("06.Migração PHC GO",            "Importação produtos, clientes/fornecedores, saldos em aberto de PHC GO",       16,   2,   1040),
    ("07.Testes + Formação + Arranque","Testes faturação end-to-end, formação 5 utilizadores, suporte go-live",        4,   0.5,  260),
]
for i, row in enumerate(fase1, start=4):
    for j, v in enumerate(row, start=1):
        w(ws, i, j, v)

# Limpar antigas R11–R15 (tarefas que já não existem na Fase 1)
for r in range(11, 16):
    clr(ws, r, max_col=5)

# R16 — subtotal Fase 1
w(ws, 16, 1, "SUBTOTAL FASE 1")
w(ws, 16, 2, None)
w(ws, 16, 3, 80)
w(ws, 16, 4, 10)
w(ws, 16, 5, 5400)

# R18 — título Fase 2 (merged A18:E18 — só A18)
w(ws, 18, 1, "FASE 2 — Operação Comercial (Indicativo — após go-live Fase 1)")

# R19 — headers (manter)

# Novos dados Fase 2: R20–R25 (6 tarefas, substituindo antigas R20–R23)
fase2 = [
    ("08.CRM",                   "Pipeline, fases, probabilidades, templates e-mail, atividades, dashboards",         8,  1,   520),
    ("09.Migração Pipedrive",    "Exportação, mapeamento e importação CRM Pipedrive → Odoo [POR VALIDAR]",           16,  2,  1040),
    ("10.Compras",               "Fornecedores europeus, preços custo, lead times, regras MTO",                       8,  1,   520),
    ("11.Inventário e Logística","Armazém Maia, rotas MTO, receção Itália, dropshipping Espanha",                    16,  2,  1040),
    ("12.Testes + Formação F2",  "Testes end-to-end Fase 2, formação, suporte go-live",                              8,  1,   520),
    ("13.Gestão de Projeto F2",  "Acompanhamento Fase 2 (taxa GP)",                                                   8,  1,   720),
]
for i, row in enumerate(fase2, start=20):
    for j, v in enumerate(row, start=1):
        w(ws, i, j, v)

# R26 — subtotal Fase 2 (era TOTAL GERAL SERVIÇOS)
w(ws, 26, 1, "SUBTOTAL FASE 2 (base, indicativo)")
w(ws, 26, 2, None)
w(ws, 26, 3, 64)
w(ws, 26, 4, 8)
w(ws, 26, 5, 4360)

# R27 — linha vazia antes Fase 3
clr(ws, 27, max_col=5)

# R28 — título Fase 3 (nova linha, sem merge — nota para usar)
w(ws, 28, 1, "FASE 3 — Expansão (Indicativo [POR VALIDAR] — após go-live Fase 2)")
for c in range(2, 6):
    w(ws, 28, c, None)

# R29 — headers Fase 3
w(ws, 29, 1, "Fase")
w(ws, 29, 2, "Descrição")
w(ws, 29, 3, "Horas")
w(ws, 29, 4, "Dias")
w(ws, 29, 5, "Valor (€)")

# R30–R33 — tarefas Fase 3
fase3 = [
    ("14.Manufatura (BoM)",             "Listas de Materiais, Ordens de Fabrico, controlo componentes",                16, 2, 1040),
    ("15.Multiempresa — Starplast SRLE","Configuração 2.ª empresa, utilizadores partilhados, intercompany",            24, 3, 1560),
    ("16.Contabilidade Analítica",      "Analítica por projeto / marca (4 marcas) / cliente",                          16, 2, 1040),
    ("17.Módulos Adicionais",           "Sign, Barcode, Documents, Helpdesk — prioridade a definir [POR VALIDAR]",     16, 2, 1040),
]
for i, row in enumerate(fase3, start=30):
    for j, v in enumerate(row, start=1):
        w(ws, i, j, v)

# R34 — subtotal Fase 3
w(ws, 34, 1, "SUBTOTAL FASE 3 (base, indicativo)")
w(ws, 34, 2, None)
w(ws, 34, 3, 72)
w(ws, 34, 4, 9)
w(ws, 34, 5, 4680)

# R35 — linha vazia
clr(ws, 35, max_col=5)

# R36 — total geral
w(ws, 36, 1, "TOTAL GERAL SERVIÇOS")
w(ws, 36, 2, None)
w(ws, 36, 3, 216)
w(ws, 36, 4, 27)
w(ws, 36, 5, 14440)


# ===========================================================================
# CRONOGRAMA
# ===========================================================================
ws = wb["Cronograma"]

# FASE 1 — reescrever R3–R15
crono_f1 = [
    ("01.Gestão de Projeto",          "Kick-off e acompanhamento",          1,   1,  6, "Transversal à Fase 1"),
    ("02.Kick-off + Configuração",    "Setup empresa e utilizadores",        1.5, 1,  1, None),
    ("03.Contactos e Produtos",       "Contactos, produtos, importação",     1,   2,  2, None),
    ("04.Vendas e Faturação",         "Orçamentos, faturas, AT",             2,   2,  3, None),
    ("05.Financeiro l10n_pt",         "Plano contas, SAF-T, banca",          2,   3,  4, None),
    ("06.Migração PHC GO",            "Produtos, clientes, saldos",          2,   4,  5, "Cliente fornece dados em Excel limpo"),
    ("07.Testes + Formação + Arranque","Testes, formação, go-live",          0.5, 5,  6, None),
]
for i, row in enumerate(crono_f1, start=3):
    for j, v in enumerate(row, start=1):
        w(ws, i, j, v)

# Limpar R10–R14 (tarefas antigas)
for r in range(10, 15):
    clr(ws, r, max_col=6)

# R15 — FASE 1 TOTAL
w(ws, 15, 1, "FASE 1 TOTAL")
w(ws, 15, 3, 10)
for c in [2, 4, 5, 6]:
    w(ws, 15, c, None)

# R16 — linha vazia
clr(ws, 16, max_col=6)

# R17 — título Fase 2 (merged A17:F17 — só A17)
w(ws, 17, 1, "FASE 2 — Operação Comercial (Indicativo — após go-live Fase 1)")

# R18–R23 — tarefas Fase 2
crono_f2 = [
    ("08.CRM",                   "Pipeline e dashboards",                  1, 1, 1, None),
    ("09.Migração Pipedrive",    "Histórico CRM Pipedrive → Odoo",         2, 1, 2, "[POR VALIDAR]"),
    ("10.Compras",               "Compras e fornecedores",                  1, 2, 2, None),
    ("11.Inventário e Logística","Armazém e rotas",                         2, 3, 4, None),
    ("12.Testes + Formação F2",  "Testes e go-live Fase 2",                 1, 4, 5, None),
    ("13.Gestão Projeto F2",     "Acompanhamento Fase 2",                   1, 1, 5, "Transversal"),
]
for i, row in enumerate(crono_f2, start=18):
    for j, v in enumerate(row, start=1):
        w(ws, i, j, v)

# R24 — FASE 2 TOTAL
w(ws, 24, 1, "FASE 2 TOTAL (indicativo)")
w(ws, 24, 3, 8)
for c in [2, 4, 5, 6]:
    w(ws, 24, c, None)

# R25 — linha vazia
clr(ws, 25, max_col=6)

# R26 — título Fase 3
w(ws, 26, 1, "FASE 3 — Expansão (Indicativo [POR VALIDAR] — após go-live Fase 2)")
for c in range(2, 7):
    w(ws, 26, c, None)

# R27 — tarefas Fase 3
crono_f3 = [
    ("14.Manufatura (BoM)",             "Listas materiais e fabrico",        2, 1, 2, "[POR VALIDAR]"),
    ("15.Multiempresa — Starplast SRLE","2.ª empresa, intercompany",         3, 2, 4, "[POR VALIDAR]"),
    ("16.Contabilidade Analítica",      "Analítica por projeto/marca",       2, 4, 5, None),
    ("17.Módulos Adicionais",           "Sign, Barcode, Helpdesk — a def.",  2, 5, 6, "[POR VALIDAR] — a definir"),
]
for i, row in enumerate(crono_f3, start=27):
    for j, v in enumerate(row, start=1):
        w(ws, i, j, v)

# R31 — FASE 3 TOTAL
w(ws, 31, 1, "FASE 3 TOTAL (indicativo)")
w(ws, 31, 3, 9)
for c in [2, 4, 5, 6]:
    w(ws, 31, c, None)

# Limpar R32–R35
for r in range(32, 36):
    clr(ws, r, max_col=6)


# ===========================================================================
# DB — reescrever tarefas (R2 em diante)
# ===========================================================================
ws = wb["DB"]

db_rows = [
    # FASE 1
    ("Virgin Details, Lda", "01.Gestão de Projeto",
     "Kick-off interno e cliente; reuniões quinzenais; reporting; encerramento de fase",
     "Gestão de Projeto", "CS_Gestor", 90, 8, 1, 720, None, None),
    ("Virgin Details, Lda", "02.Kick-off + Configuração",
     "Setup empresa; 5 utilizadores; grupos de acesso; e-mail; templates",
     "Parametrização", "CS_Consultor", 65, 12, 1.5, 780, None, None),
    ("Virgin Details, Lda", "03.Contactos e Produtos",
     "Clientes, fornecedores; produtos e serviços; importação de PHC GO e Excel",
     "Parametrização", "CS_Consultor", 65, 8, 1, 520, None, "Importação a partir de PHC GO"),
    ("Virgin Details, Lda", "04.Vendas e Faturação",
     "Templates orçamento; encomendas de venda; faturas; notas de crédito; faturação certificada AT",
     "Parametrização", "CS_Consultor", 65, 16, 2, 1040, None, None),
    ("Virgin Details, Lda", "05.Financeiro l10n_pt",
     "Plano de contas PT; grupos IVA; SAF-T; plugin certificação AT; condições pagamento; adiantamentos; banca",
     "Parametrização", "CS_Consultor", 65, 16, 2, 1040, None,
     "Plugin AT licenciado separadamente — confirmar antes do kick-off"),
    ("Virgin Details, Lda", "06.Migração PHC GO",
     "Importação catálogo produtos; clientes e fornecedores; saldos em aberto de PHC GO",
     "Migração", "CS_Consultor", 65, 16, 2, 1040, None, "Cliente fornece dados em Excel limpo"),
    ("Virgin Details, Lda", "07.Testes + Formação + Arranque",
     "Testes faturação end-to-end; formação 5 utilizadores; suporte go-live; documentação",
     "Formação", "CS_Consultor", 65, 4, 0.5, 260, None, None),
    ("Virgin Details, Lda", "SUBTOTAL — Fase 1 (Faturação)",
     None, None, None, None, 80, 10, 5400, None, None),
    # FASE 2
    ("Virgin Details, Lda", "08.CRM",
     "Pipeline; fases e probabilidades; templates e-mail por fase; atividades; dashboards",
     "Parametrização", "CS_Consultor", 65, 8, 1, 520, None, None),
    ("Virgin Details, Lda", "09.Migração Pipedrive",
     "Exportação, mapeamento e importação de contactos e histórico CRM do Pipedrive para Odoo",
     "Migração", "CS_Consultor", 65, 16, 2, 1040, "SIM",
     "Esforço depende do volume e qualidade dos dados Pipedrive"),
    ("Virgin Details, Lda", "10.Compras",
     "Fornecedores europeus (4 marcas); preços custo; lead times; regras MTO; encomenda automática",
     "Parametrização", "CS_Consultor", 65, 8, 1, 520, None, None),
    ("Virgin Details, Lda", "11.Inventário e Logística",
     "Armazém Maia PT; rotas MTO+Comprar; receção Itália (Starplast/Sunerg); dropshipping Espanha",
     "Parametrização", "CS_Consultor", 65, 16, 2, 1040, None, None),
    ("Virgin Details, Lda", "12.Testes + Formação Fase 2",
     "Testes end-to-end Fase 2; formação por perfil; suporte go-live",
     "Formação", "CS_Consultor", 65, 8, 1, 520, None, None),
    ("Virgin Details, Lda", "13.Gestão de Projeto Fase 2",
     "Acompanhamento Fase 2 (taxa GP)",
     "Gestão de Projeto", "CS_Gestor", 90, 8, 1, 720, None, None),
    ("Virgin Details, Lda", "SUBTOTAL — Fase 2 (Operação Comercial)",
     None, None, None, None, 64, 8, 4360, None, None),
    # FASE 3
    ("Expansão (Fase 3)", "14.Manufatura (BoM)",
     "Listas de Materiais; Ordens de Fabrico; controlo componentes; custo produto final",
     "Parametrização", "CS_Consultor", 65, 16, 2, 1040, "SIM",
     "Âmbito depende de levantamento dos kits mais frequentes"),
    ("Expansão (Fase 3)", "15.Multiempresa — Starplast SRLE",
     "Configuração 2.ª empresa do grupo; utilizadores partilhados; transações intercompany",
     "Parametrização", "CS_Consultor", 65, 24, 3, 1560, "SIM",
     "Depende de confirmação estrutura legal Starplast SRLE"),
    ("Expansão (Fase 3)", "16.Contabilidade Analítica",
     "Analítica por projeto / marca (4 marcas) / cliente",
     "Parametrização", "CS_Consultor", 65, 16, 2, 1040, "SIM", None),
    ("Expansão (Fase 3)", "17.Módulos Adicionais",
     "Sign (aprovação eletrónica), Barcode (armazém), Documents, Helpdesk — a definir",
     "Parametrização", "CS_Consultor", 65, 16, 2, 1040, "SIM",
     "A orçamentar após levantamento de prioridades"),
    ("Expansão (Fase 3)", "SUBTOTAL — Fase 3 (Expansão)",
     None, None, None, None, 72, 9, 4680, None, None),
    # TOTAL
    (None, "TOTAL GERAL", None, None, None, None, 216, 27, 14440, None, None),
]

for i, row in enumerate(db_rows, start=2):
    for j, v in enumerate(row, start=1):
        w(ws, i, j, v)

# Limpar linhas antigas que ficaram para além dos novos dados
last_new = 2 + len(db_rows)
for r in range(last_new, last_new + 5):
    clr(ws, r, max_col=11)


# ===========================================================================
# EMPRESA — actualizar totais
# ===========================================================================
ws = wb["Empresa"]

w(ws, 3, 1, "Virgin Details, Lda — Fase 1 (Faturação)")
w(ws, 3, 2, 80)
w(ws, 3, 3, 5400)

w(ws, 4, 1, "Virgin Details, Lda — Fase 2 (Operação Comercial, indicativo)")
w(ws, 4, 2, 64)
w(ws, 4, 3, 4360)

w(ws, 5, 1, "Expansão — Fase 3 (indicativo [POR VALIDAR])")
w(ws, 5, 2, 72)
w(ws, 5, 3, 4680)

w(ws, 6, 1, "Total Geral")
w(ws, 6, 2, 216)
w(ws, 6, 3, 14440)

clr(ws, 7, max_col=3)


# ===========================================================================
# CONFIG — actualizar data
# ===========================================================================
ws = wb["Config"]
w(ws, 8, 2, "25/05/2026")


# ===========================================================================
# GUARDAR
# ===========================================================================
wb.save(XLSX)
print("Excel guardado:", XLSX)
print()
print("Resumo dos totais:")
print("  Fase 1 — Faturação:  80h / 10 dias /  5 400 €")
print("  Fase 2 — Operação:   64h /  8 dias /  4 360 € (indicativo)")
print("  Fase 3 — Expansão:   72h /  9 dias /  4 680 € [POR VALIDAR]")
print("  Total 1.º ano s/IVA:          7 752 €")
print("  IVA 23%:                       1 782,96 €")
print("  Total 1.º ano c/IVA:          9 534,96 €")
print("  Recorrência anual s/IVA:       5 340 €")
