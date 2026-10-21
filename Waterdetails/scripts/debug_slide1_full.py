import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

SRC = "c:/Projetos/Waterdetails/docs/WaterDetails.pptx"
p = Presentation(SRC)

for si, slide in enumerate(p.slides):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        tf = shape.text_frame
        full = "".join(pa.text for pa in tf.paragraphs).strip()
        if not full:
            continue
        runs_count = sum(len(pa.runs) for pa in tf.paragraphs)
        if runs_count > 1 or len(full) > 60:
            print(f"S{si+1} shape={shape.name} runs={runs_count}")
            print(f"  full: {repr(full[:120])}")
            for pi, para in enumerate(tf.paragraphs):
                for ri, run in enumerate(para.runs):
                    if run.text.strip():
                        print(f"    p{pi}r{ri}: {repr(run.text[:80])}")
