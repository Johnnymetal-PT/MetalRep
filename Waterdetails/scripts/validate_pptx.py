import sys, zipfile
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

DST = "c:/Projetos/Waterdetails/docs/WaterDetails_Orcamento_v1.pptx"

with zipfile.ZipFile(DST) as z:
    slides = [f for f in z.namelist() if f.startswith("ppt/slides/slide") and "rels" not in f]
    print(f"ZIP valido - {len(slides)} slides")
    has_img = any("media" in f for f in z.namelist())
    print(f"Imagens no ZIP: {has_img}")

p = Presentation(DST)
print(f"Slides: {len(p.slides)}")
for i, slide in enumerate(p.slides):
    texts = []
    has_pic = False
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = " ".join(pa.text.strip() for pa in shape.text_frame.paragraphs if pa.text.strip())
            if t:
                texts.append(t[:70])
        if shape.shape_type == 13:
            has_pic = True
    sep = " | "
    print(f"Slide {i+1} [pic={has_pic}]: {sep.join(texts[:5])}")

all_text = ""
for slide in p.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            all_text += " ".join(pa.text for pa in shape.text_frame.paragraphs)

print("\nVerificacoes (DEVE EXISTIR):")
must_exist = [
    ("Virgin Details, Lda", "Nome cliente"),
    ("waterdetails | Odoo 19", "Footer correcto"),
    ("10 dias", "Fase 1 dias"),
    ("Fase 1: Fatura", "Label capa Fase 1"),
    ("Fase 2 — Opera", "Slide 5 titulo Fase 2"),
    ("Fase 3 — Expan", "Slide 5 titulo Fase 3"),
    ("Migração Pipedrive", "Migração Pipedrive"),
    ("l10n_pt", "Localizacao PT"),
    ("Stefano Biglieri", "Licenciamento separado"),
    ("PHC GO", "Migracao PHC GO"),
    ("5\xa0400", "Valor servicos Fase 1"),
    ("7\xa0752", "Total 1 ano"),
    ("5\xa0340", "Recorrencia anual"),
]
for term, label in must_exist:
    found = term in all_text
    print(f"  {label}: {'OK' if found else 'FALHOU — nao encontrado'}")

print("\nVerificacoes (NAO DEVE EXISTIR):")
must_not = [
    ("Mercado Abastecedor", "Texto antigo"),
    ("S01705", "REF antiga"),
    ("33\xa0891", "Valor monetario antigo"),
    ("TimeCloud", "TimeCloud"),
    ("Controlo de Acessos", "Acessos antigo"),
    ("240h", "Horas antigas"),
]
for term, label in must_not:
    found = term in all_text
    print(f"  {label}: {'FALHOU — ainda presente' if found else 'OK — removido'}")
