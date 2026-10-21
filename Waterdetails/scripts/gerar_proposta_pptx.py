#!/usr/bin/env python
# -*- coding: utf-8 -*-
import sys, os, tempfile
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TPL = r'C:\Projetos\_pg_template\doc_templates\assets\Template PPTX.pptx'
OUT = r'C:\Projetos\Waterdetails\docs\PROPOSTA_WATERDETAILS_V1.pptx'
FONT = "Aptos"

C_NAVY  = RGBColor(0x0E, 0x18, 0x27)
C_TITLE = RGBColor(0x14, 0x1C, 0x29)
C_BLUE  = RGBColor(0x00, 0x99, 0xE6)
C_GRAY  = RGBColor(0x67, 0x74, 0x87)
C_BODY  = RGBColor(0x20, 0x29, 0x38)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_LGRAY = RGBColor(0xF2, 0xF5, 0xF8)
C_MGRAY = RGBColor(0xD0, 0xD8, 0xE4)
C_AMBER = RGBColor(0xFF, 0xF2, 0xCC)
C_THDR  = RGBColor(0x1F, 0x38, 0x64)
C_GANTT = RGBColor(0x00, 0x99, 0xE6)
C_GANT2 = RGBColor(0x0E, 0x18, 0x27)
C_PH2   = RGBColor(0xE8, 0xF4, 0xFD)
C_TOTAL = RGBColor(0x1F, 0x38, 0x64)

SW = Inches(13.33)
SH = Inches(7.50)

# ── extract logo ──────────────────────────────────────────────────────────────
tpl_prs = Presentation(TPL)
logo_blob, logo_ext = None, 'png'
for shape in tpl_prs.slides[0].shapes:
    if shape.shape_type == 13:
        logo_blob = shape.image.blob
        logo_ext  = shape.image.ext
        break

def _logo_tmp():
    if not logo_blob: return None
    f = tempfile.NamedTemporaryFile(delete=False, suffix=f'.{logo_ext}')
    f.write(logo_blob); f.close()
    return f.name

# ── presentation ──────────────────────────────────────────────────────────────
prs  = Presentation()
prs.slide_width  = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

# ── helpers ───────────────────────────────────────────────────────────────────
def tb(slide, text, l, t, w, h, sz=10, bold=False, color=C_BODY,
       align=PP_ALIGN.LEFT, italic=False, wrap=True):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf  = box.text_frame; tf.word_wrap = wrap
    p   = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.name = FONT; run.font.size = Pt(sz)
    run.font.bold = bold; run.font.italic = italic
    run.font.color.rgb = color
    return box

def tb_lines(slide, lines, l, t, w, h, sz=9.5, wrap=True):
    """lines = [(text, bold, color|None)]"""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf  = box.text_frame; tf.word_wrap = wrap
    first = True
    for (text, bold, col) in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(1.5); p.alignment = PP_ALIGN.LEFT
        run = p.add_run(); run.text = text
        run.font.name = FONT; run.font.size = Pt(sz); run.font.bold = bold
        run.font.color.rgb = col if col else C_BODY
    return box

def rect(slide, l, t, w, h, fill=None, line=False):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = fill
    else:    s.fill.background()
    if not line: s.line.fill.background()
    return s

def hline(slide, l, t, w, color=C_MGRAY):
    c = slide.shapes.add_connector(1, Inches(l), Inches(t), Inches(l+w), Inches(t))
    c.line.color.rgb = color; c.line.width = Pt(0.5)

def logo(slide, l, t, w, h):
    p = _logo_tmp()
    if p:
        slide.shapes.add_picture(p, Inches(l), Inches(t), Inches(w), Inches(h))
        os.unlink(p)

def cell(tbl, r, c, text, sz=8.5, bold=False, color=C_BODY, bg=None,
         align=PP_ALIGN.LEFT, italic=False):
    cl = tbl.cell(r, c); cl.text = ""
    tf = cl.text_frame; tf.word_wrap = True
    p  = tf.paragraphs[0]; p.alignment = align
    rn = p.add_run(); rn.text = text
    rn.font.name = FONT; rn.font.size = Pt(sz)
    rn.font.bold = bold; rn.font.italic = italic; rn.font.color.rgb = color
    if bg: cl.fill.solid(); cl.fill.fore_color.rgb = bg

def euro(v): return f"€ {v:,.2f}".replace(",","X").replace(".",",").replace("X",".")

FOOTER = "Virgin Details, Lda  ·  Parâmetro Global  ·  20/05/2026"

def content_slide(title, sub, page):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = C_WHITE
    logo(s, 11.20, 0.35, 1.35, 0.27)
    tb(s, title, 0.70, 0.52, 10.0, 0.50, sz=20, bold=True,  color=C_TITLE)
    tb(s, sub,   0.70, 1.10, 10.0, 0.28, sz=8,  bold=False, color=C_GRAY)
    hline(s, 0.55, 1.45, 12.20)
    hline(s, 0.55, 7.02, 12.20)
    tb(s, FOOTER, 0.55, 7.12, 9.00, 0.22, sz=6, color=C_GRAY)
    tb(s, f"{page:02d}", 12.35, 7.12, 0.60, 0.22, sz=7, bold=True,
       color=C_BLUE, align=PP_ALIGN.RIGHT)
    return s

# ── SLIDE 1 — CAPA ────────────────────────────────────────────────────────────
s1 = prs.slides.add_slide(BLANK)
bg = s1.background.fill; bg.solid(); bg.fore_color.rgb = C_WHITE
rect(s1, 0.00, 7.12, 13.33, 0.38, fill=C_NAVY)
logo(s1, 0.65, 0.55, 3.40, 0.69)
tb(s1, "Proposta de Implementação Odoo Enterprise",
   0.70, 2.00, 9.50, 0.42, sz=12, color=C_BLUE)
tb(s1, "Implementação Odoo\nVirgin Details, Lda",
   0.70, 2.52, 9.50, 1.10, sz=28, bold=True, color=C_TITLE, wrap=True)
tb(s1, "CRM · Vendas · Compras · Inventário · Faturação PT · Projetos\n"
       "5 Utilizadores  ·  Odoo SH (cloud)  ·  Localização PT (l10n_pt)",
   0.70, 3.80, 9.50, 0.70, sz=11, color=C_BODY, wrap=True)
tb(s1, "Ref. WD-2026/01  ·  20 de maio de 2026",
   0.70, 4.65, 8.00, 0.28, sz=9, color=C_GRAY)
tb(s1, "Preparado por Parâmetro Global — Consultoria Informática S.A.  ·  it.odoo@parametro.global",
   0.70, 6.80, 11.00, 0.22, sz=7, color=C_WHITE)

# ── SLIDE 2 — ENQUADRAMENTO ───────────────────────────────────────────────────
s2 = content_slide("Serviços — Enquadramento",
                   "Contexto, âmbito e metodologia da proposta", 2)
tb_lines(s2, [
    ("Após a reunião de apresentação do protótipo em 20 de maio de 2026, a Parâmetro Global propõe a "
     "implementação da plataforma Odoo Enterprise na Virgin Details, Lda (Maia, Porto), representante "
     "comercial exclusiva de marcas europeias de tratamento de águas (Starplast, HASA, Jeremias, Sunerg).",
     False, C_BODY),
    ("", False, None),
    ("O ciclo de serviços inclui:", True, C_TITLE),
    ("  •  Parametrização das áreas funcionais: CRM, Vendas, Compras, Inventário, Financeiro PT", False, C_BODY),
    ("  •  Migração de dados — catálogo, clientes, fornecedores e saldos em aberto (PHC GO / Excel)", False, C_BODY),
    ("  •  Migração do histórico CRM do Pipedrive (10+ anos de dados) para o Odoo", False, C_BODY),
    ("  •  Testes end-to-end do fluxo completo (lead → orçamento → compra → entrega → fatura)", False, C_BODY),
    ("  •  Formação de utilizadores (5 pessoas) e suporte ao go-live", False, C_BODY),
    ("", False, None),
    ("Responsabilidades do cliente:", True, C_TITLE),
    ("  •  Designar interlocutor disponível para validações (mín. 4h/semana durante a implementação)", False, C_BODY),
    ("  •  Fornecer dados mestres em Excel limpo segundo template Parâmetro antes da fase de migração", False, C_BODY),
    ("", False, None),
    ("Plataforma: Odoo Enterprise saas~19.3  ·  Odoo SH (cloud)  ·  5 utilizadores  ·  l10n_pt", False, C_BLUE),
], 0.70, 1.60, 11.80, 5.30, sz=9.5)

# ── SLIDE 3 — SERVIÇOS FASE 1 ─────────────────────────────────────────────────
tasks_p1 = [
    ("01", "Gestão de Projeto",        "CS_Gestor",    32, 4,  2880.00, ""),
    ("02", "Configuração Inicial",      "CS_Consultor", 16, 2,  1040.00, ""),
    ("03", "Contactos",                 "CS_Consultor",  8, 1,   520.00, ""),
    ("04", "CRM",                       "CS_Consultor", 16, 2,  1040.00, ""),
    ("05", "Vendas",                    "CS_Consultor", 24, 3,  1560.00, ""),
    ("06", "Compras",                   "CS_Consultor", 16, 2,  1040.00, ""),
    ("07", "Inventário e Logística",    "CS_Consultor", 16, 2,  1040.00, ""),
    ("08", "Financeiro — l10n_pt",      "CS_Consultor", 24, 3,  1560.00, ""),
    ("09", "Migração de Dados",         "CS_Consultor", 24, 3,  1560.00, ""),
    ("10", "Migração Pipedrive",        "CS_Consultor", 16, 2,  1040.00, "⚠"),
    ("11", "Testes e Validação",        "CS_Consultor", 16, 2,  1040.00, ""),
    ("12", "Formação e Arranque",       "CS_Consultor", 24, 3,  1560.00, ""),
]
s3 = content_slide("Serviços — Fase 1: Core Operacional",
                   "Virgin Details, Lda  ·  232 horas  ·  29 dias  ·  € 15.880,00 (sem IVA)", 3)
NR3 = len(tasks_p1) + 2
t3 = s3.shapes.add_table(NR3, 6,
     Inches(0.55), Inches(1.60), Inches(12.20), Inches(5.25)).table
for i, w in enumerate([0.42, 5.30, 1.10, 0.80, 2.00, 0.58]):
    t3.columns[i].width = Inches(w)
for ri in range(NR3):
    t3.rows[ri].height = Inches(5.25 / NR3)

for ci, h in enumerate(["#", "Área / Rúbrica", "Perfil", "H", "Dias", "Valor (€)"]):
    cell(t3, 0, ci, h, sz=8.5, bold=True, color=C_WHITE, bg=C_THDR,
         align=PP_ALIGN.CENTER if ci != 1 else PP_ALIGN.LEFT)

for ri, (num, area, perfil, hours, days, value, pv) in enumerate(tasks_p1):
    r = ri + 1
    bg_row = C_WHITE if ri % 2 == 0 else C_LGRAY
    bg_pv  = C_AMBER if pv else bg_row
    cell(t3, r, 0, num,           sz=8, bg=bg_row,  align=PP_ALIGN.CENTER)
    cell(t3, r, 1, area + (" ⚠" if pv else ""),
                                  sz=8, bg=bg_pv if pv else bg_row,
                                  color=C_BODY)
    cell(t3, r, 2, perfil,        sz=7.5, bg=bg_row, color=C_GRAY, italic=True)
    cell(t3, r, 3, str(hours),    sz=8, bg=bg_row,  align=PP_ALIGN.CENTER)
    cell(t3, r, 4, str(days),     sz=8, bg=bg_row,  align=PP_ALIGN.CENTER)
    cell(t3, r, 5, euro(value),   sz=8, bg=bg_row,  align=PP_ALIGN.RIGHT)

total_r = len(tasks_p1) + 1
th = sum(t[3] for t in tasks_p1)
td = sum(t[4] for t in tasks_p1)
tv = sum(t[5] for t in tasks_p1)
cell(t3, total_r, 0, "TOTAL", sz=9, bold=True, color=C_WHITE, bg=C_TOTAL, align=PP_ALIGN.CENTER)
cell(t3, total_r, 1, "Fase 1 — Core Operacional",  sz=9, bold=True, color=C_WHITE, bg=C_TOTAL)
cell(t3, total_r, 2, "",        sz=9, bg=C_TOTAL, color=C_WHITE)
cell(t3, total_r, 3, str(th),   sz=9, bold=True, color=C_WHITE, bg=C_TOTAL, align=PP_ALIGN.CENTER)
cell(t3, total_r, 4, str(td),   sz=9, bold=True, color=C_WHITE, bg=C_TOTAL, align=PP_ALIGN.CENTER)
cell(t3, total_r, 5, euro(tv),  sz=9, bold=True, color=C_WHITE, bg=C_TOTAL, align=PP_ALIGN.RIGHT)
tb(s3, "⚠ = Por Validar — estimativa sujeita a revisão após análise do volume de dados Pipedrive",
   0.55, 6.90, 12.00, 0.25, sz=7, color=C_GRAY, italic=True)

# ── SLIDE 4 — CRONOGRAMA ─────────────────────────────────────────────────────
s4 = content_slide("Cronograma — Fase 1 (Indicativo)",
                   "Duração estimada: 12 semanas  ·  Início: 2–3 semanas após assinatura", 4)
gantt = [
    ("Gestão de Projeto",    4,  list(range(1, 13)), True),
    ("Configuração Inicial", 2,  [1, 2],             False),
    ("Contactos",            1,  [2],                False),
    ("CRM",                  2,  [3, 4],             False),
    ("Vendas",               3,  [3, 4, 5],          False),
    ("Compras",              2,  [4, 5],             False),
    ("Inventário",           2,  [5, 6],             False),
    ("Financeiro",           3,  [6, 7, 8],          False),
    ("Migração Dados",       3,  [8, 9, 10],         False),
    ("Pipedrive ⚠",          2,  [9, 10],            False),
    ("Testes",               2,  [10, 11],           False),
    ("Formação / Go-live",   3,  [11, 12, 12],       False),
]
NW = 12
NRG = len(gantt) + 1
NCG = 2 + NW
t4 = s4.shapes.add_table(NRG, NCG,
     Inches(0.55), Inches(1.60), Inches(12.20), Inches(5.25)).table
t4.columns[0].width = Inches(2.55)
t4.columns[1].width = Inches(0.48)
for ci in range(2, NCG):
    t4.columns[ci].width = Inches((12.20 - 2.55 - 0.48) / NW)
for ri in range(NRG):
    t4.rows[ri].height = Inches(5.25 / NRG)

cell(t4, 0, 0, "Área / Rúbrica", sz=8, bold=True, color=C_WHITE, bg=C_THDR)
cell(t4, 0, 1, "Dias",           sz=7.5, bold=True, color=C_WHITE, bg=C_THDR, align=PP_ALIGN.CENTER)
for w in range(1, NW + 1):
    cell(t4, 0, w + 1, f"S{w}", sz=7, bold=True, color=C_WHITE, bg=C_THDR, align=PP_ALIGN.CENTER)

for ri, (area, days, active, is_pm) in enumerate(gantt):
    r = ri + 1
    is_warn = "⚠" in area
    bg_row = C_AMBER if is_warn else (C_WHITE if ri % 2 == 0 else C_LGRAY)
    cell(t4, r, 0, area,     sz=8, bg=bg_row, color=C_BODY)
    cell(t4, r, 1, str(days), sz=8, bg=bg_row, align=PP_ALIGN.CENTER)
    for w in range(1, NW + 1):
        if w in active:
            gc = C_GANT2 if is_pm else C_GANTT
            cell(t4, r, w + 1, "", sz=7, bg=gc)
        else:
            cell(t4, r, w + 1, "", sz=7, bg=bg_row)

tb(s4, "■ Gestão de Projeto (todo o projeto)  ■ Implementação   ⚠ = Por Validar  |  Fase 2 inicia na semana 17 (após go-live)",
   0.55, 6.88, 12.00, 0.28, sz=7, color=C_GRAY, italic=True)

# ── SLIDE 5 — RESUMO / PAGAMENTO ─────────────────────────────────────────────
s5 = content_slide("Resumo da Proposta — Formas de Pagamento",
                   "Investimento total e modalidades de pagamento propostas", 5)
tb(s5, "Investimento — 1.º Ano", 0.65, 1.58, 5.80, 0.30, sz=10, bold=True, color=C_TITLE)
rows5 = [
    ("Serviços — Fase 1",              "€ 15.880,00",          False, C_WHITE),
    ("  Gestão de Projeto (32h·€90/h)", "€ 2.880,00",          False, C_LGRAY),
    ("  Implementação (200h·€65/h)",   "€ 13.000,00",          False, C_LGRAY),
    ("Software Odoo Enterprise (5 us.)","⚠ a confirmar",       False, C_AMBER),
    ("Plugin Certificação AT",          "⚠ a confirmar",       False, C_AMBER),
    ("TOTAL 1.º Ano (sem IVA)",         "€ 15.880 + licenças", True,  C_TOTAL),
    ("IVA 23%",                         "€ 3.652,40 + IVA",    False, C_LGRAY),
    ("TOTAL 1.º Ano (com IVA)",         "€ 19.532,40 + lic.",  True,  C_TOTAL),
]
t5 = s5.shapes.add_table(len(rows5) + 1, 2,
     Inches(0.55), Inches(1.95), Inches(5.90), Inches(4.85)).table
t5.columns[0].width = Inches(4.10); t5.columns[1].width = Inches(1.80)
for ri in range(len(rows5) + 1):
    t5.rows[ri].height = Inches(4.85 / (len(rows5) + 1))
cell(t5, 0, 0, "Rúbrica",     sz=8.5, bold=True, color=C_WHITE, bg=C_THDR)
cell(t5, 0, 1, "Valor (s/IVA)", sz=8.5, bold=True, color=C_WHITE, bg=C_THDR, align=PP_ALIGN.RIGHT)
for ri, (lbl, val, bold, bg) in enumerate(rows5):
    tc = C_WHITE if bg in (C_TOTAL,) else C_BODY
    cell(t5, ri+1, 0, lbl, sz=8, bold=bold, color=tc, bg=bg)
    vc = C_WHITE if bg == C_TOTAL else C_BODY
    cell(t5, ri+1, 1, val, sz=8, bold=bold, color=vc, bg=bg, align=PP_ALIGN.RIGHT)

tb(s5, "Formas de Pagamento", 6.65, 1.58, 5.80, 0.30, sz=10, bold=True, color=C_TITLE)
tb_lines(s5, [
    ("Pagamento por fases:", True, C_TITLE),
    ("", False, None),
    ("  ①  30%  no Kick-off — assinatura do contrato", False, C_BODY),
    ("  ②  40%  a meio — parametrização concluída e aceite", False, C_BODY),
    ("  ③  30%  no Go-live — formação e arranque aceites", False, C_BODY),
    ("", False, None),
    ("Exploração anual (a partir do 2.º Ano):", True, C_TITLE),
    ("", False, None),
    ("  Software Odoo Enterprise  →  a confirmar (Stefano)", False, C_BODY),
    ("  Plugin Certificação AT  →  a confirmar (Stefano)", False, C_BODY),
    ("  Apoio Após Venda Parâmetro Global  →  € 2.400/ano", False, C_BLUE),
    ("", False, None),
    ("Expansão — Fase 2 (indicativo, após Fase 1):", True, C_TITLE),
    ("  9 dias base (BoM + Multiempresa + Analítica)", False, C_BODY),
    ("  →  € 4.680,00 sem IVA  ⚠ Por Validar", False, C_GRAY),
], 6.65, 1.95, 5.90, 4.85, sz=9.5)

# ── SLIDE 6 — TOTAL GERAL ────────────────────────────────────────────────────
s6 = content_slide("Serviços — Total Geral",
                   "Fase 1 (Core Operacional) + Fase 2 (Expansão Indicativa)", 6)
rows6 = [
    ("FASE 1 — Core Operacional",             "Virgin Details, Lda",                  232, 29, 15880.00, True,  C_THDR, ""),
    ("  Gestão de Projeto",                   "CS_Gestor  ·  €90/h",                   32,  4,  2880.00, False, C_WHITE,""),
    ("  Parametrização + Migração + Arranque","CS_Consultor  ·  €65/h",               200, 25, 13000.00, False, C_LGRAY,""),
    ("FASE 2 — Expansão (Indicativo)",        "A confirmar após go-live Fase 1",        72,  9,  4680.00, True,  C_THDR, "⚠"),
    ("  Manufatura / BoM",                    "CS_Consultor",                           32,  4,  2080.00, False, C_PH2,  "⚠"),
    ("  Multiempresa — Starplast SRLE",       "CS_Consultor",                           24,  3,  1560.00, False, C_WHITE,"⚠"),
    ("  Contabilidade Analítica",             "CS_Consultor",                           16,  2,  1040.00, False, C_LGRAY,""),
    ("  Módulos Adicionais (Sign/Barcode/…)", "CS_Consultor  ·  âmbito a definir",      0,  0,     0.00, False, C_AMBER,"⚠"),
    ("TOTAL GERAL SERVIÇOS",                  "Fase 1 + Fase 2 base",                 304, 38, 20560.00, True,  C_NAVY, ""),
]
NR6 = len(rows6) + 1
t6 = s6.shapes.add_table(NR6, 6,
     Inches(0.55), Inches(1.60), Inches(12.20), Inches(5.20)).table
for i, w in enumerate([4.30, 3.10, 0.85, 0.75, 2.15, 0.55]):
    t6.columns[i].width = Inches(w)
for ri in range(NR6):
    t6.rows[ri].height = Inches(5.20 / NR6)
for ci, h in enumerate(["Componente", "Notas", "H", "Dias", "Valor (€)", "⚠"]):
    cell(t6, 0, ci, h, sz=8.5, bold=True, color=C_WHITE, bg=C_THDR,
         align=PP_ALIGN.CENTER if ci > 1 else PP_ALIGN.LEFT)
for ri, (lbl, note, h, d, v, bold, bg, pv) in enumerate(rows6):
    r = ri + 1
    tc = C_WHITE if bg in (C_THDR, C_NAVY) else C_BODY
    vs = euro(v) if v else "— a definir"
    hs = str(h) if h else "—"; ds = str(d) if d else "—"
    cell(t6, r, 0, lbl,  sz=8.5 if not bold else 9, bold=bold, color=tc, bg=bg)
    cell(t6, r, 1, note, sz=7.5, italic=not bold, color=(C_GRAY if tc==C_BODY else C_LGRAY), bg=bg)
    cell(t6, r, 2, hs,   sz=8.5, bold=bold, color=tc, bg=bg, align=PP_ALIGN.CENTER)
    cell(t6, r, 3, ds,   sz=8.5, bold=bold, color=tc, bg=bg, align=PP_ALIGN.CENTER)
    cell(t6, r, 4, vs,   sz=8.5, bold=bold, color=tc, bg=bg, align=PP_ALIGN.RIGHT)
    pbg = C_AMBER if pv else bg
    cell(t6, r, 5, pv,   sz=8,   color=C_BODY if pv else tc, bg=pbg, align=PP_ALIGN.CENTER)
tb(s6, "⚠ = Por Validar  ·  Licenciamento Odoo SH + Plugin AT não incluídos — proposta separada por Stefano Biglieri",
   0.55, 6.88, 12.00, 0.28, sz=7, color=C_GRAY, italic=True)

# ── SLIDE 7 — SUPORTE APÓS VENDA ─────────────────────────────────────────────
s7 = content_slide("Suporte Após a Venda",
                   "Contrato anual de suporte e manutenção — Parâmetro Global", 7)
rect(s7, 0.55, 1.62, 3.60, 1.30, fill=C_BLUE)
tb(s7, "€ 2.400 / ano",  0.55, 1.72, 3.60, 0.60, sz=22, bold=True,
   color=C_WHITE, align=PP_ALIGN.CENTER)
tb(s7, "€ 200 / mês", 0.55, 2.35, 3.60, 0.38, sz=12, color=C_WHITE,
   align=PP_ALIGN.CENTER)
rect(s7, 0.55, 3.10, 3.60, 0.80, fill=C_LGRAY)
tb(s7, "Proposta válida até\n20 de junho de 2026",
   0.65, 3.15, 3.40, 0.68, sz=9, bold=True, color=C_TITLE, wrap=True)
tb(s7, "Incluído no contrato AAV:", 4.45, 1.62, 7.90, 0.28, sz=9.5, bold=True, color=C_TITLE)
tb_lines(s7, [
    ("  •  Resolução de incidentes (tempo de resposta: 1 dia útil)", False, C_BODY),
    ("  •  Atualizações de configuração de baixa complexidade", False, C_BODY),
    ("  •  Aconselhamento funcional em modo de suporte", False, C_BODY),
    ("  •  Comunicação por e-mail: it.odoo@parametro.global", False, C_BODY),
    ("  •  Prioridade na gestão de tickets para clientes com contrato", False, C_BODY),
], 4.45, 1.98, 7.90, 2.00, sz=9.5)
tb(s7, "Não incluído no contrato de suporte:", 4.45, 4.15, 7.90, 0.28, sz=9.5, bold=True, color=C_TITLE)
tb_lines(s7, [
    ("  •  Novos módulos ou novas implementações (orçamentados separadamente)", False, C_BODY),
    ("  •  Upgrades de versão Odoo (orçamentados separadamente)", False, C_BODY),
    ("  •  Desenvolvimentos custom (orçamentados separadamente)", False, C_BODY),
], 4.45, 4.48, 7.90, 1.40, sz=9.5)

# ── SLIDE 8 — OBRIGADO ───────────────────────────────────────────────────────
s8 = prs.slides.add_slide(BLANK)
bg = s8.background.fill; bg.solid(); bg.fore_color.rgb = C_WHITE
rect(s8, 0.00, 7.12, 13.33, 0.38, fill=C_NAVY)
logo(s8, 0.65, 0.55, 3.40, 0.69)
rect(s8, 6.30, 1.80, 6.80, 4.10, fill=C_NAVY)
tb(s8, "Muito obrigado\npela confiança.",
   0.70, 2.20, 5.00, 1.80, sz=26, bold=True, color=C_TITLE, wrap=True)
tb(s8, "Próximos passos:", 0.70, 4.20, 5.20, 0.28, sz=10, bold=True, color=C_TITLE)
tb_lines(s8, [
    ("  1.  Confirmação do licenciamento Odoo SH — Stefano Biglieri", False, C_BODY),
    ("  2.  Assinatura do contrato de implementação", False, C_BODY),
    ("  3.  Kick-off — 2 a 3 semanas após assinatura", False, C_BODY),
    ("  4.  Início Fase 1 — Core Operacional (12 semanas)", False, C_BODY),
], 0.70, 4.55, 5.20, 2.00, sz=9.5)
tb(s8, "Pedro Coelho",
   7.10, 2.45, 5.70, 0.50, sz=18, bold=True, color=C_WHITE)
tb(s8, "Project Manager / Consultor Odoo",
   7.10, 3.05, 5.70, 0.30, sz=9, color=C_GRAY)
tb(s8, "it.odoo@parametro.global",
   7.10, 3.50, 5.70, 0.38, sz=12, bold=True, color=C_BLUE)
tb(s8, "Parâmetro Global — Consultoria Informática S.A.",
   7.10, 4.05, 5.70, 0.28, sz=8.5, color=C_WHITE)
tb(s8, "www.parametro.global",
   7.10, 4.40, 5.70, 0.25, sz=8, color=C_GRAY)

# ── SAVE ─────────────────────────────────────────────────────────────────────
prs.save(OUT)
print(f"OK: {OUT}")
print(f"Slides: {len(prs.slides)}")
