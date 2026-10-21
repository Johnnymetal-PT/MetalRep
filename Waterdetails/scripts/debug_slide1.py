import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

SRC = "c:/Projetos/Waterdetails/docs/WaterDetails.pptx"
p = Presentation(SRC)
slide = p.slides[0]
print("=== SLIDE 1 shapes ===")
for shape in slide.shapes:
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                t = run.text
                if t.strip():
                    # Mostrar repr para ver espaços especiais
                    print(f"  run repr: {repr(t[:100])}")
