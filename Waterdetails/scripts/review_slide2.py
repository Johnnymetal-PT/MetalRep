import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

p = Presentation("c:/Projetos/Waterdetails/docs/WaterDetails_Orcamento_v1_tmp.pptx")
slide = p.slides[1]
print("=== SLIDE 2 ===")
for shape in slide.shapes:
    if shape.has_text_frame:
        full = "".join(pa.text for pa in shape.text_frame.paragraphs).strip()
        if full:
            print(f"  {full[:100]}")
