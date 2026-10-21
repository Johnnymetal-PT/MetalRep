import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

p = Presentation("c:/Projetos/Waterdetails/docs/WaterDetails_Orcamento_v1.pptx")
for si in [0, 3, 4]:
    slide = p.slides[si]
    print(f"\n=== SLIDE {si+1} ===")
    for shape in slide.shapes:
        if shape.has_text_frame:
            full = "".join(pa.text for pa in shape.text_frame.paragraphs).strip()
            if full:
                print(f"  {full[:100]}")

# Confirmação: nenhum €___/dia restante
all_text = ""
for slide in p.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            all_text += "".join(pa.text for pa in shape.text_frame.paragraphs)
found = "€___" in all_text or "_/dia" in all_text
print(f"\n€___/dia restantes: {'SIM — VERIFICAR' if found else 'NENHUM — OK'}")
