import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

DST = "c:/Projetos/Waterdetails/docs/WaterDetails_Orcamento_v1.pptx"
p = Presentation(DST)
for i, slide in enumerate(p.slides):
    print(f"\n=== SLIDE {i+1} ===")
    for shape in slide.shapes:
        if shape.has_text_frame:
            full = "".join(pa.text for pa in shape.text_frame.paragraphs).strip()
            if full and shape.shape_type != 13:
                print(f"  {full[:110]}")
