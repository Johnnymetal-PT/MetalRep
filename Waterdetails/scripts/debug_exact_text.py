import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

SRC = "c:/Projetos/Waterdetails/docs/WaterDetails.pptx"
p = Presentation(SRC)

# Shapes específicos que não foram substituídas
targets = {
    (0, "CaixaDeTexto 5"): "slide1 descricao",
    (1, "CaixaDeTexto 10"): "slide2 resultado",
    (1, "CaixaDeTexto 21"): "slide2 parametros",
    (3, "CaixaDeTexto 14"): "slide4 licenciamento",
    (3, "CaixaDeTexto 17"): "slide4 PT+",
    (3, "CaixaDeTexto 27"): "slide4 nota",
    (3, "CaixaDeTexto 29"): "slide4 rodape",
    (4, "CaixaDeTexto 11"): "slide5 software detalhe",
    (4, "CaixaDeTexto 27"): "slide5 opcao ativa",
}

for (si, sname), label in targets.items():
    slide = p.slides[si]
    for shape in slide.shapes:
        if shape.name == sname and shape.has_text_frame:
            full = "".join(pa.text for pa in shape.text_frame.paragraphs)
            print(f"\n[{label}] repr:")
            print(repr(full))
