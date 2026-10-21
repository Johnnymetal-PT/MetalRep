#!/usr/bin/env python3
"""Migração da contabilidade a meio do ano: importação do SAF-T de contabilidade e consolidação.

Motor genérico do método TPL_MIGRACAO_CONTABILIDADE_MEIO_ANO (doc_templates). Cada subcomando é
um passo; escreve só com --confirm (e, em produção, --allow-prod) e acrescenta o resultado ao
diário do ensaio, que é o registo para replicar em produção. Tudo o que é do cliente vem do
ficheiro de configuração do projeto (ver README.md desta pasta): empresa, datas, ficheiros,
contas e regras da contabilidade.

Extraído da migração da Reckon.AI (Odoo 19, Exo PT+, SAF-T de janeiro a julho de 2026).

    O guard da framework bloqueia métodos que não estejam classificados. Os botões do importador
    da Exo (process_saft_file, action_write_records), o disparo manual da ação agendada
    (method_direct_trigger) e a fusão de contactos (action_merge) são classificados aqui como
    ESCRITA, só dentro deste processo, e continuam sujeitos às mesmas guardas.

Uso (a partir da raiz do repositório do projeto):
    python .pg_framework\\scripts\\pg_migracao_contabilidade\\motor.py <config.json> <ligação> estado
    python .pg_framework\\scripts\\pg_migracao_contabilidade\\motor.py <config.json> <ligação> importar --confirm
    <ligação> = ficheiro de ligação local (bases de ensaio) ou "producao" (canal do projeto).
"""
import argparse
import base64
import collections
import datetime
import json
import os
import sys
import time

sys.path.insert(0, os.path.join('.pg_framework', 'scripts'))
import pg_odoo_migration_guard as G  # noqa: E402

G.METODOS_ESCRITA = G.METODOS_ESCRITA | {
    'process_saft_file', 'action_write_records', 'method_direct_trigger', 'action_merge'}

PASSOS = ['estado', 'importar', 'ignorar', 'gravar', 'fundir_parceiros', 'parceiros_em_falta', 'parceiros', 'faltas',
          'grelhas', 'config_extra', 'impostos', 'recodificar', 'juntar', 'reclassificar', 'abertura', 'limpar',
          'extratos', 'taxonomia', 'config_fiscal', 'natureza', 'bloqueio', 'lancar', 'resumo_grelhas',
          'rascunho', 'repor_datas']

ap = argparse.ArgumentParser()
ap.add_argument('config', help='configuração do projeto (JSON), ex.: docs/plano_contas/migracao_contabilidade.json')
ap.add_argument('conn', help='ficheiro de ligação da base de ensaio, ou "producao"')
ap.add_argument('passo', choices=PASSOS)
ap.add_argument('--data', default=None,
                help='bloqueio: data a aplicar (defeito: véspera da abertura); "corte" = data de corte; '
                     '"repor" = o bloqueio anterior do cliente, nunca abaixo da data de corte')
ap.add_argument('--confirm', action='store_true')
ap.add_argument('--allow-prod', action='store_true', help='obrigatório para escrever em produção (conn = producao)')
ap.add_argument('--log', default=None)
args = ap.parse_args()

CFG = json.load(open(args.config, encoding='utf-8'))
PASTA = CFG['pasta']
F = {k: os.path.join(PASTA, v) for k, v in CFG.get('ficheiros', {}).items()}  # tabelas do projeto (fora do git)
SAFT = CFG['saft']
EMPRESA = CFG['empresa']
NOME = CFG['nome_importacao']
DATA_ABERTURA = CFG['data_abertura']          # data do lançamento de abertura (ex.: 2025-12-31)
DATA_CORTE = CFG['data_corte']                # último dia do histórico importado (ex.: 2026-07-31)
VESPERA = (datetime.date.fromisoformat(DATA_ABERTURA) - datetime.timedelta(days=1)).isoformat()
REF_ABERTURA = CFG['ref_abertura']
DIARIOS_APUR = CFG.get('diarios_apuramento', [])  # transferem saldos no fim do período: ficam fora das grelhas
PREFIXO_OUTROS = CFG.get('prefixo_outros_devedores_credores', '278')
RET = CFG.get('retencoes', {})
RET_PREFIXO = RET.get('prefixo', '242')
RET_MODELO10 = RET.get('contas_modelo10', [])  # só as retenções a crédito destas contas levam parceiro
DIARIO = os.path.join(PASTA, 'diario_ensaio.md')

if args.data is None:
    args.data = VESPERA
elif args.data == 'corte':
    args.data = DATA_CORTE

# "producao" liga pelo canal do projeto (credenciais no Odoo PG, sem chave em disco); qualquer
# outro valor é um ficheiro de ligação local (bases de ensaio). Em produção o guard só escreve
# com --allow-prod explícito, além de --confirm.
alvo = (G.resolver_alvo_por_projeto(CFG['project_key'], 'production') if args.conn == 'producao'
        else G.resolver_alvo(args.conn))
g = G.Guard(alvo, confirmed=args.confirm, allow_prod=args.allow_prod, log_path=args.log)
if args.conn == 'producao':  # o registo de produção não se mistura com o dos ensaios
    DIARIO = os.path.join(PASTA, 'diario_producao.md')
kw = g.kw


def w(model, method, a, k=None):
    return kw(model, method, a, k or {}, allow_write=True)


# Com impostos ligados às linhas, qualquer escrita nelas faz o Odoo recalcular o IVA (lançamento desequilibrado)
SEM_SYNC = {'context': {'skip_invoice_sync': True}}


def diario(titulo, linhas):
    os.makedirs(PASTA, exist_ok=True)
    novo = not os.path.exists(DIARIO)
    with open(DIARIO, 'a', encoding='utf-8') as fh:
        if novo:
            fh.write('# Diário do ensaio SAF-T e plano de contas\n\n'
                     f'Base: `{g.alvo.database}`\n\n')
        fh.write(f'## {datetime.datetime.now():%d/%m/%Y %H:%M} — {titulo}\n\n')
        for l in linhas:
            fh.write(f'- {l}\n')
        fh.write('\n')
    for l in linhas:
        print(' ', l)


def cron_saft():
    return kw('ir.cron', 'search_read', [[['name', '=', 'Process SAF-T Import'], ['active', 'in', [True, False]]]],
              {'fields': ['active']})[0]


def importacao():
    r = kw('l10n_pt.import.saft.configuration', 'search_read', [[['name', '=', NOME]]], {'fields': [
        'id', 'state', 'import_running', 'import_stage', 'import_progress', 'errors', 'errors_count', 'info_message',
        'total_partners', 'partners_to_create', 'partners_to_update', 'partners_to_ignore',
        'total_accounts', 'accounts_to_create', 'accounts_to_update', 'accounts_to_ignore',
        'total_journals', 'journals_to_create', 'journals_to_update', 'journals_to_ignore',
        'total_moves', 'moves_to_create', 'moves_to_ignore', 'date_start', 'date_end']}, )
    return r[0] if r else None


def resumo(i):
    return [f"estado {i['state']} | a correr {i['import_running']} ({i['import_stage'] or '-'}) | {i['import_progress'] or ''}",
            f"parceiros: total {i['total_partners']} C {i['partners_to_create']} U {i['partners_to_update']} I {i['partners_to_ignore']}",
            f"contas: total {i['total_accounts']} C {i['accounts_to_create']} U {i['accounts_to_update']} I {i['accounts_to_ignore']}",
            f"diários: total {i['total_journals']} C {i['journals_to_create']} U {i['journals_to_update']} I {i['journals_to_ignore']}",
            f"lançamentos: total {i['total_moves']} C {i['moves_to_create']} I {i['moves_to_ignore']}",
            f"período {i['date_start']} a {i['date_end']} | erros {i['errors_count']} {i['errors'] or ''}"]


if args.passo == 'estado':
    i = importacao()
    print('sem importação' if not i else '\n'.join(resumo(i)))
    if i:
        linhas = kw('l10n_pt.account.saft.import', 'read_group',
                    [[['configuration_id.name', '=', NOME]], ['rec_model'], ['rec_model', 'action', 'state', 'ignore']],
                    {'lazy': False})
        for l in linhas:
            print(f"  {l['rec_model']:28} {l['action']:7} {l['state']:8} ignore={l['ignore']} -> {l['__count']}")

elif args.passo == 'importar':
    if importacao():
        sys.exit('Já existe uma importação com este nome; usar "estado".')
    c = cron_saft()
    if not c['active']:
        w('ir.cron', 'write', [[c['id']], {'active': True}])
    dados = base64.b64encode(open(SAFT, 'rb').read()).decode()
    cid = w('l10n_pt.import.saft.configuration', 'create', [{
        'name': NOME, 'company_id': EMPRESA, 'saft_file': dados,
        'saft_filename': os.path.basename(SAFT),
        'update_partners': True, 'update_accounts': True, 'update_products': True,
        'update_journal_entries': True, 'disable_accounts': False, 'default_ignore': False}])
    w('l10n_pt.import.saft.configuration', 'process_saft_file', [[cid]])
    time.sleep(5)
    i = importacao()
    diario('Importação criada e ficheiro enviado para processamento',
           [f'ação agendada "Process SAF-T Import" ativada (id {c["id"]})',
            f'importação id {cid}, ficheiro original {os.path.basename(SAFT)}; Partners, Accounts, Products e '
            'Journal Entries ligados; Disable Accounts without Moves e Ignore By Value desligados',
            'Process File executado'] + resumo(i))

elif args.passo == 'ignorar':
    i = importacao()
    if i['state'] not in ('processed', 'error'):
        sys.exit(f"A importação está em '{i['state']}'; só se marca Ignorar depois de processada.")
    # Contas e grupos que já existem vêm do plano certificado: a designação certificada fica
    # (a atualização só traz a designação do legado, muitas vezes em maiúsculas).
    # Parceiros que já existem vêm da produção: a atualização reescreve nome, morada e contactos,
    # e em alguns passa a empresa a pessoa singular.
    alvo = kw('l10n_pt.account.saft.import', 'search_read',
              [[['configuration_id.name', '=', NOME], ['action', '=', 'update'], ['ignore', '=', False],
                ['rec_model', 'in', ['account.account', 'account.group', 'res.partner']]]],
              {'fields': ['rec_model', 'account_code', 'rec_name']})
    if alvo:
        w('l10n_pt.account.saft.import', 'write', [[x['id'] for x in alvo], {'ignore': True}])
    cont = collections.Counter(x['rec_model'] for x in alvo)
    diario('Atualizações de contas, grupos e parceiros existentes marcadas Ignorar',
           [f'{n} linhas {m}' for m, n in cont.items()] or ['nada a ignorar'])

elif args.passo == 'gravar':
    i = importacao()
    if i['state'] not in ('processed', 'error'):
        sys.exit(f"A importação está em '{i['state']}'.")
    w('l10n_pt.import.saft.configuration', 'action_write_records', [[i['id']]])
    time.sleep(5)
    diario('Create/Update Records executado', resumo(importacao()))

else:
    import csv
    import openpyxl

    def folha(nome):
        wb = openpyxl.load_workbook(F['plano'], read_only=True)
        return [r for r in wb[nome].iter_rows(min_row=2, values_only=True) if r[0]]

    def contas():
        return {a['code']: a for a in kw('account.account', 'search_read',
                                         [[['company_ids', 'in', [EMPRESA]]]],
                                         {'fields': ['code', 'name', 'account_type', 'reconcile']})}

    def n_linhas(acc_id):
        return kw('account.move.line', 'search_count', [[['account_id', '=', acc_id]]])

    def saft():
        """Entidades do MasterFiles (por subconta e por ID) e transações do SAF-T (por TransactionID)."""
        import xml.etree.ElementTree as ET
        ns = {'s': 'urn:OECD:StandardAuditFile-Tax:PT_1.04_01'}
        root = ET.parse(SAFT).getroot()

        def t(e, tag):
            x = e.find('s:' + tag, ns)
            return x.text.strip() if x is not None and x.text else None
        por_conta, por_id = {}, {}
        for tag in ('Customer', 'Supplier'):
            for e in root.findall(f'.//s:MasterFiles/s:{tag}', ns):
                ent = {'nif': t(e, f'{tag}TaxID'), 'nome': t(e, 'CompanyName')}
                por_conta[t(e, 'AccountID')] = ent
                por_id[t(e, f'{tag}ID')] = ent
        trans = {}
        for j in root.findall('.//s:GeneralLedgerEntries/s:Journal', ns):
            for tr in j.findall('s:Transaction', ns):
                trans[t(tr, 'TransactionID')] = {
                    'journal': t(j, 'JournalID'), 'date': t(tr, 'GLPostingDate') or t(tr, 'TransactionDate'),
                    'desc': t(tr, 'Description'), 'ent': t(tr, 'CustomerID') or t(tr, 'SupplierID'),
                    'linhas': [{'conta': t(l, 'AccountID'), 'desc': t(l, 'Description'),
                                'debit': float(t(l, 'DebitAmount') or 0), 'credit': float(t(l, 'CreditAmount') or 0)}
                               for grp in tr.findall('s:Lines', ns) for l in grp]}
        return por_conta, por_id, trans

    cache_parc = {}

    def parceiro(ent):
        """Parceiro do Odoo para uma entidade do SAF-T: pelo NIF, e na falta pelo nome exato.

        Só entidades principais (sem empresa-mãe): os contactos de uma empresa herdam-lhe o NIF, e
        uma pesquisa pelo NIF devolvia a empresa e os contactos. Se ainda houver mais de uma, fica a
        mais antiga, que é a que já existia em produção.
        """
        if not ent:
            return False
        chave = (ent['nif'], ent['nome'])
        if chave not in cache_parc:
            r = []
            base = [['parent_id', '=', False]]
            if ent['nif'] and ent['nif'] not in ('999999990',):
                nif = ent['nif'].replace(' ', '')
                # Também pelo NIF sem prefixo: depois de fundir_parceiros a ficha que fica pode ter o NIF
                # escrito de outra forma (Mouser: FR27512466913 no SAF-T, 27512466913 na ficha).
                k = nif_chave(nif) or nif
                r = kw('res.partner', 'search', [base + ['|', '|', ['vat', '=', nif], ['vat', '=ilike', '%' + nif],
                                                         ['vat', '=ilike', '%' + k]]], {'order': 'id'})
            if not r and ent['nome']:
                r = kw('res.partner', 'search', [base + [['name', '=ilike', ent['nome']]]], {'order': 'id'})
            cache_parc[chave] = r[0] if r else False
        return cache_parc[chave]

    def cliente_por_nif(nif):
        """Cliente pelo NIF confirmado na tabela parceiros_por_subconta.csv (2721x, 2829xx).

        Os clientes vêm da importação do SAF-T: uns com o NIF no campo NIF, outros só na referência
        (MODELO CONTINENTE HIPERMERCADOS S.A), e alguns em duplicado, com e sem prefixo de país
        (Petit Forestier, Newcastle). Fica o que é cliente; em empate, o mais antigo.
        """
        import re
        n = re.sub(r'^[A-Z]{2}(?=[A-Z0-9]{8,}$)', '', nif.replace(' ', '').upper())
        c = kw('res.partner', 'search_read', [[['parent_id', '=', False], '|', ['vat', '=ilike', '%' + n], ['ref', '=ilike', '%' + n]]],
               {'fields': ['customer_rank'], 'order': 'id'})
        return max(c, key=lambda p: (p['customer_rank'] > 0, -p['id']))['id'] if c else False

    def clientes_da_tabela():
        """{subconta: NIF confirmado} da tabela parceiros_por_subconta.csv; vazio se não existir."""
        import csv
        tabela = F.get('parceiros_por_subconta', '')
        if not os.path.exists(tabela):
            return {}
        return {r['subconta']: r['parceiro_confirmado_nif'].strip()
                for r in csv.DictReader(open(tabela, encoding='utf-8-sig'), delimiter=';')
                if (r.get('parceiro_confirmado_nif') or '').strip()}

    def nif_chave(v):
        """NIF sem espaços nem prefixo de país: "PT 500829993", "BE0448826918" e "0448826918" dão o mesmo."""
        import re
        v = re.sub(r'[^A-Z0-9]', '', (v or '').upper())
        v = re.sub(r'^[A-Z]{2}(?=[A-Z0-9]{8,}$)', '', v)
        return v if len(v) >= 8 and v != '999999990' else None

    if args.passo == 'fundir_parceiros':
        # O importador só reconhece um parceiro existente quando o NIF vem escrito igual; com espaço,
        # sem o prefixo do país ou com ele a mais, cria outra ficha. Os movimentos do mesmo parceiro
        # ficavam divididos entre duas fichas (no caso de referência, 56 NIF, com as faturas numa ficha e
        # os pagamentos na outra). Junta-se cada NIF numa só ficha com a fusão de
        # contactos do Odoo, que passa todas as referências, linhas contabilísticas incluídas, para a ficha
        # que fica. A ficha que fica mantém nome, NIF e morada; da outra só vêm os campos que lhe faltem.
        # Fica: a que tem faturas de cliente lançadas (documentos certificados não mudam de parceiro); senão
        # a que já existia antes da importação; senão a que tem mais movimentos. Grupos com duas fichas com
        # faturas de cliente, ou com utilizador ou empresa associados, ficam de fora e listados.
        # NIF errados em fichas que já existiam, corrigidos pela contabilidade (tabela correcoes_nif:
        # nome;nif_atual;nif_novo). Depois da correção as fichas passam a ter o mesmo NIF e juntam-se
        # como as restantes.
        tabela_c = F.get('correcoes_nif', '')
        corrigidas = []
        if os.path.exists(tabela_c):
            for r in csv.DictReader(open(tabela_c, encoding='utf-8-sig'), delimiter=';'):
                ids = kw('res.partner', 'search', [[['parent_id', '=', False], ['vat', '=', r['nif_atual']]]],
                         {'context': {'active_test': False}})
                if ids:
                    w('res.partner', 'write', [ids, {'vat': r['nif_novo']}], {'context': {'no_vat_validation': True}})
                corrigidas.append(f"{r['nome']}: {r['nif_atual']} -> {r['nif_novo']} em {len(ids)} ficha(s)")
        criados = {x['rec_id'] for x in kw('l10n_pt.account.saft.import', 'search_read',
                                          [[['rec_model', '=', 'res.partner'], ['state', '=', 'created']]], {'fields': ['rec_id']})}
        P = kw('res.partner', 'search_read', [[['parent_id', '=', False], ['company_id', 'in', [EMPRESA, False]]]],
               {'fields': ['name', 'vat', 'ref', 'user_ids'], 'context': {'active_test': False}})
        empresas = {c['partner_id'][0] for c in kw('res.company', 'search_read', [[]], {'fields': ['partner_id']})}
        grupos = collections.defaultdict(list)
        for x in P:
            k = nif_chave(x['vat']) or (nif_chave(x['ref']) if x['id'] in criados else None)
            if k:
                grupos[k].append(x)
        grupos = {k: v for k, v in grupos.items() if len(v) > 1}
        ids = [x['id'] for v in grupos.values() for x in v]
        linhas = collections.Counter({g['partner_id'][0]: g['__count'] for g in kw('account.move.line', 'read_group',
                    [[['partner_id', 'in', ids]], ['partner_id'], ['partner_id']], {'lazy': False})}) if ids else collections.Counter()
        faturas = {g['partner_id'][0] for g in kw('account.move', 'read_group',
                    [[['partner_id', 'in', ids], ['move_type', 'in', ['out_invoice', 'out_refund']], ['state', '=', 'posted']],
                     ['partner_id'], ['partner_id']], {'lazy': False})} if ids else set()
        feitos, fora, registo = [], [], []
        for k, v in sorted(grupos.items()):
            fixos = [x for x in v if x['id'] in faturas or x['id'] in empresas or x['user_ids']]
            if len(fixos) > 1:
                fora.append(f"{k}: " + ' | '.join(f"{x['id']} {x['name']}" for x in v) + ' (mais de uma com faturas, utilizador ou empresa)')
                continue
            dst = fixos[0] if fixos else sorted(v, key=lambda x: (x['id'] in criados, -linhas[x['id']], x['id']))[0]
            src = [x for x in v if x['id'] != dst['id']]
            for i in range(0, len(src), 2):  # o assistente aceita até 3 contactos de cada vez
                lote = src[i:i + 2]
                wid = w('base.partner.merge.automatic.wizard', 'create',
                        [{'partner_ids': [(6, 0, [dst['id']] + [x['id'] for x in lote])], 'dst_partner_id': dst['id']}],
                        {'context': {'active_model': 'res.partner', 'active_ids': [dst['id']]}})
                w('base.partner.merge.automatic.wizard', 'action_merge', [[wid]])
            feitos.append(f"{k}: fica {dst['id']} {dst['name']} ({linhas[dst['id']]} linhas); fundidas "
                          + ', '.join(f"{x['id']} {x['name']} ({linhas[x['id']]} linhas)" for x in src))
            registo += [{'nif': k, 'fica_id': dst['id'], 'fica': dst['name'], 'fundida_id': x['id'], 'fundida': x['name'],
                         'linhas_movidas': linhas[x['id']]} for x in src]
        if registo:
            with open(os.path.join(PASTA, 'fusoes_parceiros.csv'), 'w', encoding='utf-8-sig', newline='') as fh:
                wr = csv.DictWriter(fh, fieldnames=list(registo[0]), delimiter=';')
                wr.writeheader()
                wr.writerows(registo)
        diario('Fichas de parceiro duplicadas fundidas (mesmo NIF, escrito de outra forma)',
               corrigidas + [f'{len(feitos)} NIF fundidos, {len(registo)} fichas a menos; detalhe em fusoes_parceiros.csv'] + feitos
               + [f'FORA: {f}' for f in fora])

    elif args.passo == 'taxonomia':
        # Taxonomia SNC das contas (02/10: a declaração "Balanço" da Exo saiu a zeros, só com o resultado
        # líquido). No caso de referência nenhuma conta tinha taxonomia, nem em produção, e as declarações
        # da Exo (Balanço, DR, IES) e o SAF-T de contabilidade somam por ela.
        # 1) Contas finais com movimento ou saldo: a taxonomia das contas do SAF-T que lhes correspondem
        #    (TaxonomyCode, referência S). No ensaio de 02/10 cada uma das 264 tinha uma só.
        # 2) Restantes contas da empresa sem taxonomia: a que a Exo propõe para o código
        #    (l10n_pt_taxonomy_domain), só quando é uma só; as outras ficam listadas.
        import json
        import xml.etree.ElementTree as ET
        ns = {'s': 'urn:OECD:StandardAuditFile-Tax:PT_1.04_01'}
        raiz = ET.parse(SAFT).getroot()
        ref = (raiz.findtext('.//s:TaxonomyReference', 'S', ns) or 'S').strip()
        tx_saft = {}
        for c in raiz.findall('.//s:GeneralLedgerAccounts/s:Account', ns):
            cod_t = c.find('s:TaxonomyCode', ns)
            if cod_t is not None and cod_t.text:
                tx_saft[c.find('s:AccountID', ns).text.strip()] = int(cod_t.text)
        wb = openpyxl.load_workbook(F['plano'], read_only=True)
        mapa = {str(r[0]): str(r[2]) for r in wb['Tabela'].iter_rows(min_row=2, values_only=True) if r[0]}
        por_final = collections.defaultdict(set)
        for leg, cod_t in tx_saft.items():
            if leg in mapa:
                por_final[mapa[leg]].add(cod_t)
        tax_id = {t['code']: t['id'] for t in kw('l10n_pt.account.taxonomy', 'search_read', [[['reference', '=', ref]]],
                                                 {'fields': ['code']})}
        C = {a['code']: a for a in kw('account.account', 'search_read', [[['company_ids', 'in', [EMPRESA]]]],
                                      {'fields': ['code', 'l10n_pt_taxonomy_id', 'l10n_pt_taxonomy_domain']})}
        escrever, conflito, falta_conta = collections.defaultdict(list), [], []
        for final, cods in sorted(por_final.items()):
            if final not in C:
                falta_conta.append(final)
            elif len(cods) > 1:
                conflito.append(f'{final}: {sorted(cods)}')
            elif tax_id.get(next(iter(cods))) and C[final]['l10n_pt_taxonomy_id'] is False:
                escrever[tax_id[next(iter(cods))]].append(C[final]['id'])
        n_saft = sum(len(v) for v in escrever.values())
        ambiguas, unicas = [], 0
        feitas = {i for v in escrever.values() for i in v}
        for cod, a in sorted(C.items()):
            if a['l10n_pt_taxonomy_id'] or a['id'] in feitas or not a['l10n_pt_taxonomy_domain']:
                continue
            cand = kw('l10n_pt.account.taxonomy', 'search', [json.loads(a['l10n_pt_taxonomy_domain'])])
            if len(cand) == 1:
                escrever[cand[0]].append(a['id'])
                unicas += 1
            else:
                ambiguas.append(f'{cod} ({len(cand)} possíveis)')
        for t_id, ids in escrever.items():
            w('account.account', 'write', [ids, {'l10n_pt_taxonomy_id': t_id}])
        diario('Taxonomia SNC das contas',
               [f'referência {ref}; {n_saft} contas pelo SAF-T; {unicas} pela proposta única da Exo',
                f'conflitos no SAF-T (mais de uma taxonomia para a mesma conta final): {len(conflito)}',
                f'contas finais do SAF-T que não existem na base: {len(falta_conta)}',
                f'contas sem taxonomia por terem várias possíveis: {len(ambiguas)}']
               + conflito[:20] + [f'sem conta: {x}' for x in falta_conta[:20]] + [f'várias: {x}' for x in ambiguas[:40]])

    elif args.passo == 'config_fiscal':
        # Configuração fiscal confirmada pela contabilidade (configuração "config_fiscal"):
        # a) Apuramento do IVA: grupos de impostos do IVA com a conta de IVA a pagar e a de IVA a recuperar,
        #    para o fecho em Accounting › Tax Return. A passagem para reembolso pedido continua à mão.
        # b) Retenções na fonte nas subcontas, pelo tipo de rendimento do imposto (ex.: B na 24212, F e R3
        #    na 24214); as restantes ficam na conta-mãe.
        # c) As retenções não entram no apuramento do IVA: a Exo traz as repartições dos impostos RF com
        #    "use in tax closing"; com contas nos grupos do IVA, um trimestre com retenções fazia o fecho
        #    pedir contas para os grupos das retenções ou misturá-las no IVA.
        C = contas()
        feitos = []
        CF = CFG.get('config_fiscal', {})
        pagar, recuperar = CF.get('iva_a_pagar'), CF.get('iva_a_recuperar')
        grupos = kw('account.tax.group', 'search_read', [[['company_id', '=', EMPRESA], ['name', '=like', 'VAT%']]],
                    {'fields': ['name', 'tax_payable_account_id', 'tax_receivable_account_id']})
        if pagar in C and recuperar in C:
            w('account.tax.group', 'write', [[x['id'] for x in grupos],
                                             {'tax_payable_account_id': C[pagar]['id'], 'tax_receivable_account_id': C[recuperar]['id']}])
            feitos.append(f'{len(grupos)} grupos do IVA com {pagar} (a pagar) e {recuperar} (a recuperar)')
        else:
            feitos.append(f'contas {pagar}/{recuperar} não encontradas: grupos do IVA por configurar')
        conta_rf = CF.get('retencoes_por_rendimento', {})
        rf = kw('account.tax', 'search_read', [[['company_id', '=', EMPRESA], ['name', '=like', 'RF%']]],
                {'fields': ['name', 'l10n_pt_wh_income_type'], 'context': {'active_test': False}})
        por_conta = collections.Counter()
        for t in rf:
            rep = kw('account.tax.repartition.line', 'search_read', [[['tax_id', '=', t['id']], ['repartition_type', '=', 'tax']]],
                     {'fields': ['account_id', 'use_in_tax_closing']})
            vals_ids = [r['id'] for r in rep if r['use_in_tax_closing']]
            if vals_ids:
                w('account.tax.repartition.line', 'write', [vals_ids, {'use_in_tax_closing': False}])
            alvo_c = conta_rf.get(t['l10n_pt_wh_income_type'])
            com_conta = [r['id'] for r in rep if r['account_id']]
            if alvo_c and alvo_c in C and com_conta:
                w('account.tax.repartition.line', 'write', [com_conta, {'account_id': C[alvo_c]['id']}])
                por_conta[alvo_c] += 1
            else:
                por_conta[f'{RET_PREFIXO} (sem subconta)' if com_conta else 'sem conta'] += 1
        feitos.append(f'{len(rf)} impostos de retenção fora do apuramento do IVA')
        feitos += [f'impostos de retenção na {k}: {v}' for k, v in sorted(por_conta.items())]
        diario('Configuração fiscal: apuramento do IVA e contas das retenções', feitos)

    elif args.passo == 'extratos':
        # Extratos por conta e por parceiro: a contabilidade consulta em separado clientes (21) e
        # adiantamentos de clientes (218), fornecedores (221) e adiantamentos a fornecedores (228). O Extrato de Parceiro do Odoo só mostra contas a receber e a pagar, e as 278x também o são:
        # aparecem misturadas com a 211 e a 221; as 218 e 228 nem aparecem.
        # a) 278x passam a "não comerciais": o Extrato de Parceiro separa Receivable/Payable (211, 221) de
        #    Non Trade Receivable/Payable (278x) no filtro de tipo de conta.
        # b) Razão Geral com filtro de parceiro: cada conta em separado (211, 218, 221, 228, 278...) só com os
        #    movimentos do parceiro, exportável em PDF. Registo do account_reports: uma atualização desse
        #    módulo (-u) repõe o valor por omissão; nesse caso volta a correr-se este passo.
        contas_nt = kw('account.account', 'search', [[['company_ids', 'in', [EMPRESA]], ['code', '=like', PREFIXO_OUTROS + '%'],
                                                      ['account_type', 'in', ['asset_receivable', 'liability_payable']],
                                                      ['non_trade', '=', False]]])
        if contas_nt:
            w('account.account', 'write', [contas_nt, {'non_trade': True}])
        gl = kw('ir.model.data', 'search_read', [[['module', '=', 'account_reports'], ['name', '=', 'general_ledger_report']]],
                {'fields': ['res_id']})
        if gl:
            w('account.report', 'write', [[gl[0]['res_id']], {'filter_partner': True}])
        diario('Extratos por conta e parceiro',
               [f'{len(contas_nt)} contas 278x marcadas como não comerciais (Extrato de Parceiro: Non Trade)',
                'Razão Geral com filtro de parceiro' if gl else 'Razão Geral não encontrado'])

    elif args.passo == 'parceiros_em_falta':
        # Fornecedores sem entidade no SAF-T (a subconta não tem cliente nem fornecedor associado, por
        # isso não há NIF por onde ligar), indicados pelo cliente (tabela parceiros_em_falta:
        # nome;nif;subconta_saft). Criam-se se ainda não existirem; o passo parceiros associa-os às
        # linhas da subconta e o passo abertura ao saldo de abertura.
        feitos = []
        if not os.path.exists(F.get('parceiros_em_falta', '')):
            diario('Parceiros sem entidade no SAF-T', ['sem tabela parceiros_em_falta: nada a criar'])
            sys.exit(0)
        for r in csv.DictReader(open(F['parceiros_em_falta'], encoding='utf-8-sig'), delimiter=';'):
            ja = kw('res.partner', 'search', [[['name', '=ilike', r['nome']], ['parent_id', '=', False]]], {'limit': 1})
            if ja:
                feitos.append(f"{r['nome']}: já existe (id {ja[0]})")
                continue
            vals = {'name': r['nome'], 'is_company': True, 'supplier_rank': 1}
            if r['nif']:
                vals['vat'] = r['nif']
            pid = w('res.partner', 'create', [vals])
            feitos.append(f"{r['nome']}: criado (id {pid}), subconta {r['subconta_saft']}")
        diario('Parceiros sem entidade no SAF-T, criados pela tabela do cliente', feitos)

    elif args.passo == 'parceiros':
        # Contas a receber/pagar e também outros devedores e credores (278x: Uber, Via Verde, Elavon,
        # IKEA…). Aqui: subconta, depois a entidade do lançamento, e por fim, se houver, o parceiro que o
        # cliente deu à mesma linha numa base de testes (ficheiro linhas_base_testes: lançamento, conta,
        # valores).
        import json
        import re
        por_conta, por_id, trans = saft()

        # 1. Correção: o importador põe em TODAS as linhas o parceiro do
        #    cabeçalho do lançamento. Nos lançamentos de bancos agregados (um por mês, 100 a 135 movimentos
        #    de entidades diferentes) o cabeçalho é uma só entidade e as linhas ficavam todas com ela.
        #    Regras:
        #    a) linha cuja subconta identifica a entidade -> essa entidade, mesmo que já tenha parceiro;
        #    b) restantes linhas de um lançamento agregado (cabeçalho com entidade diferente de alguma
        #       subconta) -> parceiro pela descrição, se houver um único com esse nome; senão, sem parceiro.
        def nif_n(n):  # tira o prefixo do país, também quando o NIF começa por letra (ESB66854530)
            return re.sub(r'^[A-Z]{2}(?=[A-Z0-9]{8,}$)', '', (n or '').replace(' ', '').upper())
        agregados = set()
        for tid, tr in trans.items():
            cab = por_id.get(tr['ent']) if tr['ent'] else None
            if cab and any(por_conta.get(x['conta']) and nif_n(por_conta[x['conta']]['nif']) != nif_n(cab['nif'])
                           for x in tr['linhas']):
                agregados.add(tid)
        T = kw('account.move.line', 'search_read',
               [[['company_id', '=', EMPRESA], ['parent_state', '=', 'draft'], ['date', '<=', DATA_CORTE]]],
               {'fields': ['account_id', 'move_id', 'partner_id', 'name']})
        refs_t = {m['id']: m['ref'] for m in kw('account.move', 'read', [list({l['move_id'][0] for l in T}), ['ref']])}
        nif_parc = {p['id']: nif_n(p['vat']) for p in kw('res.partner', 'read',
                                                           [list({l['partner_id'][0] for l in T if l['partner_id']}), ['vat']])}
        cache_nome = {}

        def por_descricao(desc):
            d = (desc or '').strip()
            if d not in cache_nome:
                r = kw('res.partner', 'search', [[['parent_id', '=', False], ['name', '=ilike', d]]]) if len(d) > 3 else []
                cache_nome[d] = r[0] if len(r) == 1 else False
            return cache_nome[d]
        corr, limpar, corr_cont = collections.defaultdict(list), [], collections.Counter()
        for l in T:
            cod = l['account_id'][1].split(' ')[0]
            ref = refs_t.get(l['move_id'][0])
            atual = l['partner_id'][0] if l['partner_id'] else False
            if cod in por_conta:
                # só se o NIF do parceiro atual (sem prefixo do país) for outro: a mesma empresa aparece
                # às vezes com e sem prefixo, e não se troca um registo certo por outro da mesma empresa
                if atual and nif_parc.get(atual) and nif_parc[atual] == nif_n(por_conta[cod]['nif']):
                    continue
                pid = parceiro(por_conta[cod])
                if pid and pid != atual:
                    corr[pid].append(l['id'])
                    corr_cont['subconta'] += 1
            elif ref in agregados and atual:
                pid = por_descricao(l['name'])
                if pid and pid != atual:
                    corr[pid].append(l['id'])
                    corr_cont['descrição'] += 1
                elif not pid:
                    limpar.append(l['id'])
                    corr_cont[f'sem parceiro ({cod})'] += 1
        for pid, ids in corr.items():
            w('account.move.line', 'write', [ids, {'partner_id': pid}], SEM_SYNC)
        if limpar:
            w('account.move.line', 'write', [limpar, {'partner_id': False}], SEM_SYNC)
        diario('Parceiros corrigidos: o importador tinha posto o parceiro do cabeçalho em todas as linhas',
               [f'{len(agregados)} lançamentos agregados (cabeçalho com entidade diferente das subcontas)',
                f'corrigidas pela subconta: {corr_cont["subconta"]}; pela descrição: {corr_cont["descrição"]}; '
                f'deixadas sem parceiro: {len(limpar)}']
               + [f'{k}: {v}' for k, v in corr_cont.most_common() if k.startswith('sem parceiro')][:20])

        # 2. Linhas de terceiros ainda sem parceiro
        L = kw('account.move.line', 'search_read',
               [[['company_id', '=', EMPRESA], ['parent_state', '=', 'draft'], ['date', '<=', DATA_CORTE],
                 ['partner_id', '=', False], '|', '|', ['account_id.account_type', 'in', ['asset_receivable', 'liability_payable']],
                 ['account_id.code', '=like', PREFIXO_OUTROS + '%'],
                 # retenções na fonte que entram na Modelo 10: só as retenções (crédito) das contas indicadas
                 # (trabalho independente, rendas). O trabalho dependente vai na DMR do programa de salários
                 # e os pagamentos ao Estado não levam NIF
                 '&', ['account_id.code', 'in', RET_MODELO10 or ['-']], ['credit', '>', 0]]],
               {'fields': ['account_id', 'move_id', 'debit', 'credit']})
        refs = {m['id']: m['ref'] for m in kw('account.move', 'read', [list({l['move_id'][0] for l in L}), ['ref']])}
        base_testes = {}
        if os.path.exists(F.get('linhas_base_testes', '')):
            for x in json.load(open(F.get('linhas_base_testes', ''), encoding='utf-8')):
                if x['parceiro']:
                    base_testes[(x['ref'], x['conta'], round(x['debito'], 2), round(x['credito'], 2))] = x['parceiro']
        por_parc = collections.defaultdict(list)
        origem, falha = collections.Counter(), collections.Counter()
        for l in L:
            cod = l['account_id'][1].split(' ')[0]
            ent, via = por_conta.get(cod), 'subconta'
            if not ent and refs.get(l['move_id'][0]) not in agregados:  # num agregado o cabeçalho não serve
                tr = trans.get(refs.get(l['move_id'][0]))
                ent, via = (por_id.get(tr['ent']) if tr and tr['ent'] else None), 'lançamento'
            if not parceiro(ent):
                m = base_testes.get((refs.get(l['move_id'][0]), cod, round(l['debit'], 2), round(l['credit'], 2)))
                if m:
                    ent, via = m, 'base de testes do cliente'
            pid = parceiro(ent)
            if pid:
                por_parc[pid].append(l['id'])
                origem[via] += 1
            else:
                falha['sem entidade no SAF-T' if not ent else f"{ent['nif']} {ent['nome']} não encontrado"] += 1
        # Retenções que não levam NIF: as que não vão na Modelo 10 e os pagamentos ao Estado (débitos). O
        # importador põe-lhes o parceiro do cabeçalho; tira-se.
        sem_nif = kw('account.move.line', 'search', [[['company_id', '=', EMPRESA], ['parent_state', '=', 'draft'],
                                                     ['partner_id', '!=', False], ['account_id.code', '=like', RET_PREFIXO + '%'], '|',
                                                     ['account_id.code', 'not in', RET_MODELO10 or ['-']], ['debit', '>', 0]]]) if RET_MODELO10 else []
        if sem_nif:
            w('account.move.line', 'write', [sem_nif, {'partner_id': False}], SEM_SYNC)
            origem['retenções sem NIF (fora da Modelo 10 e pagamentos): parceiro retirado'] += len(sem_nif)
        # Acréscimos e rendimentos a reconhecer por cliente (2721x, 2829xx): o SAF-T não identifica o
        # cliente nestas linhas, e a subconta junta-se na conta-mãe. O cliente vem da tabela
        # parceiros_por_subconta.csv, só nas linhas com o NIF confirmado (pedido da contabilidade de
        # 01/10: consultar os rendimentos a reconhecer por cliente; correspondência confirmada a 02/10).
        # A subconta é de um só cliente, por isso o cliente confirmado substitui o que a linha tiver.
        tabela_p = F.get('parceiros_por_subconta', '')
        if os.path.exists(tabela_p):
            for r in csv.DictReader(open(tabela_p, encoding='utf-8-sig'), delimiter=';'):
                nif = (r.get('parceiro_confirmado_nif') or '').strip()
                if not nif:
                    falha[f"{r['subconta']} sem cliente confirmado"] += int(r['movimentos'] or 0)
                    continue
                pid = cliente_por_nif(nif)
                ids = kw('account.move.line', 'search', [[['company_id', '=', EMPRESA], ['parent_state', '=', 'draft'],
                                                          ['account_id.code', '=', r['subconta']], ['partner_id', '!=', pid]]])
                if pid and ids:
                    por_parc[pid].extend(ids)
                    origem['tabela de clientes por subconta'] += len(ids)
                elif not pid:
                    falha[f"{r['subconta']}: NIF {nif} não encontrado"] += len(ids)
        # Subcontas sem entidade no SAF-T, com o parceiro criado pelo passo parceiros_em_falta
        tabela_f = F.get('parceiros_em_falta', '')
        if os.path.exists(tabela_f):
            for r in csv.DictReader(open(tabela_f, encoding='utf-8-sig'), delimiter=';'):
                pid = kw('res.partner', 'search', [[['name', '=ilike', r['nome']], ['parent_id', '=', False]]], {'limit': 1})
                ids = kw('account.move.line', 'search', [[['company_id', '=', EMPRESA], ['parent_state', '=', 'draft'],
                                                          ['account_id.code', '=', r['subconta_saft']], ['partner_id', '=', False]]])
                ids = [i for i in ids if not any(i in v for v in por_parc.values())]
                if pid and ids:
                    por_parc[pid[0]].extend(ids)
                    origem['parceiros sem entidade no SAF-T'] += len(ids)
                elif ids:
                    falha[f"{r['subconta_saft']}: {r['nome']} por criar (passo parceiros_em_falta)"] += len(ids)
        for pid, ids in por_parc.items():
            w('account.move.line', 'write', [ids, {'partner_id': pid}], SEM_SYNC)
        diario('Parceiro atribuído às linhas de clientes/fornecedores que a importação deixou sem ele',
               [f'{len(L)} linhas sem parceiro em contas a receber/a pagar e em 278x',
                f'atribuídas pela subconta: {origem["subconta"]}; pela entidade do lançamento: {origem["lançamento"]}; '
                f'pela base de testes do cliente: {origem["base de testes do cliente"]}; '
                f'pela tabela de clientes por subconta (2721x, 2829xx): {origem["tabela de clientes por subconta"]}',
                f'por atribuir: {sum(falha.values())}'] + [f'{k}: {v}' for k, v in falha.most_common(15)])

    elif args.passo == 'faltas':
        # Lançamentos que a importação não criou (ex.: entidade com espaço a mais no ID do MasterFiles):
        # recriados a partir do XML, em rascunho, com o parceiro resolvido pela subconta ou pelo NIF.
        por_conta, por_id, trans = saft()
        C = contas()
        erros = kw('l10n_pt.account.saft.import', 'search_read',
                   [[['configuration_id.name', '=', NOME], ['state', '=', 'error'], ['rec_model', '=', 'account.move']]],
                   {'fields': ['saft_data', 'partner_code', 'error_message']})
        diarios = {j['code']: j['id'] for j in kw('account.journal', 'search_read', [[['company_id', '=', EMPRESA]]],
                                                  {'fields': ['code']})}
        feitos = []
        for e in erros:
            ref = e['saft_data']['ref']
            if kw('account.move', 'search_count', [[['ref', '=', ref], ['company_id', '=', EMPRESA]]]):
                feitos.append(f'{ref}: já existe, não recriado')
                continue
            tr = trans[ref]
            ent = next((v for k, v in por_id.items() if k and k.strip() == (e['partner_code'] or '').strip()), None)
            vals = [(0, 0, {'account_id': C[x['conta']]['id'], 'name': x['desc'] or tr['desc'],
                            'partner_id': parceiro(por_conta.get(x['conta']) or ent),
                            'debit': x['debit'], 'credit': x['credit']}) for x in tr['linhas']]
            jid = diarios.get(tr['journal']) or kw('account.journal', 'search', [[['company_id', '=', EMPRESA], ['type', '=', 'general']]], {'limit': 1})[0]
            mid = w('account.move', 'create', [{'move_type': 'entry', 'journal_id': jid, 'date': tr['date'],
                                                 'ref': ref, 'line_ids': vals}])
            feitos.append(f'{ref} recriado (id {mid}, {len(vals)} linhas, entidade {ent["nome"] if ent else "-"}): {e["error_message"][:90]}')
        diario('Lançamentos que a importação não criou, recriados do XML', feitos or ['nenhum'])

    elif args.passo == 'grelhas':
        # Grelhas de IVA da Exo nas linhas importadas, pela subconta de origem (tabela de
        # preparar_grelhas.py, feita da taxonomia da contabilidade). Tem de correr ANTES de
        # "recodificar" e "juntar": depois a subconta desaparece e com ela o campo da DP e a linha do
        # Anexo L. O tax_tag_ids é um campo simples da linha (não é recalculado quando a conta muda).
        # Os diários de apuramento (configuração "diarios_apuramento") ficam de fora: transferem os saldos no fim do
        # período e anulavam os campos.
        tabela = {r['subconta']: [t.strip() for t in r['grelhas'].split('|') if t.strip()]
                  for r in csv.DictReader(open(F['grelhas'], encoding='utf-8-sig'), delimiter=';')}
        tags = {t['name']: t['id'] for t in kw('account.account.tag', 'search_read',
                                               [[['applicability', '=', 'taxes']]], {'fields': ['name']})}
        faltam = sorted({t for ts in tabela.values() for t in ts if t not in tags})
        if faltam:
            sys.exit(f'Grelhas que não existem nesta base: {faltam}')
        C = contas()
        feitos, sem_conta = collections.Counter(), []
        for sub, nomes in sorted(tabela.items()):
            if sub not in C:
                sem_conta.append(sub)
                continue
            ids = kw('account.move.line', 'search', [[['account_id', '=', C[sub]['id']], ['company_id', '=', EMPRESA],
                                                     ['parent_state', '=', 'draft'], ['date', '<=', DATA_CORTE],
                                                     ['journal_id.code', 'not in', DIARIOS_APUR]]])
            for i in range(0, len(ids), 500):
                w('account.move.line', 'write', [ids[i:i + 500], {'tax_tag_ids': [(6, 0, [tags[n] for n in nomes])]}], SEM_SYNC)
            if ids:
                feitos[f'{sub}: {" | ".join(nomes)}'] = len(ids)
        diario('Grelhas de IVA gravadas nas linhas importadas, pela subconta de origem',
               [f'{sum(feitos.values())} linhas em {len(feitos)} subcontas (diários de apuramento {"/".join(DIARIOS_APUR)} excluídos)']
               + [f'{k}: {n} linhas' for k, n in feitos.items()]
               + [f'{len(sem_conta)} subcontas da tabela sem conta nesta base (sem movimentos no SAF-T)'])

    elif args.passo == 'config_extra':
        # Configuração que o cliente fez numa base de testes e que não existe em produção (ficheiro
        # config_extra: {"impostos": [...], "posicoes_fiscais": [...]}): impostos criados por ele (só os
        # que faltam, pelo nome) e posição fiscal de parceiros. Sem o ficheiro, não faz nada.
        if not os.path.exists(F.get('config_extra', '')):
            diario('Configuração extra do cliente', ['sem ficheiro config_extra: nada a fazer'])
            sys.exit(0)
        import json
        cfg = json.load(open(F['config_extra'], encoding='utf-8'))
        C = contas()
        tags = {t['name']: t['id'] for t in kw('account.account.tag', 'search_read',
                                               [[['applicability', '=', 'taxes']]], {'fields': ['name']})}
        feitos = []
        for t in cfg['impostos']:
            if kw('account.tax', 'search_count', [[['name', '=', t['name']], ['company_id', '=', EMPRESA]]], {'context': {'active_test': False}}):
                feitos.append(f"imposto {t['name']}: já existe, não criado")
                continue
            grupo = kw('account.tax.group', 'search', [[['name', '=', t['tax_group_id']], ['company_id', '=', EMPRESA]]], {'limit': 1})
            pais = kw('res.country', 'search', [[['name', '=', t['country_id']]]], {'limit': 1}) if t.get('country_id') else []
            reps = []
            for r in t['reparticao']:
                faltam = [g for g in r['grelhas'] if g not in tags]
                if faltam or (r['conta'] and r['conta'] not in C):
                    sys.exit(f"imposto {t['name']}: falta conta {r['conta']} ou grelhas {faltam}")
                reps.append((0, 0, {'document_type': r['document_type'], 'repartition_type': r['repartition_type'],
                                    'factor_percent': r['factor_percent'], 'sequence': r['sequence'],
                                    'use_in_tax_closing': r['use_in_tax_closing'],
                                    'account_id': C[r['conta']]['id'] if r['conta'] else False,
                                    'tag_ids': [(6, 0, [tags[g] for g in r['grelhas']])]}))
            vals = {k: t[k] for k in ('name', 'description', 'invoice_label', 'amount', 'amount_type', 'type_tax_use',
                                      'tax_scope', 'price_include_override', 'include_base_amount', 'is_base_affected',
                                      'sequence', 'tax_exigibility', 'invoice_legal_notes', 'analytic', 'active') if k in t}
            vals.update({'company_id': EMPRESA, 'tax_group_id': grupo[0] if grupo else False,
                         'repartition_line_ids': reps})
            if pais:
                vals['country_id'] = pais[0]
            tid = w('account.tax', 'create', [vals])
            feitos.append(f"imposto {t['name']} criado (id {tid}, {len(reps)} linhas de repartição)")
        for p in cfg['posicoes_fiscais']:
            fp = kw('account.fiscal.position', 'search', [[['name', '=', p['posicao']], ['company_id', '=', EMPRESA]]], {'limit': 1})
            pid = parceiro({'nif': p['nif'], 'nome': p['nome']})
            if not fp or not pid:
                feitos.append(f"posição fiscal {p['posicao']} em {p['nome']}: {'posição' if not fp else 'parceiro'} não encontrado")
                continue
            w('res.partner', 'write', [[pid], {'property_account_position_id': fp[0]}])
            feitos.append(f"posição fiscal {p['posicao']} em {p['nome']} (id {pid})")
        diario('Configuração extra do cliente replicada', feitos)

    elif args.passo == 'impostos':
        # A declaração periódica da Exo (ptplus_accounting_tax_statement) só conta uma linha de IVA
        # nos campos "Normal" se estiver ligada à repartição de um imposto; a grelha não chega
        # (ensaio de 30/09: E1 só com grelhas não entrou no C024, E2 com imposto entrou). As bases
        # contam só com a grelha (Anexo L de agosto: L73 com E1 e E2). Aqui liga-se cada linha de
        # IVA importada (contas 243x) à repartição "fatura / imposto" de um imposto da Exo do
        # Continente que tenha a grelha DP da linha, com a taxa dada pela linha do Anexo L ou pelo
        # campo. As grelhas da linha ficam como estão, e os valores também: skip_invoice_sync impede
        # o Odoo de recalcular ou acrescentar linhas. Regularizações (C40/C41) ficam de fora: já contam.
        TAXA_AL = {'L45': 6, 'L48': 6, 'L51': 6, 'L46': 13, 'L49': 13, 'L52': 13, 'L47': 23, 'L50': 23, 'L53': 23}
        TAXA_DP = {'C2': 6, 'C21': 6, 'C6': 13, 'C23': 13, 'C4': 23, 'C22': 23}
        tagn = {t['id']: t['name'] for t in kw('account.account.tag', 'search_read', [[['applicability', '=', 'taxes']]], {'fields': ['name']})}
        taxas = {t['id']: t for t in kw('account.tax', 'search_read', [[['company_id', '=', EMPRESA]]],
                                          {'fields': ['name', 'amount', 'type_tax_use'], 'context': {'active_test': False}})}
        reps = [r for r in kw('account.tax.repartition.line', 'search_read',
                              [[['tax_id.company_id', '=', EMPRESA], ['document_type', '=', 'invoice'], ['repartition_type', '=', 'tax']]],
                              {'fields': ['tax_id', 'tag_ids', 'factor_percent']})
                if not any(x in taxas[r['tax_id'][0]]['name'] for x in ('[AZ]', '[MD]', 'Adiantamento'))]

        def escolher(dp, al):
            cands = [r for r in reps if dp in [tagn[i] for i in r['tag_ids']]]
            campo = dp.split('-')[-1].rstrip(']').split(']')[0]
            taxa = TAXA_AL.get(al.split('-')[-1].rstrip(']')) if al else None
            taxa = taxa or TAXA_DP.get(campo)
            if taxa:
                cands = [r for r in cands if abs(taxas[r['tax_id'][0]]['amount']) == taxa] or cands
            # campos 2, 4 e 6: IVA liquidado nas vendas -> imposto de venda; restantes -> de compra
            uso = 'sale' if campo in ('C2', 'C4', 'C6') else 'purchase'
            cands = [r for r in cands if taxas[r['tax_id'][0]]['type_tax_use'] == uso] or cands
            if al:  # preferir a repartição que também tem a linha do Anexo L
                cands = [r for r in cands if al in [tagn[i] for i in r['tag_ids']]] or cands
            especial = lambda r: sum(x in taxas[r['tax_id'][0]]['name'] for x in ('Imob', 'E-COM', 'EXTRAC', 'PD50', 'DRS', 'CCV', 'EGE', 'VTA', 'Conta OBS'))
            return sorted(cands, key=lambda r: (especial(r), len(r['tag_ids']), r['id']))[0] if cands else None

        base_dom = [['company_id', '=', EMPRESA], ['parent_state', '=', 'draft'], ['date', '<=', DATA_CORTE],
                    ['account_id.code', '=like', '243%'], ['tax_tag_ids', '!=', False],
                    ['journal_id.code', 'not in', DIARIOS_APUR], ['move_id.move_type', '=', 'entry']]
        ja = kw('account.move.line', 'search', [base_dom + [['tax_repartition_line_id', '!=', False]]])
        if ja:  # voltar a correr o passo: desfaz as ligações anteriores e escolhe de novo
            kw('account.move.line', 'write', [ja, {'tax_repartition_line_id': False}],
               {'context': {'skip_invoice_sync': True}}, allow_write=True)
        L = kw('account.move.line', 'search_read', [base_dom], {'fields': ['tax_tag_ids', 'account_id']})
        por_rep, sem, regul = collections.defaultdict(list), collections.Counter(), 0
        for l in L:
            nomes = [tagn[i] for i in l['tax_tag_ids']]
            dp = next((n for n in nomes if n.startswith('IVA DP')), None)
            al = next((n for n in nomes if n.startswith('IVA AL')), None)
            if dp in ('IVA DP [Q06-C40]', 'IVA DP [Q06-C41]'):
                regul += 1
                continue
            r = escolher(dp, al) if dp else None
            if r:
                por_rep[r['id']].append(l['id'])
            else:
                sem[' | '.join(sorted(nomes))] += 1
        for rid, ids in por_rep.items():
            for i in range(0, len(ids), 500):
                kw('account.move.line', 'write', [ids[i:i + 500], {'tax_repartition_line_id': rid}],
                   {'context': {'skip_invoice_sync': True}}, allow_write=True)
        nome_rep = {r['id']: f"{r['tax_id'][1]} ({r['factor_percent']:+.0f}%)" for r in reps}
        diario('Linhas de IVA ligadas à repartição do imposto da Exo (sem mudar valores nem grelhas)',
               [f'{sum(len(v) for v in por_rep.values())} linhas ligadas; {regul} regularizações (C40/C41) deixadas como estão; '
                f'{sum(sem.values())} sem imposto correspondente']
               + [f'{nome_rep[r]}: {len(ids)} linhas' for r, ids in sorted(por_rep.items(), key=lambda x: -len(x[1]))]
               + [f'SEM IMPOSTO: {k}: {n} linhas' for k, n in sem.items()])

        # Bases com grelha da DP (campos 1, 3, 5, 7, 8, 9, 12, 14, 15, 16, 18, 97 a 105): a DP da Exo
        # só as conta com imposto na linha (1.º trimestre de 30/09: IVA ao cêntimo, bases a zero; o
        # Anexo L, sem repartição por região, conta-as só com a grelha). Imposto: o da linha de IVA do
        # mesmo lançamento, se a base desse imposto tiver a grelha DP da linha; senão, um imposto do
        # Continente cuja base tenha essa grelha, com a taxa do campo (1: 6 %, 3: 23 %, 5: 13 %).
        # Só tax_ids: sem sincronização, não se criam nem alteram linhas de IVA.
        CAMPOS_IMPOSTO = {'C2', 'C4', 'C6', 'C11', 'C13', 'C17', 'C19', 'C20', 'C21', 'C22', 'C23', 'C24', 'C40', 'C41', 'C61'}
        TAXA_BASE = {'C1': 6, 'C3': 23, 'C5': 13}
        breps = [r for r in kw('account.tax.repartition.line', 'search_read',
                               [[['tax_id.company_id', '=', EMPRESA], ['document_type', '=', 'invoice'], ['repartition_type', '=', 'base']]],
                               {'fields': ['tax_id', 'tag_ids']})
                 if not any(x in taxas[r['tax_id'][0]]['name'] for x in ('[AZ]', '[MD]', 'Adiantamento'))]
        base_tags = {r['tax_id'][0]: {tagn[i] for i in r['tag_ids']} for r in breps}
        bdom = [['company_id', '=', EMPRESA], ['parent_state', '=', 'draft'], ['date', '<=', DATA_CORTE],
                '!', ['account_id.code', '=like', '243%'], ['tax_tag_ids', '!=', False],
                ['journal_id.code', 'not in', DIARIOS_APUR], ['move_id.move_type', '=', 'entry']]
        ja = kw('account.move.line', 'search', [bdom + [['tax_ids', '!=', False]]])
        if ja:
            kw('account.move.line', 'write', [ja, {'tax_ids': [(5, 0, 0)]}], {'context': {'skip_invoice_sync': True}}, allow_write=True)
        B = kw('account.move.line', 'search_read', [bdom], {'fields': ['tax_tag_ids', 'move_id', 'credit', 'debit', 'account_id']})
        # imposto escolhido pela contabilidade para certas contas (ex.: venda de sucata isenta no campo 8
        # com o motivo M30, e não o que a regra geral escolheria)
        IMPOSTO_POR_CONTA = CFG.get('impostos_por_conta', {})
        nome_taxa = {t['name']: tid for tid, t in taxas.items()}
        iva_mov = collections.defaultdict(set)
        for l in kw('account.move.line', 'search_read', [base_dom + [['tax_line_id', '!=', False]]], {'fields': ['move_id', 'tax_line_id']}):
            iva_mov[l['move_id'][0]].add(l['tax_line_id'][0])
        por_imp, bsem = collections.defaultdict(list), collections.Counter()
        for l in B:
            nomes = [tagn[i] for i in l['tax_tag_ids']]
            dps = [n for n in nomes if n.startswith('IVA DP') and n.split('-C')[-1].split(']')[0] not in {c[1:] for c in CAMPOS_IMPOSTO}]
            if not dps:
                continue  # base só do Anexo L: conta sem imposto
            dp = dps[0]
            campo = 'C' + dp.split('-C')[-1].split(']')[0]
            doMov = [t for t in iva_mov.get(l['move_id'][0], ()) if dp in base_tags.get(t, ())]
            uso = 'sale' if dp.endswith('VND') or campo in ('C7', 'C8', 'C9') else 'purchase'
            conta_l = l['account_id'][1].split(' ')[0] if l.get('account_id') else ''
            if conta_l in IMPOSTO_POR_CONTA and IMPOSTO_POR_CONTA[conta_l] in nome_taxa:
                tid = nome_taxa[IMPOSTO_POR_CONTA[conta_l]]
            elif doMov:
                tid = sorted(doMov)[0]
            else:
                c = [t for t, tg in base_tags.items() if dp in tg and taxas[t]['type_tax_use'] == uso]
                if TAXA_BASE.get(campo):
                    c = [t for t in c if abs(taxas[t]['amount']) == TAXA_BASE[campo]] or c
                al = next((n for n in nomes if n.startswith('IVA AL')), None)
                if al:
                    c = [t for t in c if al in base_tags[t]] or c
                especial = lambda t: sum(x in taxas[t]['name'] for x in ('Imob', 'E-COM', 'EXTRAC', 'PD50', 'Conta OBS', 'Artº'))
                tid = sorted(c, key=lambda t: (especial(t), len(base_tags[t]), t))[0] if c else None
            if tid:
                por_imp[tid].append(l['id'])
            else:
                bsem[' | '.join(sorted(nomes))] += 1
        for tid, ids in por_imp.items():
            for i in range(0, len(ids), 500):
                kw('account.move.line', 'write', [ids[i:i + 500], {'tax_ids': [(6, 0, [tid])]}],
                   {'context': {'skip_invoice_sync': True}}, allow_write=True)
        diario('Bases com grelha da DP ligadas ao imposto da Exo (tax_ids, sem sincronizar)',
               [f'{sum(len(v) for v in por_imp.values())} bases ligadas; {sum(bsem.values())} sem imposto correspondente']
               + [f'{taxas[t]["name"]}: {len(ids)} linhas' for t, ids in sorted(por_imp.items(), key=lambda x: -len(x[1]))]
               + [f'SEM IMPOSTO: {k}: {n} linhas' for k, n in bsem.items()])

    elif args.passo == 'rascunho':
        # Repõe em rascunho os lançamentos até 31/07 (ensaio: para acrescentar um passo a uma base
        # já lançada). As datas e os números ficam; exige o bloqueio abaixo da data de abertura.
        ids = kw('account.move', 'search', [[['company_id', '=', EMPRESA], ['state', '=', 'posted'], ['date', '<=', DATA_CORTE]]])
        for i in range(0, len(ids), 100):
            w('account.move', 'button_draft', [ids[i:i + 100]], {'context': {'skip_invoice_sync': True}})
        diario('Lançamentos até 31/07 repostos em rascunho', [f'{len(ids)} lançamentos'])

    elif args.passo == 'resumo_grelhas':
        # Leitura: soma por grelha e por período (1.º trimestre, 2.º trimestre, julho), para as
        # contabilistas compararem com as declarações entregues. O sinal é o do saldo da linha
        # (débito - crédito): IVA liquidado e bases de vendas saem negativos.
        periodos = [tuple(p) for p in CFG['periodos_resumo']]
        tags = kw('account.account.tag', 'search_read', [[['applicability', '=', 'taxes'], ['name', '=like', 'IVA %']]],
                  {'fields': ['name']})
        linhas = []
        for t in sorted(tags, key=lambda t: t['name']):
            vals = []
            for _, d0, d1 in periodos:
                g = kw('account.move.line', 'read_group', [[['company_id', '=', EMPRESA], ['tax_tag_ids', 'in', [t['id']]],
                                                             ['date', '>=', d0], ['date', '<=', d1]], ['balance:sum'], []],
                       {'lazy': False})
                vals.append(round(g[0]['balance'] or 0, 2) if g else 0)
            if any(vals):
                linhas.append(f"{t['name']}: " + ' | '.join(f'{p} {v:,.2f}' for (p, *_), v in zip(periodos, vals)))
        diario('Resumo por grelha e período (saldo D-C das linhas com a grelha)', linhas or ['nenhuma linha com grelha'])

    elif args.passo == 'recodificar':
        C = contas()
        feitos, notas = [], []
        for de, para, nome, _ in folha('Recodificar'):
            de, para = str(de), str(para)
            if para in C:
                notas.append(f'{de} -> {para}: {para} já existe, fica para a junção')
                continue
            if de not in C:
                notas.append(f'{de} -> {para}: {de} não existe (sem movimentos no ficheiro?)')
                continue
            w('account.account', 'write', [[C[de]['id']], {'code': para}])
            feitos.append(f'{de} -> {para} ({nome}, {n_linhas(C[de]["id"])} linhas)')
        diario('Recodificações um-para-um', [f'{len(feitos)} contas recodificadas'] + feitos + notas)

    elif args.passo == 'juntar':
        # O assistente "Juntar contas" do Odoo 19 não serve: só junta a mesma conta entre empresas
        # diferentes (_apply_different_companies_constraint em account_merge_wizard.py) e recusa
        # contas da mesma empresa ("Belongs to the same company as ..."). Como a importação deixa os
        # lançamentos em rascunho, a junção faz-se mudando a conta das linhas para a conta final (o
        # parceiro fica na linha) e apagando depois a subconta vazia, ou arquivando-a se o Odoo recusar.
        C = contas()
        por_final = collections.defaultdict(list)
        for final, fonte, *_ in folha('Juntar'):
            por_final[str(final)].append(str(fonte))
        feitos, notas = [], []
        for f, fontes in sorted(por_final.items()):
            fontes = [s for s in fontes if s in C and s != f]
            if not fontes:
                continue
            if f not in C:  # conta final nova: a primeira fonte passa a ser ela
                w('account.account', 'write', [[C[fontes[0]]['id']], {'code': f}])
                notas.append(f'{f} não existia: {fontes[0]} recodificada para {f}')
                C = contas()
                fontes = fontes[1:]
                if not fontes:
                    continue
            alvo = C[f]
            movidas, apagadas, arquivadas, com_posted = 0, 0, [], []
            for s in fontes:
                sid = C[s]['id']
                ids = kw('account.move.line', 'search', [[['account_id', '=', sid], ['parent_state', '=', 'draft']]])
                if ids:
                    w('account.move.line', 'write', [ids, {'account_id': alvo['id']}], SEM_SYNC)
                    movidas += len(ids)
                if n_linhas(sid):
                    com_posted.append(s)
                    continue
                try:
                    w('account.account', 'unlink', [[sid]])
                    apagadas += 1
                except Exception as erro:  # referência noutro registo: arquiva
                    w('account.account', 'write', [[sid], {'active': False}])
                    arquivadas.append(f'{s} ({str(erro)[-80:]})')
            feitos.append(f'{f} <- {len(fontes)} subcontas | {movidas} linhas movidas | {apagadas} apagadas'
                          + (f' | arquivadas: {arquivadas}' if arquivadas else '')
                          + (f' | COM LINHAS NÃO RASCUNHO: {com_posted}' if com_posted else ''))
            C = contas()
        diario('Junções: linhas movidas para a conta final e subcontas vazias apagadas',
               [f'{len(feitos)} contas finais'] + feitos + notas)

    elif args.passo == 'reclassificar':
        C = contas()
        feitos = []
        for cert, para, nome, *_ in folha('Reclassificar'):
            cert, para = str(cert), str(para)
            if para not in C:
                base = C[cert]
                w('account.account', 'create', [{'code': para, 'name': CFG.get('nomes_contas_novas', {}).get(para) or nome.strip().title(),
                                                 'account_type': base['account_type'], 'reconcile': base['reconcile'],
                                                 'company_ids': [(6, 0, [EMPRESA])]}])
                C = contas()
                feitos.append(f'{para} criada (tipo de {cert})')
            ids = kw('account.move.line', 'search', [[['account_id', '=', C[cert]['id']], ['parent_state', '=', 'draft'],
                                                     ['date', '<=', DATA_CORTE], ['company_id', '=', EMPRESA]]])
            if ids:
                w('account.move.line', 'write', [ids, {'account_id': C[para]['id']}], SEM_SYNC)
            feitos.append(f'{cert} -> {para}: {len(ids)} linhas importadas passaram de conta')
        diario('Reclassificações (mesmo código, outro sentido)', feitos)

    elif args.passo == 'abertura':
        C = contas()
        linhas = list(csv.DictReader(open(F['abertura'], encoding='utf-8-sig'),
                                     delimiter=';'))
        emp = kw('res.company', 'read', [[EMPRESA], ['account_opening_journal_id']])[0]
        diario_id = emp['account_opening_journal_id'][0] if emp['account_opening_journal_id'] else \
            kw('account.journal', 'search', [[['company_id', '=', EMPRESA], ['type', '=', 'general']]], {'limit': 1})[0]
        falta_conta, falta_parc, vals = set(), [], []
        tabela_cli = clientes_da_tabela()
        tabela_falta = {}
        if os.path.exists(F.get('parceiros_em_falta', '')):
            tabela_falta = {r['subconta_saft']: r['nome'] for r in csv.DictReader(
                open(F.get('parceiros_em_falta', ''), encoding='utf-8-sig'), delimiter=';')}
        for l in linhas:
            if l['conta'] not in C:
                falta_conta.add(l['conta'])
                continue
            pid = False
            if not l['parceiro_nif'] and l['conta_saft'] in tabela_cli:
                # 2721x e 2829xx: o cliente confirmado pela contabilidade (02/10), como nos movimentos
                pid = cliente_por_nif(tabela_cli[l['conta_saft']])
                if not pid:
                    falta_parc.append(f"{l['conta_saft']}: NIF {tabela_cli[l['conta_saft']]} não encontrado")
            elif not l['parceiro_nif'] and l['conta_saft'] in tabela_falta:
                # subconta sem entidade no SAF-T, parceiro criado pelo passo parceiros_em_falta (02/10)
                r = kw('res.partner', 'search', [[['name', '=ilike', tabela_falta[l['conta_saft']]], ['parent_id', '=', False]]],
                       {'limit': 1})
                pid = r[0] if r else False
                if not pid:
                    falta_parc.append(f"{l['conta_saft']}: {tabela_falta[l['conta_saft']]} por criar")
            elif l['parceiro_nif']:
                # A mesma procura do passo parceiros: NIF com e sem prefixo, entidades principais (depois de
                # fundir_parceiros a ficha que fica pode ter o NIF sem o prefixo que a tabela traz)
                pid = parceiro({'nif': l['parceiro_nif'], 'nome': l['parceiro_nome']})
                if not pid:
                    falta_parc.append(f"{l['parceiro_nif']} {l['parceiro_nome']}")
            vals.append((0, 0, {'account_id': C[l['conta']]['id'], 'partner_id': pid,
                                'name': f"Saldo de abertura {l['conta_saft']}",
                                'debit': float(l['debito']), 'credit': float(l['credito'])}))
        if falta_conta:
            sys.exit(f'Contas finais em falta, correr as junções primeiro: {sorted(falta_conta)}')
        mid = w('account.move', 'create', [{'move_type': 'entry', 'journal_id': diario_id, 'date': DATA_ABERTURA,
                                             'ref': REF_ABERTURA, 'line_ids': vals}])
        tot = kw('account.move', 'read', [[mid], ['amount_total', 'state']])[0]
        diario('Lançamento de abertura com parceiro (rascunho)',
               [f'lançamento id {mid}, {len(vals)} linhas, total {tot["amount_total"]:,.2f}, {tot["state"]}',
                f'linhas de terceiros sem parceiro encontrado: {len(falta_parc)}'] + falta_parc[:20])

    elif args.passo == 'limpar':
        # Plano final = contas certificadas + contas finais do plano (com as designações finais). Tudo o
        # que a importação criou fora disso e não tem movimentos sai; os grupos do legado também (a
        # estrutura de grupos é a certificada). As contas finais novas ficam com a designação final.
        # Designações finais: tabela nomes_finais (codigo;nome;renomear) ou, configuração
        # "nomes_finais_excel", uma coluna do Excel do plano com "<código> <nome> [(nova|renomear)]".
        import json
        import re
        certificadas = {a['code'] for a in json.load(open(F['contas_certificadas'], encoding='utf-8-sig'))}
        nome_final, renomear = {}, set()
        NX = CFG.get('nomes_finais_excel')
        if NX:
            wb = openpyxl.load_workbook(NX['ficheiro'], read_only=True, data_only=True)
            col = NX['coluna'] - 1
            for r in wb[NX['folha']].iter_rows(min_row=2, values_only=True):
                m = re.match(r'(\d+) (.+?)( \((nova|renomear)\))?$', str(r[col] or ''))
                if m and not str(r[col]).startswith('grupo'):
                    nome_final[m.group(1)] = m.group(2).strip()
                    if m.group(4) == 'renomear':
                        renomear.add(m.group(1))
        elif os.path.exists(F.get('nomes_finais', '')):
            for r in csv.DictReader(open(F['nomes_finais'], encoding='utf-8-sig'), delimiter=';'):
                nome_final[r['codigo']] = r['nome'].strip()
                if (r.get('renomear') or '').strip().lower() in ('1', 'sim', 'true', 'x'):
                    renomear.add(r['codigo'])
        C = contas()
        usadas = {x['account_id'][0] for x in kw('account.move.line', 'read_group',
                                                 [[['company_id', '=', EMPRESA]], ['account_id'], ['account_id']],
                                                 {'lazy': False})}
        renomeadas = []
        for cod, nome in nome_final.items():
            if cod in C and (cod not in certificadas or cod in renomear) and C[cod]['name'].strip() != nome:
                w('account.account', 'write', [[C[cod]['id']], {'name': nome}])
                renomeadas.append(f'{cod}: "{C[cod]["name"].strip()}" -> "{nome}"')
        finais_plano = {str(r[2]) for r in folha('Tabela')}  # inclui as contas próprias sem movimento
        sair = [a for c, a in C.items() if c not in certificadas and c not in nome_final
                and c not in finais_plano and a['id'] not in usadas]
        apagadas, arquivadas = 0, []
        for i in range(0, len(sair), 100):
            lote = [a['id'] for a in sair[i:i + 100]]
            try:
                w('account.account', 'unlink', [lote])
                apagadas += len(lote)
            except Exception:
                for a in sair[i:i + 100]:
                    try:
                        w('account.account', 'unlink', [[a['id']]])
                        apagadas += 1
                    except Exception as erro:
                        w('account.account', 'write', [[a['id']], {'active': False}])
                        arquivadas.append(f"{a['code']} ({str(erro)[-70:]})")
        grupos = kw('account.group', 'search', [[['company_id', '=', EMPRESA], ['create_date', '>=', datetime.date.today().isoformat()]]])
        if grupos:
            w('account.group', 'unlink', [grupos])
        diario('Limpeza do plano: designações finais, contas e grupos do legado',
               [f'{len(renomeadas)} contas finais novas com a designação da coluna E'] + renomeadas
               + [f'{apagadas} contas do legado sem movimentos e fora do plano final apagadas',
                  f'{len(arquivadas)} arquivadas por não se poderem apagar'] + arquivadas[:20]
               + [f'{len(grupos)} grupos criados pela importação apagados'])

    elif args.passo == 'lancar':
        # Lança os lançamentos importados (e a abertura), em lotes; regista os que o Odoo recusar.
        # ATENÇÃO: com uma data de bloqueio posterior à dos lançamentos, o Odoo MUDA-LHES A DATA para o
        # primeiro dia aberto ao lançar. Por isso recusa correr com bloqueio igual ou posterior à data de abertura.
        emp = kw('res.company', 'read', [[EMPRESA], ['fiscalyear_lock_date', 'hard_lock_date']])[0]
        for campo in ('fiscalyear_lock_date', 'hard_lock_date'):
            if emp[campo] and emp[campo] >= DATA_ABERTURA:
                sys.exit(f'Bloqueio {campo} = {emp[campo]}: correr "bloqueio" primeiro, ou as datas mudam ao lançar.')
        ids = kw('account.move', 'search', [[['company_id', '=', EMPRESA], ['state', '=', 'draft'],
                                              ['date', '<=', DATA_CORTE]]], {'order': 'date, id'})
        ok, recusas = 0, collections.Counter()
        exemplos = {}
        for i in range(0, len(ids), 100):
            lote = ids[i:i + 100]
            try:
                w('account.move', 'action_post', [lote], {'context': {'skip_invoice_sync': True}})
                ok += len(lote)
            except Exception:
                for mid in lote:
                    try:
                        w('account.move', 'action_post', [[mid]], {'context': {'skip_invoice_sync': True}})
                        ok += 1
                    except Exception as erro:
                        msg = str(erro).split('\n')[-2][-160:] if '\n' in str(erro) else str(erro)[-160:]
                        recusas[msg] += 1
                        exemplos.setdefault(msg, mid)
        diario('Lançamento (action_post) dos movimentos importados e da abertura',
               [f'{len(ids)} em rascunho até 31/07; lançados {ok}; recusados {sum(recusas.values())}']
               + [f'{n}x (ex. id {exemplos[m]}): {m}' for m, n in recusas.most_common(10)])

    elif args.passo == 'natureza':
        # O ptplus só lança linhas de capital próprio (classe 5) com a natureza da alteração
        # (l10n_pt_equity_change) preenchida, para a IES quadro 04-A. A natureza é decisão da
        # contabilidade (configuração "natureza"). As opções a1-a7 são as alterações reconhecidas no
        # capital próprio (1.1-1.7) e d1-d5 as operações com detentores de capital (2.1-2.5); ex.:
        # 1.7 = a7, 2.5 = d5.
        NATUREZA = CFG['natureza']['padrao']
        NATUREZA_INTERNA = CFG['natureza']['abertura_e_resultado']
        L = kw('account.move.line', 'search_read',
               [[['company_id', '=', EMPRESA], ['parent_state', '=', 'draft'], ['date', '<=', DATA_CORTE],
                 ['account_id.code', '=like', '5%'], ['l10n_pt_equity_change', '=', False]]],
               {'fields': ['move_id', 'account_id']})
        moves = kw('account.move', 'read', [list({l['move_id'][0] for l in L}), ['ref', 'line_ids']]) if L else []
        contas_mov = {m['id']: {x['account_id'][1].split(' ')[0] for x in kw('account.move.line', 'read', [m['line_ids'], ['account_id']])}
                      for m in moves}
        feitos = collections.Counter()
        for m in moves:
            abertura = (m['ref'] or '') == REF_ABERTURA
            resultado = CFG['natureza'].get('conta_resultado', '818') in contas_mov[m['id']]
            nat = NATUREZA_INTERNA if abertura or resultado else NATUREZA
            ids = [l['id'] for l in L if l['move_id'][0] == m['id']]
            w('account.move.line', 'write', [ids, {'l10n_pt_equity_change': nat}], SEM_SYNC)
            feitos[f"{m['ref']}: {nat} em {len(ids)} linhas"] += 1
        diario(f'Natureza da alteração no capital próprio ({NATUREZA_INTERNA} na abertura e na aplicação do resultado, {NATUREZA} nas restantes)',
               list(feitos) or ['nada a preencher'])

    elif args.passo == 'bloqueio':
        # Baixa o bloqueio global para a véspera da abertura (defeito) ou põe-no na data indicada. Em
        # produção, baixar o bloqueio exige autorização do cliente. No fim da passagem o bloqueio fica na
        # data de corte (--data corte): a contabilidade lança os meses seguintes e a sincronização
        # bancária começa no dia a seguir ao bloqueio.
        antes = kw('res.company', 'read', [[EMPRESA], ['fiscalyear_lock_date', 'hard_lock_date']])[0]
        if antes['hard_lock_date']:
            sys.exit(f"Bloqueio definitivo (hard) em {antes['hard_lock_date']}: não se pode baixar.")
        # O bloqueio que o cliente já tinha guarda-se na primeira vez que se baixa; "--data repor" volta a
        # pô-lo, nunca abaixo da data de corte. Só se usa se o cliente quiser manter um bloqueio posterior
        # ao corte (impede a contabilidade de lançar esses meses).
        import json
        guardado = os.path.join(PASTA, f'bloqueio_anterior_{g.alvo.database}.json')
        if args.data == 'repor':
            anterior = json.load(open(guardado, encoding='utf-8'))['fiscalyear_lock_date'] if os.path.exists(guardado) else None
            args.data = max(anterior or DATA_CORTE, DATA_CORTE)
        elif not os.path.exists(guardado) and antes['fiscalyear_lock_date'] and antes['fiscalyear_lock_date'] > args.data:
            json.dump({'fiscalyear_lock_date': antes['fiscalyear_lock_date']}, open(guardado, 'w', encoding='utf-8'))
        if args.data > VESPERA and kw('account.move', 'search_count', [[['company_id', '=', EMPRESA], ['state', '=', 'draft'],
                                                                       ['date', '<=', args.data], ['date', '>=', DATA_ABERTURA]]]):
            sys.exit(f'Há lançamentos em rascunho até {args.data}: lançar primeiro, ou ficam presos atrás do bloqueio.')
        w('res.company', 'write', [[EMPRESA], {'fiscalyear_lock_date': args.data}])
        depois = kw('res.company', 'read', [[EMPRESA], ['fiscalyear_lock_date']])[0]
        diario(f'Bloqueio global definido em {args.data}',
               [f"fiscalyear_lock_date {antes['fiscalyear_lock_date']} -> {depois['fiscalyear_lock_date']}"])

    elif args.passo == 'repor_datas':
        # Recuperação: se os lançamentos foram lançados com o bloqueio ativo e ficaram com a data do
        # dia, repõe-os em rascunho e devolve-lhes a data do SAF-T (GLPostingDate), pelo TransactionID.
        _, _, trans = saft()
        M = kw('account.move', 'search_read', [[['company_id', '=', EMPRESA], ['state', '=', 'posted']]],
               {'fields': ['ref', 'date']})
        certo = {}
        for m in M:
            if m['ref'] in trans:
                certo[m['id']] = trans[m['ref']]['date']
            elif (m['ref'] or '') == REF_ABERTURA:
                certo[m['id']] = DATA_ABERTURA
        errados = [mid for mid, d in certo.items() if next(x['date'] for x in M if x['id'] == mid) != d]
        for i in range(0, len(errados), 100):
            w('account.move', 'button_draft', [errados[i:i + 100]])
        por_data = collections.defaultdict(list)
        for mid in errados:
            por_data[certo[mid]].append(mid)
        for d, ids in por_data.items():
            # o número atribuído no primeiro lançamento é do mês errado: limpa-se ('/') para o Odoo
            # voltar a numerar pela data certa quando se lançar outra vez
            w('account.move', 'write', [ids, {'date': d, 'name': '/'}], SEM_SYNC)
        diario('Datas repostas a partir do SAF-T (recuperação)',
               [f'{len(errados)} lançamentos com data errada repostos em rascunho e com a data do SAF-T',
                f'{len(por_data)} datas distintas'])
