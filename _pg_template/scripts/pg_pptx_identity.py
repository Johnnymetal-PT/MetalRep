#!/usr/bin/env python3
"""
Identidade visual Parâmetro Global para decks PowerPoint (python-pptx).

FONTE DE VERDADE
    doc_templates/assets/Template_Parametro_PowerPoint_Completo 1.pptx

Toda a paleta, tipografia e geometria deste módulo foi extraída desse ficheiro
(feito à mão no PowerPoint pela Parâmetro Global) e não reinventada. O deck base
`assets/pg_base_deck.pptx` é esse mesmo ficheiro sem slides, para que o tema, os
masters e os tipos de letra venham da marca e não do template genérico do Office.

REGRA DE UTILIZAÇÃO
    Este módulo é a única fonte de identidade para geração de decks. Os templates
    executáveis (Ponto de Situação, Orçamento, ...) compõem slides a partir dos
    componentes daqui e nunca redesenham cores, margens ou cabeçalhos por sua conta.
    É isto que garante que dois decks gerados em projetos diferentes saem iguais.

Dependências:
    pip install python-pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor

# A paleta vem toda de pg_brand_palette.py, fonte unica para todos os formatos.
# Ver la a prova de qual e a cor da marca, e porque havia duas.
import sys as _sys
from pathlib import Path as _Path
_sys.path.insert(0, str(_Path(__file__).resolve().parent))
from pg_brand_palette import (  # noqa: E402
    ACENTO, TITULO_CAPA, TINTA, BARRA_ESCURA, CORPO, SUAVE, para_rgb,
)

from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

# --------------------------------------------------------------------------
# Assets da marca
# --------------------------------------------------------------------------

ASSETS_DIR = Path(__file__).resolve().parent / "assets"
BASE_DECK = ASSETS_DIR / "pg_base_deck.pptx"
LOGO_WORDMARK = ASSETS_DIR / "pg_logo_wordmark.png"  # 940x190 — capa
LOGO_MARK = ASSETS_DIR / "pg_logo_mark.png"          # 57x57 — cabeçalho de conteúdo

LAYOUT_COVER = "Blank"
LAYOUT_CONTENT = "PARAMETRO_CONTENT_MASTER"

# --------------------------------------------------------------------------
# Paleta — valores exactos do ficheiro da marca
# --------------------------------------------------------------------------

# Valores identicos aos que este modulo tinha: a paleta do PowerPoint e a da marca, e
# foi dela que pg_brand_palette.py partiu. Nada muda no aspeto dos decks.
COVER_TITLE = RGBColor(*para_rgb(TITULO_CAPA))  # azul do título da capa
ACCENT = RGBColor(*para_rgb(ACENTO))            # azul de acento (nº página, numeração, barras)
INK = RGBColor(*para_rgb(TINTA))                # texto principal e títulos
BAR_DARK = RGBColor(*para_rgb(BARRA_ESCURA))    # barra inferior da capa
BODY = RGBColor(*para_rgb(CORPO))               # texto corrente da capa
MUTED = RGBColor(0x64, 0x74, 0x8B)         # subtítulos, rodapé, legendas
CARD_BG = RGBColor(0xFF, 0xFF, 0xFF)       # fundo de cartão
BAND = RGBColor(0xF5, 0xF9, 0xFC)          # banda de cabeçalho de tabela / faixa final
RULE = RGBColor(0xE7, 0xEE, 0xF5)          # linhas e contornos
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# acentos de estado (slide de gráfico do ficheiro da marca)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
RED = RGBColor(0xEF, 0x44, 0x44)

GREEN_BG = RGBColor(0xDC, 0xFC, 0xE7)
ORANGE_BG = RGBColor(0xFF, 0xED, 0xD5)
BLUE_BG = RGBColor(0xE0, 0xF2, 0xFE)
RED_BG = RGBColor(0xFE, 0xE2, 0xE2)

FONT_HEAD = "Aptos Display"
FONT_BODY = "Aptos"

# --------------------------------------------------------------------------
# Geometria (cm) — valores exactos do ficheiro da marca
# --------------------------------------------------------------------------

SLIDE_W = 33.867
SLIDE_H = 19.05

CONTENT_L = 1.83     # margem esquerda dos blocos de conteúdo
CONTENT_W = 30.23    # largura útil dos blocos de conteúdo
CONTENT_TOP = 3.86   # topo da área de conteúdo (abaixo da régua do cabeçalho)
CONTENT_BOTTOM = 17.10

_RULE_X, _RULE_W = 1.07, 31.37


def _cm(v):
    return Cm(v)


# Calibração de texto (medida sobre render real das fontes Aptos a 1600px de largura):
# ~62 caracteres por cm de largura, dividido pelo corpo em pt; altura de linha ~0,052 cm/pt.
#
# A capacidade bruta medida é ~67, mas a quebra em fronteira de palavra consome até 10%
# e é imprevisível — o mesmo número de caracteres tanto cabe numa linha como salta para
# duas, consoante onde caem os espaços. O valor está deliberadamente abaixo da capacidade
# real: a estimativa peca por excesso de altura, o que produz folga a mais em vez de texto
# colado ou cortado. Subir este valor volta a colar blocos (testado a 65).
_CHARS_CM_PT = 62.0
_ALTURA_LINHA_PT = 0.052


def estimar_linhas(texto, largura_cm, size):
    """Nº de linhas que um texto ocupa numa caixa de `largura_cm` ao corpo `size`."""
    por_linha = max(12, int(largura_cm * _CHARS_CM_PT / size))
    return max(1, -(-len(str(texto)) // por_linha))


def altura_linha(size):
    """Altura de uma linha de texto, em cm, para o corpo `size`."""
    return size * _ALTURA_LINHA_PT


class _Slide:
    """Wrapper de um slide de conteúdo com os componentes da identidade."""

    def __init__(self, slide, deck):
        self._s = slide
        self._deck = deck

    # -- primitivas -------------------------------------------------------

    def text(self, txt, x, y, w, h, size=11, bold=False, color=INK,
             font=FONT_BODY, align=None, wrap=True, anchor=None, line_spacing=None):
        box = self._s.shapes.add_textbox(_cm(x), _cm(y), _cm(w), _cm(h))
        tf = box.text_frame
        tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if anchor is not None:
            tf.vertical_anchor = anchor
        for i, line in enumerate(str(txt).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = font
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = color
            if align is not None:
                p.alignment = align
            if line_spacing is not None:
                p.line_spacing = line_spacing
        return box

    def rect(self, x, y, w, h, fill=None, line=None, line_pt=1.0):
        sh = self._s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                      _cm(x), _cm(y), _cm(w), _cm(h))
        sh.shadow.inherit = False
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(line_pt)
        return sh

    def hline(self, x, y, w, color=RULE, pt=1.0):
        return self.rect(x, y, w, 0.0, fill=None, line=color, line_pt=pt)

    def card(self, x, y, w, h):
        """Cartão branco com contorno — bloco de conteúdo base da marca."""
        return self.rect(x, y, w, h, fill=CARD_BG, line=RULE, line_pt=1.2)

    def section_label(self, txt, x, y, w=8.0):
        """Etiqueta de secção a azul (9pt bold), como 'Notas / Destaques'."""
        return self.text(txt, x, y, w, 0.64, size=9, bold=True, color=ACCENT)

    def bullets(self, items, x, y, w, size=11, color=INK, gap=0.20):
        """
        Bullets com altura adaptativa: um item que ocupe duas linhas empurra o
        seguinte para baixo, em vez de lhe ficar por cima. Devolve o Y final.
        """
        yy = y
        for it in items:
            # medir o texto que é mesmo desenhado, com o marcador: um item que
            # enche a linha a dois caracteres do fim era contado como uma linha
            # e desenhado em duas, ficando por cima do item seguinte
            n = estimar_linhas(f"• {it}", w, size)
            altura = altura_linha(size) * n
            self.text(f"• {it}", x, yy, w, altura + 0.06, size=size, color=color,
                      line_spacing=1.1)
            yy += altura + gap
        return yy

    # -- componentes compostos -------------------------------------------

    def status_badge(self, text, x, y, w=3.6, h=0.68, status="ok"):
        """
        Pílula/Badge visual de estado para dashboards.
        status: 'ok' (verde), 'warning' (laranja), 'alert' (vermelho), 'info' (azul).
        """
        bg_map = {"ok": GREEN_BG, "warning": ORANGE_BG, "alert": RED_BG, "info": BLUE_BG}
        fg_map = {"ok": GREEN, "warning": ORANGE, "alert": RED, "info": ACCENT}
        bg = bg_map.get(status, BLUE_BG)
        fg = fg_map.get(status, ACCENT)
        self.rect(x, y, w, h, fill=bg, line=None)
        self.text(text, x, y + 0.12, w, h - 0.24, size=8, bold=True, color=fg,
                  align=PP_ALIGN.CENTER, wrap=False)
        return x + w

    def progress_bar(self, x, y, w, h=0.38, pct=0, fill=ACCENT, bg=BAND, show_label=True, label_right=True,
                     label=None):
        """
        Barra de progresso visual com fundo e preenchimento proporcional.
        pct: valor de 0 a 100.
        `label` substitui a percentagem pelo texto dado — serve para barras que
        comparam grandezas reais (horas, contagens) em vez de progresso estimado,
        onde escrever «%» afirmaria o que ninguém mediu.
        """
        pct_norm = max(0, min(100, float(pct)))
        bar_w = w - (2.4 if show_label and label_right else 0)
        self.rect(x, y, bar_w, h, fill=bg)
        filled_w = max(0.12, bar_w * (pct_norm / 100.0)) if pct_norm > 0 else 0
        if filled_w > 0:
            self.rect(x, y, filled_w, h, fill=fill)
        if show_label:
            lbl_x = x + bar_w + 0.35 if label_right else x
            lbl_w = 2.0 if label_right else w
            self.text(label if label else f"{int(pct_norm)}%", lbl_x, y - 0.04, lbl_w, h + 0.08,
                      size=8.5, bold=True, color=INK, wrap=False)
        return y + h

    def kpi_row(self, items, y=CONTENT_TOP, h=2.90, gap=0.42):
        """
        Fila de indicadores. `items` = [(valor, etiqueta), ...] ou
        [(valor, etiqueta, cor_do_valor), ...] ou
        [(valor, etiqueta, cor_do_valor, badge_text, badge_status), ...]

        Composto só a partir de primitivas da marca: cartão branco + barra de
        acento superior (igual à barra da tabela) + número + etiqueta + badge opcional.
        """
        n = len(items)
        if n == 0:
            return y
        w = (CONTENT_W - gap * (n - 1)) / n
        for i, item in enumerate(items):
            value, label = str(item[0]), str(item[1])
            color = item[2] if len(item) > 2 and item[2] else INK
            badge_text = item[3] if len(item) > 3 else None
            badge_status = item[4] if len(item) > 4 else "ok"
            x = CONTENT_L + i * (w + gap)
            self.card(x, y, w, h)
            self.rect(x, y, w, 0.10, fill=ACCENT)
            self.text(value, x + 0.50, y + 0.48, w - 1.00, 1.15,
                      size=22, bold=True, color=color, font=FONT_HEAD, wrap=False)
            self.text(label, x + 0.50, y + 1.80, w - 1.00, 0.55,
                      size=8.5, color=MUTED)
            if badge_text:
                bw = min(3.8, max(2.5, len(badge_text) * 0.26))
                self.status_badge(badge_text, x + w - bw - 0.45, y + 0.50, w=bw, h=0.60, status=badge_status)
        return y + h

    def table(self, headers, rows, y=CONTENT_TOP, widths=None,
              total_row=False, x=CONTENT_L, w=CONTENT_W, pitch=1.32):
        """
        Tabela leve da marca: banda F5F9FC no cabeçalho, barra de acento em cima,
        separadores E7EEF5 entre linhas. `total_row=True` destaca a última linha.

        Devolve o Y a seguir à tabela.
        """
        head_h, inset = 0.97, 0.25
        ncols = len(headers)
        if widths is None:
            widths = [w / ncols] * ncols
        scale = w / sum(widths)
        widths = [cw * scale for cw in widths]

        xs, acc = [], x + inset
        for cw in widths:
            xs.append(acc)
            acc += cw

        # o espaçamento cresce se alguma célula de texto precisar de mais de uma linha
        max_linhas = 1
        for row in rows:
            for i, cell in enumerate(row):
                if i < len(widths) and not _is_num_col(rows, i):
                    max_linhas = max(max_linhas,
                                     estimar_linhas(cell, widths[i] - inset * 2, 8.5))
        pitch = max(pitch, 0.87 + altura_linha(8.5) * max_linhas)

        self.rect(x, y, w, head_h, fill=BAND)
        self.rect(x, y, w, 0.10, fill=ACCENT)
        for i, htxt in enumerate(headers):
            align = PP_ALIGN.RIGHT if i and _is_num_col(rows, i) else PP_ALIGN.LEFT
            self.text(htxt, xs[i], y + 0.25, widths[i] - inset * 2, 0.33,
                      size=7.2, bold=True, color=INK, align=align, wrap=False)

        ry = y + 1.42
        for r, row in enumerate(rows):
            last = total_row and r == len(rows) - 1
            for i, cell in enumerate(row):
                numerica = _is_num_col(rows, i)
                align = PP_ALIGN.RIGHT if i and numerica else PP_ALIGN.LEFT
                altura = altura_linha(8.5) * estimar_linhas(cell, widths[i] - inset * 2, 8.5)
                self.text(cell, xs[i], ry, widths[i] - inset * 2, max(0.41, altura),
                          size=8.5, bold=last, color=INK, align=align,
                          wrap=not numerica, line_spacing=1.1)
            if r < len(rows) - 1:
                self.hline(x, ry + pitch * 0.66, w)
            ry += pitch
        if total_row and rows:
            self.hline(x, y + 1.42 + (len(rows) - 1) * pitch - pitch * 0.34, w,
                       color=ACCENT, pt=1.0)
        return ry - pitch + pitch * 0.66

    def hbars(self, items, y=CONTENT_TOP, label_w=9.0, x=CONTENT_L, w=CONTENT_W,
              pitch=1.15, bar_h=0.50):
        """
        Barras horizontais proporcionais. `items` = [(etiqueta, valor, texto_valor)].
        Usado para distribuição de horas por área.
        """
        if not items:
            return y
        vmax = max(v for _, v, _ in items) or 1
        track_x = x + label_w
        track_w = w - label_w - 4.2
        for i, (label, value, vtext) in enumerate(items):
            by = y + i * pitch
            self.text(label, x, by + 0.06, label_w - 0.4, 0.45, size=9, color=INK, wrap=False)
            self.rect(track_x, by, track_w, bar_h, fill=BAND)
            bw = max(0.12, track_w * (value / vmax))
            self.rect(track_x, by, bw, bar_h, fill=ACCENT)
            self.text(vtext, track_x + track_w + 0.35, by + 0.06, 3.8, 0.45,
                      size=9, bold=True, color=INK, wrap=False)
        return y + len(items) * pitch

    def numbered(self, items, y=CONTENT_TOP, x=CONTENT_L, w=CONTENT_W, pitch=2.54):
        """
        Lista numerada vertical (quadrado azul + título + descrição), tal como o
        slide de próximos passos do ficheiro da marca.
        `items` = [(título, descrição), ...]
        """
        for i, (title, desc) in enumerate(items):
            iy = y + i * pitch
            self.rect(x + 0.15, iy, 0.97, 0.97, fill=ACCENT)
            self.text(str(i + 1), x + 0.15, iy + 0.27, 0.97, 0.28, size=6.5,
                      bold=True, color=WHITE, align=PP_ALIGN.CENTER, wrap=False)
            self.text(title, x + 1.60, iy + 0.15, w - 1.75, 0.64, size=14, bold=True, color=INK)
            if desc:
                self.text(desc, x + 1.60, iy + 1.06, w - 1.75, 0.46, size=8.5, color=MUTED)
        return y + len(items) * pitch

    def steps(self, labels, y=7.75, x=CONTENT_L, w=CONTENT_W):
        """Fila horizontal de passos numerados (slide de processo da marca)."""
        n = len(labels)
        if n == 0:
            return y
        pitch = w / n
        for i, label in enumerate(labels):
            sx = x + i * pitch
            self.rect(sx, y, 1.40, 1.40, fill=ACCENT)
            self.text(str(i + 1), sx, y + 0.35, 1.40, 0.30, size=8, bold=True,
                      color=WHITE, align=PP_ALIGN.CENTER, wrap=False)
            self.text(label, sx + 1.83, y + 0.12, pitch - 2.10, 1.20, size=11,
                      bold=True, color=INK)
        return y + 1.40

    def band(self, txt, y=14.27, x=1.98, w=29.46, h=1.98):
        """Faixa de fecho (mensagem final / contacto)."""
        self.rect(x, y, w, h, fill=BAND)
        self.text(txt, x + 0.56, y + 0.77, w - 1.12, 0.41, size=13, bold=True, color=INK)
        return y + h


def _is_num_col(rows, i):
    """Coluna é numérica se todas as células começarem por dígito."""
    vals = [str(r[i]).strip() for r in rows if i < len(r) and str(r[i]).strip()]
    return bool(vals) and all(v[0].isdigit() for v in vals)


class Deck:
    """Deck PowerPoint com a identidade Parâmetro Global."""

    def __init__(self, doc_label):
        if not BASE_DECK.exists():
            raise FileNotFoundError(f"Deck base da marca não encontrado: {BASE_DECK}")
        self.prs = Presentation(str(BASE_DECK))
        self.doc_label = doc_label
        self._page = 1
        self._layouts = {l.name: l for l in self.prs.slide_layouts}

    def _new(self, layout_name):
        return self.prs.slides.add_slide(self._layouts[layout_name])

    def cover(self, title, subtitle, description="", meta=None, prepared_by="Preparado por Parâmetro Global"):
        """
        Capa oficial da marca: logotipo em cima à esquerda, título azul, subtítulo,
        descrição, bloco de metadados e barra escura no fundo.

        `meta` = lista de linhas, ex.: ["Cliente / Projeto: X", "Data: 27/07/2026"]
        """
        s = _Slide(self._new(LAYOUT_COVER), self)
        s._s.shapes.add_picture(str(LOGO_WORDMARK), _cm(2.49), _cm(1.68), width=_cm(8.59))
        s.text(title, 2.62, 5.49, 17.27, 1.47, size=32, bold=True,
               color=COVER_TITLE, font=FONT_HEAD)
        s.text(subtitle, 2.62, 7.42, 19.50, 1.02, size=22, bold=True,
               color=INK, font=FONT_HEAD)
        if description:
            s.text(description, 2.62, 8.81, 19.50, 1.60, size=16, color=BODY)
        if meta:
            # bloco ancorado por baixo, para nunca colidir com "Preparado por ..."
            bloco_h = len(meta) * 0.46
            s.text("\n".join(meta), 2.62, max(11.60, 14.90 - bloco_h), 20.00, bloco_h,
                   size=10, color=INK, line_spacing=1.3)
        s.text(prepared_by, 2.62, 15.54, 7.62, 0.56, size=9, color=BODY)
        s.rect(0, 18.39, SLIDE_W, 0.66, fill=BAR_DARK)
        self._page += 1
        return s

    def content(self, title, subtitle=""):
        """Slide de conteúdo com cabeçalho, marca, régua e rodapé numerado."""
        s = _Slide(self._new(LAYOUT_CONTENT), self)
        s.text(title, 1.40, 0.97, 24.38, 1.07, size=21, bold=True, color=INK, font=FONT_HEAD)
        if subtitle:
            s.text(subtitle, 1.40, 2.24, 19.05, 0.51, size=6.8, color=MUTED)
        s.hline(_RULE_X, 2.87, _RULE_W)
        s._s.shapes.add_picture(str(LOGO_MARK), _cm(30.91), _cm(0.73), width=_cm(1.51))
        s.hline(_RULE_X, 17.58, _RULE_W)
        s.text(self.doc_label, 1.07, 17.86, 12.00, 0.41, size=6.5, color=MUTED)
        s.text(f"{self._page:02d}", 31.55, 17.78, 0.89, 0.51, size=7, bold=True,
               color=ACCENT, align=PP_ALIGN.RIGHT, wrap=False)
        self._page += 1
        return s

    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path
