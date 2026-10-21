"""Converte um tutorial markdown para Word (.docx) — sem depender de browser/pandoc.

Suporta: H1/H2/H3, paragrafos, listas (- e 1.), tabelas pipe, blockquotes (>),
linhas horizontais (---), e inline **negrito**, *italico*, `codigo`.
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

SRC = sys.argv[1] if len(sys.argv) > 1 else 'docs/TUTORIAL_QC_HMD.md'
OUT = os.path.splitext(SRC)[0] + '.docx'

BLUE = RGBColor(0x2F, 0x54, 0x96)
GREY = RGBColor(0x44, 0x54, 0x6A)
CODE = RGBColor(0xB5, 0x36, 0x0A)

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))')


def add_runs(paragraph, text):
    """Escreve texto com formatacao inline num paragrafo."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*') and part.endswith('*'):
            r = paragraph.add_run(part[1:-1]); r.italic = True
        elif part.startswith('`') and part.endswith('`'):
            r = paragraph.add_run(part[1:-1]); r.font.name = 'Consolas'
            r.font.color.rgb = CODE; r.font.size = Pt(9.5)
        elif part.startswith('['):
            m = re.match(r'\[([^\]]+?)\]\(([^)]+?)\)', part)
            r = paragraph.add_run(m.group(1) if m else part)
            r.font.color.rgb = BLUE
        else:
            paragraph.add_run(part)


def strip_inline(text):
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*([^*]+?)\*', r'\1', text)
    text = re.sub(r'`([^`]+?)`', r'\1', text)
    text = re.sub(r'\[([^\]]+?)\]\([^)]+?\)', r'\1', text)
    return text


doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

with open(SRC, encoding='utf-8') as f:
    lines = f.read().split('\n')

i = 0
while i < len(lines):
    line = lines[i].rstrip()

    # tabela (pipe)
    if line.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|?\s*$', lines[i + 1]):
        header = [c.strip() for c in line.strip('|').split('|')]
        i += 2
        rows = []
        while i < len(lines) and lines[i].lstrip().startswith('|'):
            rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
            i += 1
        t = doc.add_table(rows=1, cols=len(header))
        t.style = 'Light Grid Accent 1'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(header):
            cell = t.rows[0].cells[j]
            cell.paragraphs[0].text = ''
            add_runs(cell.paragraphs[0], h)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in rows:
            cells = t.add_row().cells
            for j in range(len(header)):
                cells[j].paragraphs[0].text = ''
                add_runs(cells[j].paragraphs[0], row[j] if j < len(row) else '')
        doc.add_paragraph()
        continue

    if not line.strip():
        i += 1
        continue

    if line.startswith('# '):
        p = doc.add_paragraph(); r = p.add_run(strip_inline(line[2:]))
        r.bold = True; r.font.size = Pt(22); r.font.color.rgb = BLUE
    elif line.startswith('## '):
        p = doc.add_paragraph(); r = p.add_run(strip_inline(line[3:]))
        r.bold = True; r.font.size = Pt(15); r.font.color.rgb = BLUE
        p.space_before = Pt(12)
    elif line.startswith('### '):
        p = doc.add_paragraph(); r = p.add_run(strip_inline(line[4:]))
        r.bold = True; r.font.size = Pt(12.5); r.font.color.rgb = GREY
    elif line.startswith('#### '):
        p = doc.add_paragraph(); r = p.add_run(strip_inline(line[5:]))
        r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = GREY
    elif line.startswith('> '):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
        add_runs(p, line[2:])
        for run in p.runs:
            run.italic = True; run.font.color.rgb = GREY
    elif re.match(r'^\s*[-*]\s+', line):
        p = doc.add_paragraph(style='List Bullet')
        add_runs(p, re.sub(r'^\s*[-*]\s+', '', line))
    elif re.match(r'^\s*\d+\.\s+', line):
        p = doc.add_paragraph(style='List Number')
        add_runs(p, re.sub(r'^\s*\d+\.\s+', '', line))
    elif re.match(r'^-{3,}$', line):
        doc.add_paragraph()
    else:
        p = doc.add_paragraph()
        add_runs(p, line)
    i += 1

doc.save(OUT)
print(f'DOCX: {OUT} ({os.path.getsize(OUT)//1024} KB)')
