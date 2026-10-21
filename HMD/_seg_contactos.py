import openpyxl, re, unicodedata
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz, process
from collections import defaultdict

SRC = 'comparacao_clientes_odoo_phc.xlsx'
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)


def sheet(name):
    ws = wb[name]
    it = ws.iter_rows(values_only=True)
    h = next(it)
    return h, list(it)


# ---------- normalizacao replicando a Metodologia ----------
LEGAL = {'LDA', 'LTD', 'LTDA', 'GMBH', 'MBH', 'KG', 'SL', 'SLU', 'SRL', 'SARL', 'SA', 'SAS',
         'BV', 'INC', 'CO', 'AG', 'SPA', 'SP', 'PLC', 'OY', 'AB', 'AS', 'NV', 'EURL', 'SC',
         'SCI', 'SNC', 'GES', 'EU', 'UNIPESSOAL', 'SOCIEDADE'}


def norm(s):
    if not s:
        return ''
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode()
    s = s.upper()
    s = re.sub(r'[^A-Z0-9 ]', ' ', s)
    toks = [t for t in s.split() if t and t not in LEGAL]
    return ' '.join(sorted(toks))


hP, phc_miss = sheet('PHC_sem_Odoo')      # 1760
hO, odoo_free = sheet('Odoo_sem_PHC')     # 1507
hC, corr = sheet('Correspondências')
IP = {h: i for i, h in enumerate(hP)}
IO = {h: i for i, h in enumerate(hO)}
IC = {h: i for i, h in enumerate(hC)}

# ---------- pool de candidatos Odoo nao emparelhados ----------
odoo_norm, odoo_ref = [], []
for r in odoo_free:
    nm = r[IO['Odoo Nome']] or r[IO['Odoo Nome Legal']]
    n = norm(nm)
    if n:
        odoo_norm.append(n)
        odoo_ref.append(r)

# sugestoes bidirecionais ja presentes no ficheiro original
back_sugg = defaultdict(list)
for r in odoo_free:
    pn = r[IO['PHC Nº Cliente possível']]
    if pn:
        back_sugg[str(pn).strip()].append((r[IO['Odoo Nome']], r[IO['Score possível']]))

JUNK = re.compile(r'\b(testes?|test|definir|dummy|xxx+|aaa+|prova|exemplo)\b|^\s*[-.\s]+\s*$', re.I)


def is_junk(name):
    return bool(name and JUNK.search(str(name)))


def has_id_data(r):
    return any([r[IP['PHC NIF']], r[IP['PHC Email']], r[IP['PHC Telefone']], r[IP['PHC Telemóvel']]])


bucketA, bucketB, bucketC = [], [], []
WEAK_MIN = 80
for r in phc_miss:
    name = r[IP['PHC Cliente']]
    nr = str(r[IP['PHC Nº Cliente']]).strip() if r[IP['PHC Nº Cliente']] is not None else ''
    nn = norm(name)
    best, best_sc = None, 0
    if nn and odoo_norm:
        m = process.extractOne(nn, odoo_norm, scorer=fuzz.token_sort_ratio)
        if m:
            best_sc = m[1]
            best = odoo_ref[m[2]]
    fs = back_sugg.get(nr, [])
    fs_best = max((s for _, s in fs if s is not None), default=0)
    fs_name = fs[0][0] if fs else None

    cand_score = max(best_sc, fs_best or 0)
    cand_odoo_name = best[IO['Odoo Nome']] if best else fs_name
    cand_odoo_id = best[IO['Odoo ID']] if best else None

    row_out = [r[IP['PHC Nº Cliente']], name, r[IP['PHC NIF']], r[IP['PHC Email']],
               r[IP['PHC Telefone']] or r[IP['PHC Telemóvel']], r[IP['PHC País']],
               round(cand_score, 1), cand_odoo_name, cand_odoo_id,
               round(fs_best or 0, 1), fs_name]

    if is_junk(name) or not has_id_data(r):
        row_out.append('nome teste/placeholder' if is_junk(name) else 'sem NIF/email/telefone')
        bucketC.append(row_out)
    elif best_sc >= WEAK_MIN or (fs_best or 0) >= 75:
        bucketB.append(row_out)
    else:
        bucketA.append(row_out)

bucketB.sort(key=lambda x: -x[6])


# ---------- suspeitos de falso positivo nos matches por nome ----------
def g(r, k):
    v = r[IC[k]]
    return str(v).strip().lower() if v else ''


susp = []
for r in corr:
    if r[IC['Método cruzamento']] != 'Nome semelhante':
        continue
    pc, oc = g(r, 'PHC País'), g(r, 'Odoo País')
    sc = r[IC['Score nome']] or 0
    reason = []
    if pc and oc and pc != oc:
        reason.append('país diferente')
    if sc < 86:
        reason.append(f'score baixo ({sc})')
    if reason:
        susp.append([r[IC['PHC Nº Cliente']], r[IC['PHC Cliente']], r[IC['PHC País']],
                     r[IC['Odoo Nome']], r[IC['Odoo País']], sc, r[IC['Confiança']], '; '.join(reason)])
susp.sort(key=lambda x: x[5])

print("bucketA (alta certeza em falta):", len(bucketA))
print("bucketB (provável match perdido):", len(bucketB))
print("bucketC (ruído/incompleto):", len(bucketC))
print("total:", len(bucketA) + len(bucketB) + len(bucketC), "(esperado 1760)")
print("suspeitos FP nos matches por nome:", len(susp))
print("\nTop 12 bucketB:")
for x in bucketB[:12]:
    print(f"  {x[1]!r} ({x[5]}) ~ {x[7]!r} score={x[6]}")
print("\nAmostra bucketC:")
for x in bucketC[:8]:
    print(f"  {x[1]!r} -> {x[-1]}")

# ================= WRITE OUTPUT WORKBOOK =================
out = openpyxl.Workbook()
H = Font(bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='2F5496')
TITLE = Font(bold=True, size=13)


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.font = H
        cell.fill = FILL
        cell.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'


ws = out.active
ws.title = 'Resumo_segmentacao'
ws['A1'] = 'Segmentação de PHC_sem_Odoo por fiabilidade (refinamento de comparacao_clientes_odoo_phc.xlsx)'
ws['A1'].font = TITLE
data = [
    ('', ''),
    ('Balde', 'Registos'),
    ('A — Em falta (alta certeza): cliente real, com dados identificadores, sem candidato plausível no Odoo', len(bucketA)),
    ('B — Provável match perdido: existe candidato Odoo por nome (rever antes de importar)', len(bucketB)),
    ('C — Ruído / incompleto: nome de teste/placeholder OU sem NIF, email e telefone', len(bucketC)),
    ('TOTAL PHC_sem_Odoo', len(bucketA) + len(bucketB) + len(bucketC)),
    ('', ''),
    ('Controlo cruzado (folha D) — matches por NOME suspeitos de falso positivo', len(susp)),
    ('   se confirmados como FP, somam-se ao balde A (em falta real)', ''),
    ('', ''),
    ('Critérios aplicados', ''),
    ('Normalização de nome', 'sem acentos, maiúsculas, sem pontuação, sem sufixos legais, tokens ordenados (igual à Metodologia original)'),
    ('Candidato "provável match perdido"', 'token_sort_ratio >= 80 vs universo Odoo NÃO emparelhado (1507), OU sugestão bidirecional do ficheiro original >= 75'),
    ('Ruído/incompleto', 'nome com teste/placeholder, ou sem qualquer NIF/email/telefone'),
    ('Suspeito de falso positivo', 'match por "Nome semelhante" com país diferente OU score < 86'),
    ('', ''),
    ('NOTA BLOQUEANTE', 'Esta segmentação assume que a extração Odoo (3550) é o universo correto. Confirmar o filtro da extração antes de importar.'),
]
for r in data:
    ws.append(r)
ws.column_dimensions['A'].width = 95
ws.column_dimensions['B'].width = 70


def write_bucket(title, rows, extra_col=None):
    ws = out.create_sheet(title)
    hdr = ['PHC Nº Cliente', 'PHC Cliente', 'PHC NIF', 'PHC Email', 'PHC Telefone/Telemóvel',
           'PHC País', 'Score candidato Odoo', 'Candidato Odoo (nome)', 'Candidato Odoo ID',
           'Score sugestão ficheiro', 'Sugestão ficheiro (nome)']
    if extra_col:
        hdr = hdr + [extra_col]
    ws.append(hdr)
    style_header(ws, len(hdr))
    for r in rows:
        ws.append(r)
    widths = [14, 40, 14, 30, 26, 14, 14, 40, 42, 16, 40, 28]
    for i, w in enumerate(widths[:len(hdr)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


write_bucket('A_em_falta_alta_certeza', bucketA)
write_bucket('B_provavel_match_perdido', bucketB)
write_bucket('C_ruido_incompleto', bucketC, extra_col='Motivo')

ws = out.create_sheet('D_matches_suspeitos_FP')
ws.append(['PHC Nº Cliente', 'PHC Cliente', 'PHC País', 'Odoo Nome', 'Odoo País',
           'Score nome', 'Confiança', 'Motivo da suspeita'])
style_header(ws, 8)
for r in susp:
    ws.append(r)
for i, w in enumerate([14, 40, 16, 40, 16, 11, 11, 30], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

out.save('comparacao_clientes_segmentacao.xlsx')
print("\nGRAVADO: comparacao_clientes_segmentacao.xlsx")
