import openpyxl, re, unicodedata
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from collections import defaultdict

# ---------------- load ----------------
def load(f, sheet):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    rows = [r for r in it if any(c is not None for c in r)]
    return h, rows

oh, od = load('data_local/odoo_partners.xlsx', 'Sheet1')
ph, pdt = load('data_local/phc_clientes.xlsx', 'Folha1')
OI = {c: i for i, c in enumerate(oh)}
PI = {c: i for i, c in enumerate(ph)}

# ---------------- normalizers ----------------
def clean(v):
    if v is None:
        return ''
    s = str(v).strip()
    if s in ('', '.', '-', '/'):
        return ''
    return s

def vat_alnum(v):
    s = clean(v)
    if not s:
        return None
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().upper()
    s = re.sub(r'[^A-Z0-9]', '', s)
    return s or None

def vat_digits(v):
    a = vat_alnum(v)
    if not a:
        return None
    d = re.sub(r'\D', '', a)
    return d if len(d) >= 8 else None   # >=8 digitos para evitar colisoes curtas

def emails(v):
    s = clean(v)
    if not s:
        return []
    parts = re.split(r'[ ,;/]+', s.lower())
    return [p.strip() for p in parts if '@' in p and '.' in p.split('@')[-1]]

LEGAL = {'LDA','LTD','LTDA','GMBH','MBH','KG','SL','SLU','SRL','SARL','SA','SAS','BV','INC',
         'CO','AG','SPA','SP','PLC','OY','AB','AS','NV','EURL','SC','SCI','SNC','GES','EU',
         'UNIPESSOAL','SOCIEDADE'}
def name_key(s):
    s = clean(s)
    if not s:
        return ''
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().upper()
    s = re.sub(r'[^A-Z0-9 ]', ' ', s)
    toks = [t for t in s.split() if t and t not in LEGAL]
    return ' '.join(sorted(toks))

GENERIC_LOCAL = {'info','geral','mail','email','office','contact','contacto','contacts','admin',
                 'sales','vendas','comercial','commercial','hello','hi','support','suporte',
                 'reception','reservas','booking','general','enquiries','export','compras',
                 'purchasing','accounts','contabilidade','financeiro','rgpd','noreply','no-reply'}
def is_generic_email(e):
    return e.split('@')[0] in GENERIC_LOCAL

# ---------------- build Odoo indexes ----------------
odoo_vat_alnum = defaultdict(list)
odoo_vat_dig = defaultdict(list)
odoo_email = defaultdict(list)
odoo_email_owners = defaultdict(set)   # email -> distinct normalized names
odoo_name_set = set()
for r in od:
    oid = clean(r[OI['id']]); onm = clean(r[OI['name']])
    va = vat_alnum(r[OI['vat']]); vd = vat_digits(r[OI['vat']])
    if va: odoo_vat_alnum[va].append((oid, onm))
    if vd: odoo_vat_dig[vd].append((oid, onm))
    for e in emails(r[OI['email']]):
        odoo_email[e].append((oid, onm))
        odoo_email_owners[e].add(name_key(onm))
    nk = name_key(onm)
    if nk: odoo_name_set.add(nk)

# emails partilhados no Odoo por >1 entidade => tratar como genericos/fracos
shared_odoo_email = {e for e, owners in odoo_email_owners.items() if len(owners) > 1}

# odoo name -> representative (id,name) para anotar matches por nome
odoo_name_rep = {}
for r in od:
    nk = name_key(r[OI['name']])
    if nk and nk not in odoo_name_rep:
        odoo_name_rep[nk] = (clean(r[OI['id']]), clean(r[OI['name']]))

# ---------------- classify PHC ----------------
# Prioridade determinística: VAT > email não-genérico > nome exato normalizado
#                            > email genérico (a rever) > em falta
PRES_VAT, PRES_EMAIL, PRES_NAME, AMBIG, MISS_RAW, SEMCHAVE = [], [], [], [], [], []

def phc_view(r, extra=None):
    base = [r[PI['NR_CLIENTE']], clean(r[PI['CLIENTE']]), clean(r[PI['NIF']]),
            clean(r[PI['E-MAIL']]), clean(r[PI['TELEFONE']]), clean(r[PI['TELEMOVEL']]),
            clean(r[PI['PAIS']]), clean(r[PI['LOCALIDADE']]), clean(r[PI['MORADA']])]
    return base + (extra or [])

for r in pdt:
    va = vat_alnum(r[PI['NIF']]); vd = vat_digits(r[PI['NIF']])
    es = emails(r[PI['E-MAIL']])
    nk = name_key(r[PI['CLIENTE']])
    has_key = bool(va or es)

    # 1) VAT match (prioridade)
    vat_hit = None
    if vd and vd in odoo_vat_dig:
        vat_hit = odoo_vat_dig[vd][0]
    elif va and va in odoo_vat_alnum:
        vat_hit = odoo_vat_alnum[va][0]
    if vat_hit:
        PRES_VAT.append(phc_view(r, [vat_hit[0], vat_hit[1], vd or va]))
        continue

    # 2) email forte (não genérico, não partilhado no Odoo)
    strong = [e for e in es if e in odoo_email and e not in shared_odoo_email and not is_generic_email(e)]
    if strong:
        e = strong[0]; hit = odoo_email[e][0]
        PRES_EMAIL.append(phc_view(r, [hit[0], hit[1], e]))
        continue

    # 3) nome exato normalizado (determinístico)
    if nk and nk in odoo_name_rep:
        oid, onm = odoo_name_rep[nk]
        PRES_NAME.append(phc_view(r, [oid, onm, 'nome exato']))
        continue

    # 4) email genérico/partilhado -> a rever
    weak = [e for e in es if e in odoo_email and (e in shared_odoo_email or is_generic_email(e))]
    if weak:
        e = weak[0]; owners = odoo_email[e]
        AMBIG.append(phc_view(r, [';'.join(f"{i}:{n}" for i, n in owners[:4]), e,
                                  'generico' if is_generic_email(e) else 'partilhado no Odoo']))
        continue

    # 5) sem qualquer sinal
    if has_key:
        MISS_RAW.append((r, va, vd, es))
    else:
        SEMCHAVE.append(phc_view(r))

# ---------------- dedup Em_falta (sem duplicar import) ----------------
seen = {}
miss_dedup = []
phc_internal_dup = []
for r, va, vd, es in MISS_RAW:
    keys = []
    if vd: keys.append('V:' + vd)
    elif va: keys.append('A:' + va)
    for e in es: keys.append('E:' + e)
    if not keys:
        keys = ['N:' + name_key(r[PI['CLIENTE']])]
    matched_existing = next((seen[k] for k in keys if k in seen), None)
    if matched_existing is not None:
        miss_dedup[matched_existing][-1] += 1   # incrementa contador de colapsados
        phc_internal_dup.append(phc_view(r, ['dup de NR ' + str(miss_dedup[matched_existing][0])]))
    else:
        idx = len(miss_dedup)
        for k in keys: seen[k] = idx
        miss_dedup.append(phc_view(r, [1]))

# ---------------- reconciliacao ----------------
presentes = len(PRES_VAT)+len(PRES_EMAIL)+len(PRES_NAME)
total = presentes+len(AMBIG)+len(MISS_RAW)+len(SEMCHAVE)
print("PHC total:", len(pdt))
print("PRESENTES (alta confianca):", presentes,
      f"[VAT {len(PRES_VAT)} | email forte {len(PRES_EMAIL)} | nome exato {len(PRES_NAME)}]")
print("A_rever (so email generico/partilhado):", len(AMBIG))
print("EM FALTA (bruto):", len(MISS_RAW), "-> deduplicado:", len(miss_dedup),
      "(colapsados", len(phc_internal_dup), "dups PHC internos)")
print("Sem_chave (sem NIF e sem email e sem nome no Odoo):", len(SEMCHAVE))
print("SOMA particao:", total, "(esperado", len(pdt), ")")

# ================= WRITE WORKBOOK =================
out = openpyxl.Workbook()
H = Font(bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='2F5496'); TITLE = Font(bold=True, size=13)
def style(ws, n):
    for c in range(1, n+1):
        cell = ws.cell(1, c); cell.font = H; cell.fill = FILL; cell.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'

PHC_COLS = ['NR_CLIENTE','CLIENTE','NIF','E-MAIL','TELEFONE','TELEMOVEL','PAIS','LOCALIDADE','MORADA']

ws = out.active; ws.title = 'Resumo'
ws['A1'] = 'PHC vs Odoo — matching determinístico por VAT/email (sem duplicar)'; ws['A1'].font = TITLE
for row in [
    ('', ''),
    ('Resultado', 'Registos PHC'),
    ('PRESENTE no Odoo (alta confiança)', len(PRES_VAT)+len(PRES_EMAIL)+len(PRES_NAME)),
    ('   — por VAT', len(PRES_VAT)),
    ('   — por email forte (não genérico)', len(PRES_EMAIL)),
    ('   — por nome exato normalizado', len(PRES_NAME)),
    ('A REVER — só bate em email genérico/partilhado', len(AMBIG)),
    ('EM FALTA (candidatos a importar, já deduplicados)', len(miss_dedup)),
    ('   — em falta antes de dedup PHC interno', len(MISS_RAW)),
    ('Sem chave — sem NIF, sem email e sem nome no Odoo', len(SEMCHAVE)),
    ('TOTAL PHC', len(pdt)),
    ('', ''),
    ('Duplicados PHC internos colapsados na lista de importação', len(phc_internal_dup)),
    ('', ''),
    ('AVISO DE ÂMBITO (bloqueante)', 'Extração Odoo = 3550 registos, o MESMO recorte filtrado de antes '
     '(sem parent_id/company_type/customer_rank; IDs 46566-125007 com ~74,9k gaps; nenhum ID baixo). '
     'Logo "em falta" = em falta NESTE recorte, nao no res.partner completo. '
     'Para resposta definitiva, exportar res.partner sem filtro.'),
    ('Chaves preenchidas — Odoo', "VAT 19% (682/3550), email 86% (3041/3550)"),
    ('Chaves preenchidas — PHC', "NIF 56% (2283/4106), email 73% (3010/4106)"),
    ('Regra VAT', 'alfanumérico normalizado OU dígitos (>=8) iguais; trata prefixo de país (PT/FR/DE)'),
    ('Regra email forte', 'lowercase; só email não-genérico e não-partilhado no Odoo'),
    ('Regra nome', 'sem acentos/pontuação/sufixos legais, tokens ordenados; igualdade EXATA (sem fuzzy)'),
]:
    ws.append(row)
ws.column_dimensions['A'].width = 70; ws.column_dimensions['B'].width = 95

def sheet_out(title, rows, extra):
    ws = out.create_sheet(title)
    hdr = PHC_COLS + extra
    ws.append(hdr); style(ws, len(hdr))
    for r in rows: ws.append(r)
    for i, w in enumerate([12,40,16,32,20,18,16,22,40] + [28]*len(extra), 1):
        ws.column_dimensions[get_column_letter(i)].width = w

sheet_out('Em_falta', miss_dedup, ['Dups PHC colapsados'])
sheet_out('Presente_VAT', PRES_VAT, ['Odoo ID', 'Odoo Nome', 'Chave VAT'])
sheet_out('Presente_email', PRES_EMAIL, ['Odoo ID', 'Odoo Nome', 'Email match'])
sheet_out('Presente_nome_exato', PRES_NAME, ['Odoo ID', 'Odoo Nome', 'Tipo'])
sheet_out('A_rever_email_generico', AMBIG, ['Odoo (donos do email)', 'Email', 'Tipo'])
sheet_out('Sem_chave', SEMCHAVE, [])
sheet_out('Duplicados_PHC_interno', phc_internal_dup, ['Observação'])

out.save('comparacao_phc_odoo_vat_email.xlsx')
print("\nGRAVADO: comparacao_phc_odoo_vat_email.xlsx")
