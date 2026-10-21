"""
Gera odoo_corrigido.xlsx aplicando as regras de prioridade:
  vendedor              -> backup (fallback: PHC actual; ultimo: Odoo)
  responsavel interno   -> Odoo (com 'Default' tratado como vazio)
  restantes campos      -> PHC actual (fallback: backup; ultimo: Odoo)

Chave imutavel: referencia (int). Encoding artifacts (?, �) preservados tal
qual aparecem na fonte priorizada.
"""
import openpyxl, re, unicodedata
from openpyxl import Workbook

VENDEDOR_NR_TO_USER = {
    1:  'ESSANI® | N/Empresa',
    2:  'ESSANI® | Joel Pereira',
    3:  'ESSANI | Joaquim Cascais',
    4:  'ESSANI | Anselmo',
    5:  'ESSANI | Profit',
    6:  'ESSANI | Comercial Zona Centro',
    7:  'ESSANI | Alcino',
    8:  'ESSANI | Italia Hardware',
    9:  'ESSANI® | Bruno Bessa',
    10: 'ESSANI® | Carlos Santos',
    11: 'Essani | Francisco Navarro',
}

PRICELIST_INT_TO_LABEL = {1:'Tab', 2:'P01', 3:'P02', 4:'RV1', 5:'RV2'}

def norm_text(s):
    if s is None: return ''
    s = str(s).strip()
    s = re.sub(r'^ESSANI\S*\s*\|\s*', '', s, flags=re.I)
    s = unicodedata.normalize('NFKD', s).encode('ascii','ignore').decode('ascii')
    s = re.sub(r'\s+', ' ', s).strip().lower()
    return s

def load(fn, sheet, ref_col, cols):
    wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
    ws = wb[sheet]; it = ws.iter_rows(values_only=True); next(it)
    out = {}
    for r in it:
        ref = r[ref_col]
        if ref is None or (isinstance(ref,str) and not str(ref).strip()): continue
        try: key = int(str(ref).strip())
        except: continue
        out[key] = {k: r[i] for k,i in cols.items()}
    wb.close(); return out

# Carregar todas as colunas que vamos precisar
odoo = load('lista_clientes_odoo.xlsx','Sheet1',1,{
    'id':0,
    'nome':2,'email':3,'telefone':4,'telemovel':5,
    'morada':6,'local':7,'codpost':8,
    'Pais':9,'codigopais':10,'ncont':11,
    'vendedor':12,'nrvendedor':13,'tabelaprecos':14,
    'responsavel_interno_odoo':15,'respint':16,
    'portesgratis':17,'termospagamento':18,'motivoisencao':19,
    'obs':20,'obs_internas':21,
})
phc = load('lista_clientes_phc.xlsx','lista_clientes_phc',0,{
    'nome':1,'ncont':2,'codigopais':3,'pais':4,
    'morada':5,'local':6,'codpost':7,
    'telefone':8,'email':9,'telemovel':10,
    'nrvendedor':11,'vendedor':12,
    'termospagamento':13,'tabelaprecos':14,'portesgratuitos':15,
    'obs':16,'codmotiseimp':17,'motiseimp':18,
    'obsinternas':19,'respinterno':20,
})
bkp = load('backup_phc_20251117.xlsx','clientes',1,{
    'pncont':0,'nome':2,'ncont':3,'pais':4,
    'morada':5,'local':6,'codpost':7,'telefone':8,'email':9,'telemovel':10,
    'vendedor':11,'termospagamento':12,
})

# Mapa nome PHC -> nrvendedor (construido a partir do PHC actual)
phc_vendor_name_to_nr = {}
for r in phc.values():
    nr = r.get('nrvendedor'); v = r.get('vendedor')
    if nr is not None and v:
        phc_vendor_name_to_nr.setdefault(norm_text(v), nr)
# Garantir mapeamentos do catalogo Odoo nao presentes no PHC (improvavel)
for nr, user_name in VENDEDOR_NR_TO_USER.items():
    phc_vendor_name_to_nr.setdefault(norm_text(user_name), nr)

def resolve_vendor_nr_from_name(name):
    if not name: return None
    return phc_vendor_name_to_nr.get(norm_text(name))

def first_non_empty(*vals):
    for v in vals:
        if v is None: continue
        if isinstance(v,str) and not v.strip(): continue
        return v
    return ''

# Universo final: tudo o que esta no PHC actual ou backup (PHC e master para existencia)
# 3 apagados no PHC -> separados
all_in_p_or_b = set(phc) | set(bkp)
removed_in_phc = (set(odoo) & set(bkp)) - set(phc)  # 3 casos (excluidos do ficheiro principal)
to_correct = sorted(all_in_p_or_b - removed_in_phc)

# Header alinhado com a estrutura do Excel Odoo (re-import friendly)
HEADERS = ['id','referencia','nome','email','telefone','telemovel',
           'morada','local','codpost','Pais','codigopais','ncont',
           'vendedor','nrvendedor','tabelaprecos',
           'responsavel interno odoo','respint',
           'portesgratis','termospagamento','motivoisencao',
           'obs','obs internas',
           '_origem_vendedor','_origem_respinterno','_situacao']

rows_out = []
changes_log = []  # detalhe de campos alterados vs Odoo original

for ref in to_correct:
    o = odoo.get(ref) or {}
    p = phc.get(ref) or {}
    b = bkp.get(ref) or {}

    # --- VENDEDOR: backup -> phc -> odoo ---
    nr_chosen = None; vendor_label_chosen = ''; vendor_origin = ''
    if b.get('vendedor'):
        nr_chosen = resolve_vendor_nr_from_name(b['vendedor'])
        if nr_chosen is not None:
            vendor_origin = 'backup'
            vendor_label_chosen = VENDEDOR_NR_TO_USER.get(nr_chosen, b['vendedor'])
        else:
            # nome do backup nao reconhecido
            vendor_origin = 'backup (nome nao mapeado)'
            vendor_label_chosen = b['vendedor']
    if nr_chosen is None and p.get('nrvendedor') is not None:
        nr_chosen = p['nrvendedor']
        vendor_origin = vendor_origin or 'phc_actual'
        vendor_label_chosen = VENDEDOR_NR_TO_USER.get(nr_chosen, p.get('vendedor') or '')
    if nr_chosen is None and o.get('nrvendedor') is not None:
        nr_chosen = o['nrvendedor']
        vendor_origin = 'odoo (fallback)'
        vendor_label_chosen = VENDEDOR_NR_TO_USER.get(nr_chosen, o.get('vendedor') or '')

    # --- RESPONSAVEL INTERNO: Odoo, com Default -> vazio ---
    respint_value = (o.get('respint') or '').strip() if o.get('respint') else ''
    resp_user_label = o.get('responsavel_interno_odoo') or ''
    if norm_text(respint_value) == 'default':
        respint_value = ''
        resp_user_label = ''
    respinterno_origin = 'odoo'
    if not o:
        respint_value = ''
        resp_user_label = ''
        respinterno_origin = 'novo_cliente (vazio)'
    # Normalizar contra catalogo actual de pg_phc_responsavel_interno_name:
    #  - 'Tânia Pacheco' herdado de antes do rename -> 'Tânia'
    #  - 'hu' typo orfao -> 'Hugo Rios' (preencher tambem o m2o)
    key = norm_text(respint_value)
    if key == 'tania pacheco':
        respint_value = 'Tânia'
        respinterno_origin = 'odoo (normalizado: Tânia Pacheco -> Tânia)'
    elif key == 'hu':
        respint_value = 'Hugo Rios'
        if not resp_user_label:
            resp_user_label = 'ESSANI® | Hugo Rios'
        respinterno_origin = 'odoo (normalizado: hu -> Hugo Rios)'

    # --- RESTO: PHC actual -> backup -> odoo ---
    nome  = first_non_empty(p.get('nome'),  b.get('nome'),  o.get('nome'))
    email = first_non_empty(p.get('email'), b.get('email'), o.get('email'))
    telefone  = first_non_empty(p.get('telefone'),  b.get('telefone'),  o.get('telefone'))
    telemovel = first_non_empty(p.get('telemovel'), b.get('telemovel'), o.get('telemovel'))
    morada = first_non_empty(p.get('morada'), b.get('morada'), o.get('morada'))
    local  = first_non_empty(p.get('local'),  b.get('local'),  o.get('local'))
    codpost = first_non_empty(p.get('codpost'), b.get('codpost'), o.get('codpost'))
    pais   = first_non_empty(p.get('pais'),   b.get('pais'),   o.get('Pais'))
    codigopais = first_non_empty(p.get('codigopais'), o.get('codigopais'))
    ncont = first_non_empty(p.get('ncont'), b.get('ncont'), o.get('ncont'))
    termospag = first_non_empty(p.get('termospagamento'), b.get('termospagamento'), o.get('termospagamento'))
    # tabela de precos: PHC int -> Odoo label
    tab_int = p.get('tabelaprecos')
    if isinstance(tab_int,int) and tab_int in PRICELIST_INT_TO_LABEL:
        tabelaprecos = PRICELIST_INT_TO_LABEL[tab_int]
    else:
        tabelaprecos = o.get('tabelaprecos') or 'Tab'
    # portes gratuitos: bool no PHC
    portesgratis = p.get('portesgratuitos') if p.get('portesgratuitos') is not None else o.get('portesgratis', False)
    # motivo isencao: PHC codmotiseimp -> codigo Odoo
    motivoisencao = first_non_empty(p.get('codmotiseimp'), o.get('motivoisencao'))
    obs           = first_non_empty(p.get('obs'), o.get('obs'))
    obs_internas  = first_non_empty(p.get('obsinternas'), o.get('obs_internas'))

    situacao = []
    if ref in odoo and ref in phc and ref in bkp: situacao.append('nos_3')
    elif ref in odoo and ref in phc and ref not in bkp: situacao.append('novo_apos_backup')
    elif ref in phc and ref in bkp and ref not in odoo: situacao.append('importar_para_odoo (PHC+backup)')
    elif ref in phc and ref not in odoo and ref not in bkp: situacao.append('importar_para_odoo (so_PHC)')

    row = [
        o.get('id') or '',           # id - vazio para novos
        ref,
        nome, email, telefone, telemovel,
        morada, local, codpost, pais, codigopais, ncont,
        vendor_label_chosen, nr_chosen or '', tabelaprecos,
        resp_user_label, respint_value,
        portesgratis if portesgratis is not None else False,
        termospag, motivoisencao,
        obs, obs_internas,
        vendor_origin, respinterno_origin, '; '.join(situacao),
    ]
    rows_out.append(row)

    # Log de mudancas vs Odoo original (apenas para clientes existentes em Odoo)
    if o:
        # vendedor
        if o.get('nrvendedor') != nr_chosen and nr_chosen is not None:
            changes_log.append((ref, nome, 'nrvendedor', o.get('nrvendedor'), nr_chosen, vendor_origin))
        # respint
        old_resp = (o.get('respint') or '').strip()
        if old_resp != respint_value:
            changes_log.append((ref, nome, 'respint', old_resp, respint_value, respinterno_origin))
        # tabela precos
        if (o.get('tabelaprecos') or '') != tabelaprecos:
            changes_log.append((ref, nome, 'tabelaprecos', o.get('tabelaprecos'), tabelaprecos, 'phc_actual'))
        # termos pagamento
        if norm_text(o.get('termospagamento')) != norm_text(termospag) and termospag:
            changes_log.append((ref, nome, 'termospagamento', o.get('termospagamento'), termospag, 'phc_actual'))
        # nome
        if norm_text(o.get('nome')) != norm_text(nome) and nome:
            changes_log.append((ref, nome, 'nome', o.get('nome'), nome, 'phc_actual'))
        # email
        if norm_text(o.get('email')) != norm_text(email) and email:
            changes_log.append((ref, nome, 'email', o.get('email'), email, 'phc_actual'))
        # morada
        if norm_text(o.get('morada')) != norm_text(morada) and morada:
            changes_log.append((ref, nome, 'morada', o.get('morada'), morada, 'phc_actual'))

# ---- Escrever workbook ----
wb = Workbook()
ws = wb.active; ws.title='clientes_corrigidos'
ws.append(HEADERS)
for row in rows_out:
    ws.append(row)

# Folha novos clientes (a criar no Odoo)
ws_new = wb.create_sheet('clientes_novos_para_importar')
ws_new.append(HEADERS)
for row in rows_out:
    if not row[0]:  # id vazio = novo
        ws_new.append(row)

# Folha de revisao: clientes apagados no PHC apos backup
ws_rev = wb.create_sheet('clientes_a_revisar_apagados_phc')
ws_rev.append(['referencia','nome_odoo','vendedor_no_backup','observacao'])
for ref in sorted(removed_in_phc):
    o = odoo.get(ref) or {}
    b = bkp.get(ref) or {}
    ws_rev.append([ref, o.get('nome') or b.get('nome'), b.get('vendedor'),
                   'cliente foi apagado/desativado no PHC apos 2025-11-17; decidir se deve permanecer no Odoo'])

# Folha log de mudancas
ws_log = wb.create_sheet('mudancas_aplicadas')
ws_log.append(['referencia','nome','campo','valor_antigo_odoo','valor_novo_proposto','origem'])
for r in changes_log:
    ws_log.append(list(r))

# Folha regras
ws_reg = wb.create_sheet('regras_aplicadas')
ws_reg.append(['regra','detalhe'])
ws_reg.append(['Chave imutavel','referencia (int) - so colunas com este campo preenchido sao processadas'])
ws_reg.append(['Vendedor','prioridade: backup > PHC actual > Odoo. Backup tem so nome, traduzido para nrvendedor via tabela do PHC actual.'])
ws_reg.append(['Responsavel interno','prioridade: Odoo. Valor "Default" tratado como vazio (instrucao do utilizador).'])
ws_reg.append(['Restantes campos','prioridade: PHC actual > backup > Odoo.'])
ws_reg.append(['Tabela de precos','traduzida de int PHC (1..5) para label Odoo (Tab/P01/P02/RV1/RV2).'])
ws_reg.append(['Clientes apagados no PHC apos backup','excluidos do ficheiro principal, listados em "clientes_a_revisar_apagados_phc".'])
ws_reg.append(['Clientes so no PHC actual ou PHC+backup','incluidos com id vazio para criar novos em Odoo, tambem listados em "clientes_novos_para_importar".'])
ws_reg.append(['Encoding','caracteres ? e � vindos das fontes sao preservados; nao foram substituidos.'])
ws_reg.append([])
ws_reg.append(['CATALOGO VENDEDOR (res.users.x_studio_salesperson_id_1)'])
for nr in sorted(VENDEDOR_NR_TO_USER): ws_reg.append([nr, VENDEDOR_NR_TO_USER[nr]])
ws_reg.append([])
ws_reg.append(['CATALOGO RESPONSAVEL INTERNO (res.users.pg_phc_responsavel_interno_name)'])
for phc_name, user_label in [
    ('Ana Martins', 'ESSANI® | Ana Martins'),
    ('Bruno Bessa', 'ESSANI® | Bruno Bessa'),
    ('Carlos',      'ESSANI® | Carlos Santos'),
    ('Egídio',      'ESSANI® | Egídio Santos'),
    ('Elsa',        'ESSANI® | Elsa Silva'),
    ('Hugo Rios',   'ESSANI® | Hugo Rios'),
    ('Hugo Rios',   'ESSANI® | Hugo Rios Export'),
    ('Nuno',        'ESSANI® | Nuno Leão'),
    ('Tânia',       'ESSANI® | Tânia Pacheco'),
    ('Armazém',     'ESSANI® | João Neves'),
]:
    ws_reg.append([phc_name, user_label])

wb.save('odoo_corrigido.xlsx')

# Sumario
total = len(rows_out)
novos = sum(1 for r in rows_out if not r[0])
existentes = total - novos
print(f'[OK] odoo_corrigido.xlsx')
print(f'Total: {total} clientes')
print(f'  - {existentes} a actualizar em Odoo (id preenchido)')
print(f'  - {novos} novos a criar em Odoo (id vazio)')
print(f'  - {len(removed_in_phc)} clientes a revisar (apagados no PHC apos backup)')
print(f'  - {len(changes_log)} mudancas registadas vs Odoo original')

# Detalhe das mudancas
from collections import Counter
fc = Counter(r[2] for r in changes_log)
print('\nMudancas por campo:')
for f, n in fc.most_common():
    print(f'  {f:25s} {n}')
