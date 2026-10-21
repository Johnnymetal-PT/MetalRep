"""
Comparacao Odoo / PHC actual / backup PHC, com referencia (int) como chave imutavel.

API (pg_phc_api/models/res_partner.py):
  - vendedor: transporta vendedornr (int). Nomes diferem por desenho.
  - respinterno: transporta a string pg_phc_responsavel_interno_name do user.
  - tabela de precos: Odoo Selection Tab/P01/P02/RV1/RV2 -> 1..5.

Mapeamento autoritativo extraido dos screenshots da configuracao Odoo:
  - VENDEDOR_NR_TO_USER: nrvendedor 1..11 -> nome do user Odoo
  - RESPINT_VALID_NAMES: lista dos pg_phc_responsavel_interno_name validos
"""
import openpyxl, re, unicodedata
from openpyxl import Workbook

# ---- Configuracao Odoo (screenshots fornecidos pelo Joao) ----
VENDEDOR_NR_TO_USER = {
    1:  'ESSANI® | N/Empresa',
    2:  'ESSANI® | Joel Pereira',
    3:  'ESSANI | Joaquim Cascais',
    4:  'ESSANI | Anselmo',
    5:  'ESSANI | Profit',
    6:  'ESSANI | Comercial Zona Centro',
    7:  'ESSANI | Alcino',
    8:  'ESSANI | Italia Hardware',
    9:  'ESSANI® | Bruno Bessa',
    10: 'ESSANI® | Carlos Santos',
    11: 'Essani | Francisco Navarro',
}

# Valores autorizados para pg_phc_responsavel_interno_name (catalogo do utilizador)
RESPINT_VALID_NAMES = {
    'Ana Martins', 'Armazém', 'Bruno Bessa', 'Carlos',
    'Egídio', 'Elsa', 'Hugo Rios', 'Nuno', 'Tânia',
}
RESPINT_VALID_NORM = {unicodedata.normalize('NFKD', n).encode('ascii','ignore').decode().lower() for n in RESPINT_VALID_NAMES}

PRICELIST_LABEL_TO_CODE = {'tab':'1','p01':'2','p02':'3','rv1':'4','rv2':'5'}
PRICELIST_CODES = {'1','2','3','4','5'}

def norm(s):
    if s is None: return ''
    s = str(s).strip()
    if not s: return ''
    s = re.sub(r'^ESSANI\S*\s*\|\s*', '', s, flags=re.I)
    s = unicodedata.normalize('NFKD', s).encode('ascii','ignore').decode('ascii')
    s = re.sub(r'\s+', ' ', s).strip().lower()
    return s

def norm_pricelist(v):
    if v is None: return ''
    s = str(v).strip()
    if not s: return ''
    if s in PRICELIST_CODES: return s
    return PRICELIST_LABEL_TO_CODE.get(s.lower(), '')

def load(fn, sheet, ref_col, cols):
    wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
    ws = wb[sheet]; it = ws.iter_rows(values_only=True); next(it)
    out = {}; nulls = 0
    for r in it:
        ref = r[ref_col]
        if ref is None or (isinstance(ref,str) and not str(ref).strip()):
            nulls += 1; continue
        try: key = int(str(ref).strip())
        except: nulls += 1; continue
        out[key] = {k: r[i] for k,i in cols.items()}
    wb.close(); return out, nulls

odoo, odoo_null = load('lista_clientes_odoo.xlsx','Sheet1',1,
    {'nome':2,'ncont':11,'vendedor':12,'nrvendedor':13,'tabelaprecos':14,'respint':16,'termospagamento':18})
phc, phc_null = load('lista_clientes_phc.xlsx','lista_clientes_phc',0,
    {'nome':1,'ncont':2,'nrvendedor':11,'vendedor':12,'termospagamento':13,'tabelaprecos':14,'respinterno':20})
bkp, bkp_null = load('backup_phc_20251117.xlsx','clientes',1,
    {'nome':2,'ncont':3,'vendedor':11,'termospagamento':12})

in_o = set(odoo); in_p = set(phc); in_b = set(bkp)
all_refs = sorted(in_o | in_p | in_b)
all_three = in_o & in_p & in_b
only_p = in_p - in_o - in_b
only_o = in_o - in_p - in_b
only_b = in_b - in_o - in_p
op_not_b = (in_o & in_p) - in_b
ob_not_p = (in_o & in_b) - in_p
pb_not_o = (in_p & in_b) - in_o

def user_for_nr(nr):
    return VENDEDOR_NR_TO_USER.get(nr, f'<nr={nr} sem mapeamento>') if nr is not None else '<vazio>'

# Vendedor REAL (por nrvendedor) com utilizador resolvido
vendor_id_diff = []
for ref in sorted(in_o & in_p):
    a = odoo[ref].get('nrvendedor'); b = phc[ref].get('nrvendedor')
    if a is None or b is None: continue
    if a != b:
        vendor_id_diff.append({
            'referencia': ref,
            'nome': odoo[ref].get('nome') or phc[ref].get('nome'),
            'nrvendedor_odoo': a, 'user_odoo': user_for_nr(a), 'label_odoo': odoo[ref].get('vendedor'),
            'nrvendedor_phc': b,  'user_esperado_para_nr_phc': user_for_nr(b), 'label_phc': phc[ref].get('vendedor'),
        })

# Backup PHC -> PHC actual: vendedor alterado
phc_vendor_changes = []
for ref in sorted(in_p & in_b):
    vp = phc[ref].get('vendedor'); vb = bkp[ref].get('vendedor')
    if norm(vp) != norm(vb) and norm(vp) and norm(vb):
        phc_vendor_changes.append({
            'referencia': ref, 'nome': phc[ref].get('nome') or bkp[ref].get('nome'),
            'vendedor_backup': vb, 'vendedor_phc_actual': vp,
            'nrvendedor_phc_actual': phc[ref].get('nrvendedor'),
            'nrvendedor_odoo': odoo.get(ref,{}).get('nrvendedor') if ref in odoo else None,
        })

# Tabela de precos
pricelist_diff = []
for ref in sorted(in_o & in_p):
    co = norm_pricelist(odoo[ref].get('tabelaprecos'))
    cp = norm_pricelist(phc[ref].get('tabelaprecos'))
    if co and cp and co != cp:
        pricelist_diff.append({
            'referencia': ref, 'nome': odoo[ref].get('nome'),
            'tabela_odoo': odoo[ref].get('tabelaprecos'), 'tabela_phc': phc[ref].get('tabelaprecos'),
            'codigo_odoo': co, 'codigo_phc': cp,
        })

# Termos pag.
payterm_diff = []
for ref in sorted(in_o & in_p):
    tpo = norm(odoo[ref].get('termospagamento'))
    tpp = norm(phc[ref].get('termospagamento'))
    if tpo and tpp and tpo != tpp:
        payterm_diff.append({
            'referencia': ref, 'nome': odoo[ref].get('nome'),
            'termos_odoo': odoo[ref].get('termospagamento'),
            'termos_phc': phc[ref].get('termospagamento'),
        })

payterm_phc_changes = []
for ref in sorted(in_p & in_b):
    tpp = norm(phc[ref].get('termospagamento'))
    tpb = norm(bkp[ref].get('termospagamento'))
    if tpp and tpb and tpp != tpb:
        payterm_phc_changes.append({
            'referencia': ref, 'nome': phc[ref].get('nome') or bkp[ref].get('nome'),
            'termos_backup': bkp[ref].get('termospagamento'),
            'termos_phc_actual': phc[ref].get('termospagamento'),
        })

# Respinterno: comparar contra catalogo de nomes validos
resp_diff = []
odoo_anomalous_values = {}  # respint que nao corresponde a nenhum user configurado
phc_anomalous_values = {}   # respinterno que nao corresponde a nenhum user configurado
for ref in sorted(in_o & in_p):
    ro_raw = odoo[ref].get('respint'); rp_raw = phc[ref].get('respinterno')
    ro = norm(ro_raw); rp = norm(rp_raw)
    if ro and ro not in RESPINT_VALID_NORM:
        odoo_anomalous_values[ro_raw] = odoo_anomalous_values.get(ro_raw, 0) + 1
    if rp and rp not in RESPINT_VALID_NORM:
        phc_anomalous_values[rp_raw] = phc_anomalous_values.get(rp_raw, 0) + 1
    if ro == rp: continue
    if not ro and not rp: continue
    o_first = ro.split(' ')[0] if ro else ''
    p_first = rp.split(' ')[0] if rp else ''
    if ro == 'default' and rp and rp != 'default':
        cat = 'odoo_default_phc_tem_pessoa'
    elif ro and ro != 'default' and not rp:
        cat = 'odoo_tem_pessoa_phc_vazio'
    elif o_first and p_first and o_first == p_first:
        cat = 'mesmo_primeiro_nome (formato)'
    elif ro == 'default' and not rp:
        cat = 'odoo_default_phc_vazio'
    else:
        cat = 'pessoa_diferente (REAL)'
    resp_diff.append({
        'referencia': ref, 'nome': odoo[ref].get('nome'),
        'respint_odoo': ro_raw, 'respinterno_phc': rp_raw, 'categoria': cat,
        'odoo_valor_valido_no_catalogo': 'sim' if ro in RESPINT_VALID_NORM else ('na' if not ro else 'NAO'),
        'phc_valor_valido_no_catalogo':  'sim' if rp in RESPINT_VALID_NORM else ('na' if not rp else 'NAO'),
    })

# Vendedor por nrvendedor: detectar nr fora do catalogo (1..11)
vendor_nr_invalid_odoo = {}
vendor_nr_invalid_phc  = {}
for ref,r in odoo.items():
    nr = r.get('nrvendedor')
    if nr is not None and nr not in VENDEDOR_NR_TO_USER:
        vendor_nr_invalid_odoo[nr] = vendor_nr_invalid_odoo.get(nr,0)+1
for ref,r in phc.items():
    nr = r.get('nrvendedor')
    if nr is not None and nr not in VENDEDOR_NR_TO_USER:
        vendor_nr_invalid_phc[nr] = vendor_nr_invalid_phc.get(nr,0)+1

# WORKBOOK
wb = Workbook()
ws = wb.active; ws.title='resumo'
ws.append(['metrica','valor','nota'])
ws.append(['Odoo - clientes', len(odoo), f'{odoo_null} ref nula descartada'])
ws.append(['PHC actual', len(phc), ''])
ws.append(['Backup PHC (2025-11-17)', len(bkp), 'so tem vendedor + termospagamento'])
ws.append([])
ws.append(['PRESENCA'])
ws.append(['Refs nos 3 ficheiros', len(all_three), ''])
ws.append(['So no PHC actual', len(only_p), 'nunca chegaram ao Odoo nem ao backup'])
ws.append(['So no Odoo', len(only_o), ''])
ws.append(['So no backup', len(only_b), ''])
ws.append(['Em Odoo+PHC, ausente backup', len(op_not_b), 'criados apos 2025-11-17'])
ws.append(['Em Odoo+backup, ausente PHC', len(ob_not_p), 'apagados no PHC apos 2025-11-17'])
ws.append(['Em PHC+backup, ausente Odoo', len(pb_not_o), 'nunca exportados para Odoo'])
ws.append([])
ws.append(['VENDEDOR (validado contra catalogo de 11 users Odoo)'])
ws.append(['nrvendedor diferente Odoo vs PHC', len(vendor_id_diff), 'REAL'])
ws.append(['Backup -> PHC actual: vendedor alterado', len(phc_vendor_changes), ''])
ws.append(['Odoo: clientes com nrvendedor fora do catalogo', sum(vendor_nr_invalid_odoo.values()), str(vendor_nr_invalid_odoo)])
ws.append(['PHC: clientes com nrvendedor fora do catalogo', sum(vendor_nr_invalid_phc.values()),  str(vendor_nr_invalid_phc)])
ws.append([])
ws.append(['TABELA DE PRECOS (Odoo vs PHC, mapeado)'])
ws.append(['Discrepancias', len(pricelist_diff), ''])
ws.append([])
ws.append(['TERMOS DE PAGAMENTO'])
ws.append(['Odoo vs PHC actual', len(payterm_diff), ''])
ws.append(['Backup -> PHC actual', len(payterm_phc_changes), ''])
ws.append([])
ws.append(['RESPONSAVEL INTERNO (validado contra catalogo de nomes configurados em res.users)'])
for cat in ['pessoa_diferente (REAL)','odoo_default_phc_tem_pessoa','odoo_tem_pessoa_phc_vazio','odoo_default_phc_vazio','mesmo_primeiro_nome (formato)']:
    n = sum(1 for d in resp_diff if d['categoria']==cat)
    ws.append([cat, n, ''])
ws.append([])
ws.append(['ANOMALIAS DE CATALOGO RESPINTERNO'])
ws.append(['Valores em Odoo respint NAO presentes no catalogo', sum(odoo_anomalous_values.values()), str(odoo_anomalous_values)])
ws.append(['Valores em PHC respinterno NAO presentes no catalogo', sum(phc_anomalous_values.values()), str(phc_anomalous_values)])

# presenca
ws2 = wb.create_sheet('presenca')
ws2.append(['referencia','nome','odoo','phc_actual','backup_phc','categoria'])
for ref in all_refs:
    o=ref in in_o; p=ref in in_p; b=ref in in_b
    if o and p and b: continue
    nome = ((odoo.get(ref) or phc.get(ref) or bkp.get(ref)) or {}).get('nome')
    if   o and p and not b: c='novo_apos_backup'
    elif o and b and not p: c='apagado_no_phc_apos_backup'
    elif p and b and not o: c='nunca_exportado_para_odoo'
    elif p and not o and not b: c='so_no_phc_actual'
    elif o and not p and not b: c='so_no_odoo'
    elif b and not o and not p: c='so_no_backup'
    else: c='?'
    ws2.append([ref, nome, 'X' if o else '', 'X' if p else '', 'X' if b else '', c])

ws3 = wb.create_sheet('vendedor_diff_REAL')
ws3.append(['referencia','nome','nrvendedor_odoo','user_odoo','label_excel_odoo','nrvendedor_phc','user_esperado_para_nr_phc','label_excel_phc'])
for d in vendor_id_diff:
    ws3.append([d['referencia'],d['nome'],d['nrvendedor_odoo'],d['user_odoo'],d['label_odoo'],
                d['nrvendedor_phc'],d['user_esperado_para_nr_phc'],d['label_phc']])

ws4 = wb.create_sheet('phc_vendedor_alterado')
ws4.append(['referencia','nome','vendedor_backup_17_11_2025','vendedor_phc_actual','nrvendedor_phc_actual','user_esperado_para_nr','nrvendedor_odoo','alinhado_com_odoo?'])
for d in phc_vendor_changes:
    nr_phc = d['nrvendedor_phc_actual']
    nr_od  = d['nrvendedor_odoo']
    aligned = 'sim' if nr_phc is not None and nr_od is not None and nr_phc == nr_od else ('odoo_nao_tem_cliente' if nr_od is None else 'NAO')
    ws4.append([d['referencia'],d['nome'],d['vendedor_backup'],d['vendedor_phc_actual'],
                nr_phc, user_for_nr(nr_phc), nr_od, aligned])

ws5 = wb.create_sheet('tabela_precos_diff')
ws5.append(['referencia','nome','tabela_odoo','tabela_phc','codigo_odoo','codigo_phc'])
for d in pricelist_diff:
    ws5.append([d['referencia'],d['nome'],d['tabela_odoo'],d['tabela_phc'],d['codigo_odoo'],d['codigo_phc']])

ws6 = wb.create_sheet('termos_pag_diff')
ws6.append(['referencia','nome','termos_odoo','termos_phc_actual'])
for d in payterm_diff:
    ws6.append([d['referencia'],d['nome'],d['termos_odoo'],d['termos_phc']])

ws7 = wb.create_sheet('phc_termos_pag_alterado')
ws7.append(['referencia','nome','termos_backup','termos_phc_actual'])
for d in payterm_phc_changes:
    ws7.append([d['referencia'],d['nome'],d['termos_backup'],d['termos_phc_actual']])

ws8 = wb.create_sheet('respinterno_diff')
ws8.append(['referencia','nome','respint_odoo','respinterno_phc','categoria','odoo_valor_no_catalogo','phc_valor_no_catalogo'])
for d in sorted(resp_diff, key=lambda x: (x['categoria'], x['referencia'])):
    ws8.append([d['referencia'],d['nome'],d['respint_odoo'],d['respinterno_phc'],d['categoria'],
                d['odoo_valor_valido_no_catalogo'], d['phc_valor_valido_no_catalogo']])

ws9 = wb.create_sheet('catalogo_referencia')
ws9.append(['VENDEDORES (configurados em res.users.x_studio_salesperson_id_1)'])
ws9.append(['nrvendedor','user_odoo'])
for nr, u in sorted(VENDEDOR_NR_TO_USER.items()):
    ws9.append([nr, u])
ws9.append([])
ws9.append(['RESPONSAVEIS INTERNOS (configurados em res.users.pg_phc_responsavel_interno_name)'])
for n in sorted(RESPINT_VALID_NAMES):
    ws9.append([n])

wb.save('discrepancias_clientes.xlsx')

print('[OK] discrepancias_clientes.xlsx')
print()
print('=== PRESENCA ===')
print(f'  {len(only_p):4d} so PHC actual | {len(op_not_b):3d} novos pos-backup | {len(ob_not_p):3d} removidos | {len(pb_not_o):3d} PHC+bkp sem Odoo')
print()
print('=== VENDEDOR ===')
print(f'  {len(vendor_id_diff)} nrvendedor diferente Odoo vs PHC')
print(f'  {sum(vendor_nr_invalid_odoo.values())} clientes Odoo com nrvendedor fora do catalogo 1..11: {vendor_nr_invalid_odoo}')
print(f'  {sum(vendor_nr_invalid_phc.values())} clientes PHC  com nrvendedor fora do catalogo 1..11: {vendor_nr_invalid_phc}')
print(f'  {len(phc_vendor_changes)} alteracoes backup -> PHC actual')
print()
print('=== TABELA DE PRECOS ===')
print(f'  {len(pricelist_diff)} discrepancias Odoo vs PHC')
print()
print('=== TERMOS PAG. ===')
print(f'  Odoo vs PHC: {len(payterm_diff)}   backup -> PHC actual: {len(payterm_phc_changes)}')
print()
print('=== RESPINTERNO ===')
for cat in ['pessoa_diferente (REAL)','odoo_default_phc_tem_pessoa','odoo_tem_pessoa_phc_vazio','odoo_default_phc_vazio','mesmo_primeiro_nome (formato)']:
    n = sum(1 for d in resp_diff if d['categoria']==cat)
    print(f'  {n:4d} {cat}')
print('  Valores Odoo respint NAO no catalogo:', odoo_anomalous_values)
print('  Valores PHC respinterno NAO no catalogo:', phc_anomalous_values)
