"""Add a pg_responsavel_interno_user_id column to odoo_corrigido.xlsx.

The original ``respint`` column holds the PHC user name (e.g. ``Tânia``).
The Odoo import wizard cannot resolve those values directly to res.users
because:
  * the Char field ``responsavel_interno`` is now related/readonly;
  * the Many2one ``pg_responsavel_interno_user_id`` would need the Odoo
    display_name (``ESSANI® | Tânia Pacheco``), not the PHC short name.

This script writes a new column with the exact Odoo display_name so the
wizard can resolve every row via name_search. The user maps the new column
to ``pg_responsavel_interno_user_id`` and skips the original ``respint``.

It also normalises the U+FFFD replacement char that appears in the source
file (e.g. ``T�nia`` -> ``Tânia``).

Output: odoo_corrigido_remapped.xlsx (original left untouched).
"""

import shutil

import openpyxl


SRC = "odoo_corrigido.xlsx"
DST = "odoo_corrigido_remapped.xlsx"
SHEET = "clientes_corrigidos"

NEW_HEADER = "pg_responsavel_interno_user_id"

# Replacement-char normalisation seen in the source file:
RESPINT_NORMALIZE = {
    "T�nia": "Tânia",
}

# Final mapping PHC name -> Odoo display_name.
# "Hugo Rios" defaults to "ESSANI® | Hugo Rios" (base account); rows that
# should point to "Hugo Rios Export" need manual review after import.
PHC_TO_ODOO_DISPLAY_NAME = {
    "Ana Martins": "ESSANI® | Ana Martins",
    "Bruno Bessa": "ESSANI® | Bruno Bessa",
    "Carlos": "ESSANI® | Carlos Santos",
    "Egídio": "ESSANI® | Egídio Santos",
    "Elsa": "ESSANI® | Elsa Silva",
    "Hugo Rios": "ESSANI® | Hugo Rios",
    "Nuno": "ESSANI® | Nuno Leão",
    "Tânia": "ESSANI® | Tânia Pacheco",
    "Armazém": "ESSANI® | João Neves",
}


def main():
    shutil.copy(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    ws = wb[SHEET]

    headers = [c.value for c in ws[1]]
    idx_respint = headers.index("respint") + 1

    new_col = ws.max_column + 1
    ws.cell(row=1, column=new_col, value=NEW_HEADER)

    total = 0
    mapped = 0
    unmapped = {}
    for row in range(2, ws.max_row + 1):
        raw = ws.cell(row=row, column=idx_respint).value
        if not raw:
            continue
        total += 1
        normalized = RESPINT_NORMALIZE.get(str(raw).strip(), str(raw).strip())
        target = PHC_TO_ODOO_DISPLAY_NAME.get(normalized)
        if target:
            ws.cell(row=row, column=new_col, value=target)
            mapped += 1
        else:
            unmapped[normalized] = unmapped.get(normalized, 0) + 1

    wb.save(DST)

    print(f"Wrote {DST}")
    print(f"Rows with respint: {total}")
    print(f"  mapped:   {mapped}")
    print(f"  unmapped: {total - mapped}")
    if unmapped:
        print("Unmapped distinct values (need manual handling):")
        for k, c in sorted(unmapped.items(), key=lambda x: -x[1]):
            print(f"  {c:5d}x  {k!r}")


if __name__ == "__main__":
    main()
